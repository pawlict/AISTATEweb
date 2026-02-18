
from __future__ import annotations

import io
import asyncio
import json
import os
import shutil
import sys
import threading
import time
import tempfile
import uuid
import zipfile
from dataclasses import dataclass, asdict, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from collections import deque

from fastapi import FastAPI, File, Form, UploadFile, Request, HTTPException, Body
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.concurrency import run_in_threadpool

from backend.settings import APP_NAME, APP_VERSION, AUTHOR_EMAIL
from backend.settings_store import load_settings, save_settings, _local_config_dir
from backend.legacy_adapter import diarize_text_simple
from backend.prompts.manager import PromptManager

# --- Multi-user auth system ---
from webapp.auth.deployment_store import DeploymentStore
from webapp.auth.user_store import UserStore
from webapp.auth.session_store import SessionStore
from webapp.auth.permissions import (
    get_user_modules, is_route_allowed,
    PUBLIC_ROUTES, PUBLIC_PREFIXES,
)
from webapp.routers import auth as auth_router
from webapp.routers import users as users_router
from webapp.routers import setup as setup_router

# Analysis: documents + Ollama
from backend.document_processor import extract_text, DocumentProcessingError, SUPPORTED_EXTS
from backend.ollama_client import OllamaClient, OllamaError, quick_analyze, deep_analyze, stream_analyze

# Analysis report generator (HTML/DOCX/MD)
from backend.report_generator import save_report, ReportSaveError

# Financial intelligence pipeline
from backend.finance.pipeline import run_finance_pipeline, run_multi_statement_pipeline, build_enriched_prompt as build_finance_enriched_prompt
from backend.models_info import MODELS_INFO, MODELS_GROUPS, DEFAULT_MODELS

from generators import generate_txt_report, generate_html_report, generate_pdf_report

# Routers (refactored modules)
from webapp.routers import chat as chat_router
from webapp.routers import admin as admin_router
from webapp.routers import tasks as tasks_router
from webapp.routers import aml as aml_router
from webapp.routers import messages as messages_router
from webapp.routers import workspaces as workspaces_router

try:
    from markdown import markdown as md_to_html  # type: ignore
except Exception:  # pragma: no cover
    md_to_html = None


ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = Jinja2Templates(directory=str(Path(__file__).resolve().parent / "templates"))

DATA_DIR = Path(os.environ.get("AISTATEWEB_DATA_DIR") or os.environ.get("AISTATEWWW_DATA_DIR") or os.environ.get("AISTATE_DATA_DIR") or str(ROOT / "data_www")).resolve()
PROJECTS_DIR = DATA_DIR / "projects"
PROJECTS_DIR.mkdir(parents=True, exist_ok=True)
PROMPTS = PromptManager(PROJECTS_DIR)

# --- Multi-user auth stores ---
from webapp.auth.message_store import MessageStore
from webapp.auth.audit_store import AuditStore

from webapp.auth.passwords import init_blacklist

_AUTH_CONFIG_DIR = _local_config_dir()
DEPLOYMENT_STORE = DeploymentStore(_AUTH_CONFIG_DIR)
USER_STORE = UserStore(_AUTH_CONFIG_DIR)
SESSION_STORE = SessionStore(_AUTH_CONFIG_DIR)
MESSAGE_STORE = MessageStore(_AUTH_CONFIG_DIR)
AUDIT_STORE = AuditStore(_AUTH_CONFIG_DIR)
PASSWORD_BLACKLIST = init_blacklist(_AUTH_CONFIG_DIR)

from webapp.auth.workspace_store import WorkspaceStore
WORKSPACE_STORE = WorkspaceStore()


def _get_session_timeout() -> int:
    """Read configurable session timeout (hours) from settings."""
    try:
        s = load_settings()
        return int(getattr(s, "session_timeout_hours", 8) or 8)
    except Exception:
        return 8

# ---------------------------
# Admin file logs (persistent)
# ---------------------------
# These logs are NOT stored inside projects. They are intended for administrators.
# Default location: AISTATEweb/backend/logs/YYYY-MM-DD/aistateHH-HH.log
# Override with env var: AISTATEWEB_ADMIN_LOG_DIR

ADMIN_LOG_DIR = Path(os.environ.get("AISTATEWEB_ADMIN_LOG_DIR") or str(ROOT / "backend" / "logs")).resolve()

class _AdminFileLogger:
    """Hourly file logger with per-day folders (simple, dependency-free).

    Output example:
      backend/logs/2026-01-17/aistate14-15.log
      backend/logs/2026-01-17/aistate_user14-15.log  (paired audit log)
    """

    def __init__(self, root_dir: Path, prefix: str = "aistate", pair: Optional["_AdminFileLogger"] = None) -> None:
        self.root_dir = Path(root_dir)
        self._prefix = prefix
        self._pair = pair
        self._lock = threading.Lock()
        self._cur_key: Optional[str] = None
        self._fh: Optional[io.TextIOWrapper] = None

    def _target_path(self, dt: datetime) -> Path:
        date_str = dt.strftime("%Y-%m-%d")
        h0 = dt.hour
        h1 = (dt + timedelta(hours=1)).hour
        folder = self.root_dir / date_str
        return folder / f"{self._prefix}{h0:02d}-{h1:02d}.log"

    def _ensure_pair(self, dt: datetime) -> None:
        """Ensure the paired logger also has a file for this hour."""
        if not self._pair:
            return
        try:
            pair_path = self._pair._target_path(dt)
            if not pair_path.exists():
                pair_path.parent.mkdir(parents=True, exist_ok=True)
                pair_path.write_text("No events\n", encoding="utf-8")
        except Exception:
            pass

    def write_line(self, line: str) -> None:
        line = (line or "").rstrip("\n")
        if not line:
            return

        dt = datetime.now()
        key = dt.strftime("%Y-%m-%d_%H")

        try:
            with self._lock:
                if self._fh is None or self._cur_key != key:
                    try:
                        if self._fh:
                            self._fh.flush()
                            self._fh.close()
                    except Exception:
                        pass

                    path = self._target_path(dt)
                    path.parent.mkdir(parents=True, exist_ok=True)
                    # If file only has "No events" placeholder, overwrite it
                    mode = "a"
                    if path.exists():
                        try:
                            content = path.read_text(encoding="utf-8").strip()
                            if content == "No events":
                                mode = "w"
                        except Exception:
                            pass
                    self._fh = open(path, mode, encoding="utf-8", errors="replace")
                    self._cur_key = key

                    # Ensure paired log file exists for this hour
                    self._ensure_pair(dt)

                assert self._fh is not None
                self._fh.write(line + "\n")
                self._fh.flush()
        except Exception:
            # Never break app flow because of logging.
            return

    def ensure_file_for_hour(self, dt: Optional[datetime] = None) -> None:
        """Create the file for the current hour if it doesn't exist (with 'No events' placeholder)."""
        dt = dt or datetime.now()
        try:
            path = self._target_path(dt)
            if not path.exists():
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text("No events\n", encoding="utf-8")
        except Exception:
            pass


# System log writes to aistateHH-HH.log, audit log writes to aistate_userHH-HH.log.
# They are "paired" — when one creates a new hourly file, the other is also created.
AUDIT_FILE_LOG = _AdminFileLogger(ADMIN_LOG_DIR, prefix="aistate_user")
ADMIN_FILE_LOG = _AdminFileLogger(ADMIN_LOG_DIR, prefix="aistate", pair=AUDIT_FILE_LOG)
AUDIT_FILE_LOG._pair = ADMIN_FILE_LOG
AUDIT_STORE._file_logger = AUDIT_FILE_LOG

# Ollama client (local LLM). Used by Analysis endpoints.
OLLAMA = OllamaClient()

# Ollama model install tasks (background pulls).
# Stored so the UI can query /api/ollama/install/status and the Logs tab can show progress.
# Shape:
#   {"status": "idle|running|done|error", "progress": int(0..100), "stage": str, "scope": "quick|deep|unknown", "error": str|None, "started_at": iso, "updated_at": iso}
OLLAMA_INSTALL_TASKS: Dict[str, Dict[str, Any]] = {}
_OLLAMA_INSTALL_LOCK = threading.Lock()

def _set_install_state(model: str, **updates: Any) -> None:
    with _OLLAMA_INSTALL_LOCK:
        st = OLLAMA_INSTALL_TASKS.get(model) or {
            "status": "idle",
            "progress": 0,
            "stage": "",
            "scope": "unknown",
            "error": None,
            "started_at": None,
            "updated_at": None,
        }
        st.update(updates)
        st["updated_at"] = now_iso()
        OLLAMA_INSTALL_TASKS[model] = st

def _ollama_pull_stream(model: str, base_url: str) -> None:
    """Pull model using Ollama HTTP API: POST {base}/api/pull with stream=true.

    Streams JSON lines with fields like: status, completed, total.
    """
    import urllib.request

    url = base_url.rstrip("/") + "/api/pull"
    payload = json.dumps({"name": model, "stream": True}).encode("utf-8")
    req = urllib.request.Request(url, data=payload, method="POST", headers={"Content-Type": "application/json"})

    last_logged_bucket = -1
    last_stage = None

    with urllib.request.urlopen(req, timeout=60) as resp:
        for raw in resp:
            line = raw.decode("utf-8", errors="ignore").strip()
            if not line:
                continue
            try:
                evt = json.loads(line)
            except Exception:
                evt = {"status": line}

            stage = str(evt.get("status") or evt.get("message") or "").strip()
            completed = evt.get("completed")
            total = evt.get("total")

            progress = None
            if isinstance(completed, (int, float)) and isinstance(total, (int, float)) and total:
                try:
                    progress = int((float(completed) / float(total)) * 100)
                    progress = max(0, min(100, progress))
                except Exception:
                    progress = None

            if stage and stage != last_stage:
                last_stage = stage
                _set_install_state(model, stage=stage)
                app_log(f"Ollama pull stage: model={model} | {stage}")

            if progress is not None:
                _set_install_state(model, progress=progress)
                # log every 5% to avoid log spam
                bucket = progress // 5
                if bucket != last_logged_bucket:
                    last_logged_bucket = bucket
                    app_log(f"Ollama pull progress: model={model} | {progress}%")

            if evt.get("error"):
                raise RuntimeError(str(evt["error"])[:4000])

def _ollama_install_worker(model: str, scope: str = "unknown") -> None:
    """Background worker to install/pull a model.

    Important: it continues even if user leaves the Settings tab (server-side thread).
    Also: it logs progress to the system Logs tab.
    """
    base_url = os.getenv("OLLAMA_HOST") or "http://127.0.0.1:11434"
    _set_install_state(model, status="running", progress=0, stage="starting", scope=scope, error=None, started_at=now_iso())
    app_log(f"Ollama install requested: model={model}, scope={scope}")

    try:
        # Preferred: streaming HTTP pull (gives progress).
        try:
            _ollama_pull_stream(model, base_url=base_url)
        except Exception as e:
            # Fallback to client ensure_model (best-effort) if streaming pull fails.
            app_log(f"Ollama pull stream failed: model={model} | {e} | fallback to ensure_model()")
            asyncio.run(OLLAMA.ensure_model(model))

        _set_install_state(model, status="done", progress=100, stage="done", error=None)
        app_log(f"Ollama install done: model={model}, scope={scope}")
    except Exception as e:
        _set_install_state(model, status="error", error=str(e)[:4000], stage="error")
        app_log(f"Ollama install error: model={model}, scope={scope} | {e}")


# ---------- Helpers: default name + secure delete (best-effort) ----------
from datetime import datetime, timedelta

def default_project_name() -> str:
    # Default name: AISTATE_YYYY-MM-DD (editable in UI)
    return f"AISTATE_{datetime.now().date().isoformat()}"

def _overwrite_path_bytes(path: Path, pattern: bytes, chunk_size: int = 1024 * 1024) -> None:
    size = path.stat().st_size
    if size <= 0:
        return
    with open(path, "r+b", buffering=0) as f:
        remaining = size
        if not pattern:
            raise ValueError("Empty overwrite pattern")
        buf = (pattern * (chunk_size // len(pattern) + 1))[:chunk_size]
        while remaining > 0:
            n = min(chunk_size, remaining)
            f.write(buf[:n])
            remaining -= n
        f.flush()
        os.fsync(f.fileno())

def _overwrite_path_random(path: Path, passes: int = 1, chunk_size: int = 1024 * 1024) -> None:
    size = path.stat().st_size
    if size <= 0:
        return
    for _ in range(passes):
        with open(path, "r+b", buffering=0) as f:
            remaining = size
            while remaining > 0:
                n = min(chunk_size, remaining)
                f.write(os.urandom(n))
                remaining -= n
            f.flush()
            os.fsync(f.fileno())

def _gutmann_patterns() -> List[bytes]:
    # Gutmann 35-pass sequence (classic). Patterns for passes 5-31 are based on Gutmann's table.
    # Passes 1-4 and 32-35 are random.
    # Source: Peter Gutmann, USENIX Security '96. citeturn0search2
    P: List[bytes] = []
    P.append(b"\x55")
    P.append(b"\xAA")
    P.append(bytes([0x92, 0x49, 0x24]))
    P.append(bytes([0x49, 0x24, 0x92]))
    P.append(bytes([0x24, 0x92, 0x49]))
    P.extend([bytes([x]) for x in [
        0x00, 0x11, 0x22, 0x33, 0x44, 0x55, 0x66, 0x77, 0x88, 0x99,
        0xAA, 0xBB, 0xCC, 0xDD, 0xEE, 0xFF
    ]])
    P.append(bytes([0x92, 0x49, 0x24]))
    P.append(bytes([0x49, 0x24, 0x92]))
    P.append(bytes([0x24, 0x92, 0x49]))
    P.append(bytes([0x6D, 0xB6, 0xDB]))
    P.append(bytes([0xB6, 0xDB, 0x6D]))
    P.append(bytes([0xDB, 0x6D, 0xB6]))
    return P  # 27 pattern passes

def secure_delete_project_dir(pdir: Path, method: str) -> None:
    # Best-effort file wiping inside project directory, then removal.
    # NOTE: On SSD/VM/CoW filesystems, overwriting may not guarantee secure erase.
    # For HMG IS5 Enhanced (0x00, 0xFF, Random) see e.g. LSoft KillDisk manual. citeturn0search31
    method = (method or "none").lower().strip()

    if method == "none":
        shutil.rmtree(pdir)
        return

    files = [fp for fp in pdir.rglob("*") if fp.is_file()]

    if method == "random1":
        for fp in files:
            try:
                _overwrite_path_random(fp, passes=1)
            except Exception:
                pass
            try:
                fp.unlink()
            except Exception:
                pass

    elif method == "hmg_is5":
        for fp in files:
            try:
                _overwrite_path_bytes(fp, b"\x00")
                _overwrite_path_bytes(fp, b"\xFF")
                _overwrite_path_random(fp, passes=1)
            except Exception:
                pass
            try:
                fp.unlink()
            except Exception:
                pass

    elif method == "gutmann":
        patterns = _gutmann_patterns()
        for fp in files:
            try:
                _overwrite_path_random(fp, passes=4)
                for pat in patterns:
                    _overwrite_path_bytes(fp, pat)
                _overwrite_path_random(fp, passes=4)
            except Exception:
                pass
            try:
                fp.unlink()
            except Exception:
                pass

    else:
        shutil.rmtree(pdir)
        return

    shutil.rmtree(pdir, ignore_errors=True)


WHISPER_MODELS = [
    # NOTE: OpenAI Whisper also exposes a "turbo" model (optimized large-v3 variant)
    # which is very fast for transcription tasks.
    "tiny", "base", "small", "medium", "large", "turbo", "large-v2", "large-v3",
]

# Optional ASR engines: Whisper / NVIDIA NeMo / pyannote
# NOTE: We do not persist defaults yet (per user request) — selection is UI-only.
NEMO_MODELS = [
    # English (fast / lightweight)
    "nvidia/stt_en_conformer_ctc_small",
    # English (highest accuracy)
    "nvidia/stt_en_conformer_transducer_large",
    # Multilingual (best overall; strong Polish)
    "nvidia/stt_multilingual_fastconformer_hybrid_large_pc",
]

# NeMo diarization models (speaker diarization / embeddings)
# NOTE: These are UI-only presets for caching in ASR Settings.
NEMO_DIARIZATION_MODELS = [
    # MSDD diarization (telephonic) – best quality for overlap + multi-language
    "diar_msdd_telephonic",
]

PYANNOTE_PIPELINES = [
    "pyannote/speaker-diarization-community-1",
    "pyannote/speaker-diarization-3.1",
    "pyannote/speaker-diarization",
]


# --- NLLB translation models (offline via HuggingFace cache) ---
# These are UI presets for the Translation "Fast" / "Accurate" modes.
NLLB_FAST_MODELS = [
    # Best speed/size ratio
    "facebook/nllb-200-distilled-600M",
    # Better quality while still reasonably fast
    "facebook/nllb-200-distilled-1.3B",
]

NLLB_ACCURATE_MODELS = [
    # Higher quality (full)
    "facebook/nllb-200-1.3B",
    # Highest quality (very large)
    "facebook/nllb-200-3.3B",
]


def now_iso() -> str:
    import datetime as _dt
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def safe_filename(name: str) -> str:
    # best-effort sanitization for uploads
    name = name.replace("\\", "_").replace("/", "_")
    return "".join(c for c in name if c.isalnum() or c in "._- ").strip()[:180] or "audio"


@dataclass
class TaskState:
    task_id: str
    kind: str
    project_id: str
    status: str = "queued"  # queued|running|done|error
    progress: int = 0
    logs: List[str] = field(default_factory=list)
    result: Optional[Dict[str, Any]] = None
    error: Optional[str] = None
    started_at: str = field(default_factory=now_iso)
    finished_at: Optional[str] = None

    def add_log(self, line: str) -> None:
        line = line.rstrip("\n")
        if not line:
            return
        self.logs.append(line)
        # Persist logs for administrators (separate from project logs)
        try:
            for _ln in str(line).splitlines() or []:
                _ln = (_ln or "").rstrip("\n")
                if not _ln:
                    continue
                if self.task_id == "system":
                    ADMIN_FILE_LOG.write_line(_ln)
                else:
                    ADMIN_FILE_LOG.write_line(f"{now_iso()} | task={self.task_id} kind={self.kind} project={self.project_id} | {_ln}")
        except Exception:
            pass

        # keep last N lines to avoid memory blow-up
        if len(self.logs) > 2000:
            self.logs = self.logs[-2000:]


class TaskManager:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._tasks: Dict[str, TaskState] = {}

        # Track running subprocess handles for optional cancellation from Admin panel.
        self._procs: Dict[str, Any] = {}
        self._procs_lock = threading.Lock()

        # A persistent "system" task to collect server-side application events.
        # This lets the Logs tab show more than just worker stderr.
        sys_task = TaskState(
            task_id="system",
            kind="system",
            project_id="-",
            status="running",
            progress=0,
        )
        self._tasks[sys_task.task_id] = sys_task

    def system_log(self, msg: str) -> None:
        """Append an application-level log line (English only)."""
        try:
            t = self.get("system")
            t.add_log(f"{now_iso()} | {msg}")
        except Exception:
            # Best-effort: never break the request because of logging.
            pass

    def create_task(self, kind: str, project_id: str) -> TaskState:
        """Create an in-memory task (used for async/streaming operations)."""
        task_id = uuid.uuid4().hex
        t = TaskState(task_id=task_id, kind=kind, project_id=project_id, status="running", progress=0)
        self._set(t)
        t.add_log(f"Task started: kind={kind}, project_id={project_id}, task_id={task_id}")
        try:
            self.system_log(f"Task started: {kind} (project_id={project_id}, task_id={task_id})")
        except Exception:
            pass
        return t

    def create_queued_task(self, kind: str, project_id: str, task_id: Optional[str] = None) -> TaskState:
        """Create a queued task placeholder (used by the GPU Resource Manager)."""
        tid = task_id or uuid.uuid4().hex
        t = TaskState(task_id=tid, kind=kind, project_id=project_id, status="queued", progress=0)
        self._set(t)
        t.add_log(f"Task queued: kind={kind}, project_id={project_id}, task_id={tid}")
        try:
            self.system_log(f"Task queued: {kind} (project_id={project_id}, task_id={tid})")
        except Exception:
            pass
        return t

    def cancel_task(self, task_id: str, reason: str = "Canceled by admin") -> bool:
        """Try to cancel a queued/running task. Best-effort."""
        try:
            t = self.get(task_id)
        except KeyError:
            return False

        # If there's an active subprocess, terminate it.
        with getattr(self, "_procs_lock", threading.Lock()):
            p = getattr(self, "_procs", {}).get(task_id)

        if p is not None:
            try:
                p.terminate()
                t.add_log("Task cancellation requested (terminate).")
                try:
                    self.system_log(f"Task cancel requested: {t.kind} (task_id={task_id})")
                except Exception:
                    pass
                # SIGKILL fallback if SIGTERM doesn't work within 5s
                def _force_kill() -> None:
                    try:
                        p.wait(timeout=5)
                    except Exception:
                        try:
                            p.kill()
                            t.add_log("Process did not exit after SIGTERM, sent SIGKILL.")
                        except Exception:
                            pass
                threading.Thread(target=_force_kill, daemon=True).start()
                return True
            except Exception as e:
                t.add_log(f"Task cancellation failed: {e}")
                return False

        # Otherwise: mark as error if still queued/running
        if t.status in ("queued", "running"):
            t.status = "error"
            t.error = reason
            t.finished_at = now_iso()
            t.add_log(reason)
            try:
                self.system_log(f"Task canceled: {t.kind} (task_id={task_id})")
            except Exception:
                pass
            return True
        return False

    def start_subprocess_with_id(
        self,
        task_id: str,
        kind: str,
        project_id: str,
        cmd: List[str],
        cwd: Path,
        env: Optional[Dict[str, str]] = None,
    ) -> TaskState:
        """Start a subprocess but reuse a pre-created task_id (queued placeholder)."""
        import subprocess

        try:
            t = self.get(task_id)
            # Keep existing logs (queue history), but reset state for run.
            t.kind = kind
            t.project_id = project_id
            t.status = "running"
            t.progress = 0
            t.error = None
            t.result = None
            t.finished_at = None
        except KeyError:
            t = TaskState(task_id=task_id, kind=kind, project_id=project_id, status="running", progress=0)
            self._set(t)

        t.add_log(f"Task started: kind={kind}, project_id={project_id}, task_id={task_id}")
        try:
            self.system_log(f"Task started: {kind} (project_id={project_id}, task_id={task_id})")
        except Exception:
            pass

        proc_env = os.environ.copy()
        if env:
            proc_env.update(env)

        def _redact_cmd(argv: List[str]) -> str:
            out: List[str] = []
            skip_next = False
            for a in argv:
                if skip_next:
                    out.append("******")
                    skip_next = False
                    continue
                if a in ("--hf_token", "--token", "--hf-token"):
                    out.append(a)
                    skip_next = True
                    continue
                if "hf_" in a and len(a) > 20:
                    out.append("******")
                else:
                    out.append(a)
            return " ".join(out)

        t.add_log(f"Command: {_redact_cmd(cmd)}")

        p = subprocess.Popen(
            cmd,
            cwd=str(cwd),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            bufsize=1,
            env=proc_env,
        )

        # Track proc for optional cancellation
        try:
            with self._procs_lock:
                self._procs[task_id] = p
        except Exception:
            pass

        stdout_buf: List[str] = []

        def read_stderr() -> None:
            assert p.stderr is not None
            for line in p.stderr:
                line = line.rstrip("\n")
                # Progress markers from workers: "PROGRESS: <0-100>"
                if line.startswith("PROGRESS"):
                    try:
                        pct = int(line.split(":")[-1].strip())
                        t.progress = max(0, min(100, pct))
                    except Exception:
                        pass
                    continue

                if line.strip():
                    t.add_log(line)
                    try:
                        if t.task_id != "system":
                            self.system_log(f"[{kind}:{task_id[:8]}] {line}")
                    except Exception:
                        pass
            t.add_log("(stderr stream closed)")

        def read_stdout() -> None:
            assert p.stdout is not None
            for line in p.stdout:
                s = line.rstrip("\n")
                if s.strip().startswith("{") or s.strip().startswith("["):
                    stdout_buf.append(line)
                else:
                    if s.strip():
                        t.add_log(f"STDOUT: {s}")
                        try:
                            if t.task_id != "system" and s.strip():
                                self.system_log(f"[{kind}:{task_id[:8]}] STDOUT: {s}")
                        except Exception:
                            pass
            t.add_log("(stdout stream closed)")

        th1 = threading.Thread(target=read_stderr, daemon=True)
        th2 = threading.Thread(target=read_stdout, daemon=True)
        th1.start()
        th2.start()


        def finalize() -> None:
            try:
                rc = p.wait()
            except Exception as e:
                t.status = "error"
                t.error = f"Failed to wait for subprocess: {e}"
                t.finished_at = now_iso()
                return

            # Try to release proc handle
            try:
                with self._procs_lock:
                    self._procs.pop(task_id, None)
            except Exception:
                pass

            th1.join(timeout=1)
            th2.join(timeout=1)
            t.finished_at = now_iso()

            out = "".join(stdout_buf).strip()
            if rc != 0:
                t.status = "error"
                t.error = f"Worker exited with code {rc}"
                if out:
                    t.add_log("STDOUT (tail): " + out[-2000:])
                try:
                    self.system_log(f"Task finished: {kind} (task_id={task_id}) -> ERROR (exit code {rc})")
                except Exception:
                    pass
                return

            # Parse JSON result
            try:
                data = json.loads(out) if out else {}
                t.status = "done"
                t.result = data if isinstance(data, dict) else {"data": data}

                # If this was an ASR install/predownload, rescan caches so UI updates without manual refresh.
                try:
                    if isinstance(kind, str) and (kind.startswith("asr_predownload_") or kind.startswith("asr_install_")):
                        scan_and_persist_asr_model_cache()
                except Exception as e:
                    try:
                        t.add_log(f"ASR rescan failed: {e}")
                    except Exception:
                        pass

                # If this was an NLLB install/predownload, rescan caches so UI updates.
                try:
                    if isinstance(kind, str) and (kind.startswith("nllb_predownload_") or kind.startswith("nllb_install_")):
                        scan_and_persist_nllb_model_cache()
                except Exception as e:
                    try:
                        t.add_log(f"NLLB rescan failed: {e}")
                    except Exception:
                        pass

                # If transcription, persist Whisper's detected language to project meta.
                try:
                    if isinstance(kind, str) and kind.startswith("transcribe") and isinstance(data, dict):
                        det = data.get("detected_lang")
                        if det and project_id and project_id != "-":
                            meta = read_project_meta(project_id)
                            meta["detected_lang"] = str(det)
                            write_project_meta(project_id, meta)
                except Exception:
                    pass

                try:
                    self.system_log(f"Task finished: {kind} (task_id={task_id}) -> DONE")
                except Exception:
                    pass
            except Exception as e:
                # Try to recover JSON from stdout (in case of garbage output)
                import re as _re
                
                recovered = None
                try:
                    # Try multiple recovery strategies
                    # 1. Find last JSON object (allow garbage before it)
                    matches = list(_re.finditer(r"(\{[\s\S]*?\})\s*$", out))
                    if matches:
                        recovered = json.loads(matches[-1].group(1))
                    
                    # 2. If that fails, try line-by-line from end
                    if recovered is None:
                        lines = out.strip().splitlines()
                        for line in reversed(lines):
                            line = line.strip()
                            if line.startswith('{') and line.endswith('}'):
                                try:
                                    recovered = json.loads(line)
                                    break
                                except Exception:
                                    continue
                except Exception:
                    recovered = None
                
                if isinstance(recovered, dict):
                    t.status = "done"
                    t.result = recovered
                    try:
                        self.system_log(f"Task finished: {kind} (task_id={task_id}) -> DONE (recovered JSON)")
                    except Exception:
                        pass
                else:
                    t.status = "error"
                    t.error = f"Invalid JSON from worker: {e}"
                    t.add_log("STDOUT (tail): " + out[-2000:])
                    try:
                        self.system_log(f"Task finished: {kind} (task_id={task_id}) -> ERROR (invalid JSON)")
                    except Exception:
                        pass

        threading.Thread(target=finalize, daemon=True).start()
        return t

    def list_tasks(self) -> List[TaskState]:
        with self._lock:
            tasks = list(self._tasks.values())
            system = [t for t in tasks if t.task_id == "system"]
            rest = [t for t in tasks if t.task_id != "system"][::-1]
            return system + rest

    def get(self, task_id: str) -> TaskState:
        with self._lock:
            if task_id not in self._tasks:
                raise KeyError(task_id)
            return self._tasks[task_id]

    def clear(self) -> None:
        with self._lock:
            # Clear all tasks and reset the persistent system log task.
            self._tasks.clear()
            self._tasks["system"] = TaskState(
                task_id="system",
                kind="system",
                project_id="-",
                status="running",
                progress=0,
            )

    def _set(self, t: TaskState) -> None:
        with self._lock:
            self._tasks[t.task_id] = t

    def start_subprocess(
        self,
        kind: str,
        project_id: str,
        cmd: List[str],
        cwd: Path,
        env: Optional[Dict[str, str]] = None,
    ) -> TaskState:
        import subprocess

        task_id = uuid.uuid4().hex
        t = TaskState(task_id=task_id, kind=kind, project_id=project_id, status="running", progress=0)
        self._set(t)

        # Task header (always visible even if worker is quiet for a while)
        t.add_log(f"Task started: kind={kind}, project_id={project_id}, task_id={task_id}")
        try:
            self.system_log(f"Task started: {kind} (project_id={project_id}, task_id={task_id})")
        except Exception:
            pass

        proc_env = os.environ.copy()
        if env:
            proc_env.update(env)

        # Redact obvious secrets (HF token) from command line in logs.
        def _redact_cmd(argv: List[str]) -> str:
            out: List[str] = []
            skip_next = False
            for a in argv:
                if skip_next:
                    out.append("***")
                    skip_next = False
                    continue
                if a in ("--hf_token", "--token", "--api_key", "--apikey", "--key"):
                    out.append(a)
                    skip_next = True
                    continue
                out.append(a)
            return " ".join(out)

        t.add_log(f"Command: {_redact_cmd(cmd)}")

        p = subprocess.Popen(
            cmd,
            cwd=str(cwd),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            bufsize=1,
        )

        stdout_buf: List[str] = []

        def read_stderr() -> None:
            assert p.stderr is not None
            for line in p.stderr:
                line = line.rstrip("\n")
                # Progress markers from workers:
                # transcribe_worker: "PROGRESS: 12"
                # voice_worker: "PROGRESS:12"
                if line.startswith("PROGRESS"):
                    try:
                        pct = int(line.split(":")[-1].strip())
                        t.progress = max(0, min(100, pct))
                    except Exception:
                        pass
                else:
                    t.add_log(line)
                    try:
                        if t.task_id != "system" and line.strip():
                            self.system_log(f"[{kind}:{task_id[:8]}] {line}")
                    except Exception:
                        pass

            # Always flush a final marker for readability.
            t.add_log("(stderr stream closed)")

        def read_stdout() -> None:
            assert p.stdout is not None
            for line in p.stdout:
                s = line.rstrip("\n")
                # Worker stdout should end with JSON. Any other stdout lines
                # are treated as logs (many libs print to stdout).
                if s.strip().startswith("{") or s.strip().startswith("["):
                    stdout_buf.append(line)
                else:
                    if s.strip():
                        t.add_log(f"STDOUT: {s}")
                        try:
                            if t.task_id != "system" and s.strip():
                                self.system_log(f"[{kind}:{task_id[:8]}] STDOUT: {s}")
                        except Exception:
                            pass

            t.add_log("(stdout stream closed)")

        th1 = threading.Thread(target=read_stderr, daemon=True)
        th2 = threading.Thread(target=read_stdout, daemon=True)
        th1.start()
        th2.start()


        def finalize() -> None:
            rc = p.wait()
            th1.join(timeout=1)
            th2.join(timeout=1)
            t.finished_at = now_iso()

            out = "".join(stdout_buf).strip()
            if rc != 0:
                t.status = "error"
                t.error = f"Process exited with code {rc}"
                if out:
                    t.add_log("STDOUT (last): " + out[-2000:])
                try:
                    self.system_log(f"Task finished: {kind} (task_id={task_id}) -> ERROR (rc={rc})")
                except Exception:
                    pass
                return

            def _asr_rescan() -> None:
                try:
                    if isinstance(kind, str) and (kind.startswith("asr_predownload_") or kind.startswith("asr_install_")):
                        scan_and_persist_asr_model_cache()
                except Exception as e:
                    try:
                        t.add_log(f"ASR rescan failed: {e}")
                    except Exception:
                        pass

            def _nllb_rescan() -> None:
                try:
                    if isinstance(kind, str) and (kind.startswith("nllb_predownload_") or kind.startswith("nllb_install_")):
                        scan_and_persist_nllb_model_cache()
                except Exception as e:
                    try:
                        t.add_log(f"NLLB rescan failed: {e}")
                    except Exception:
                        pass

            # parse JSON result
            if not out:
                t.result = {}
                t.status = "done"
                t.progress = 100
                _asr_rescan()
                _nllb_rescan()
                try:
                    self.system_log(f"Task finished: {kind} (task_id={task_id}) -> DONE")
                except Exception:
                    pass
                return

            # Be robust: if stdout contains any accidental extra output,
            # try to recover the last JSON object from the stream.
            try:
                t.result = json.loads(out)
                t.status = "done"
                t.progress = 100
                try:
                    _persist_task_outputs(t)
                except Exception as e:
                    t.add_log(f"PERSIST ERROR: {e}")
                _asr_rescan()
                _nllb_rescan()
                try:
                    self.system_log(f"Task finished: {kind} (task_id={task_id}) -> DONE")
                except Exception:
                    pass
            except Exception as e:
                import re as _re

                recovered = None
                try:
                    # Try multiple recovery strategies
                    # 1. Find last JSON object (allow garbage before it)
                    matches = list(_re.finditer(r"(\{[\s\S]*?\})\s*$", out))
                    if matches:
                        recovered = json.loads(matches[-1].group(1))
                    
                    # 2. If that fails, try line-by-line from end
                    if recovered is None:
                        lines = out.strip().splitlines()
                        for line in reversed(lines):
                            line = line.strip()
                            if line.startswith('{') and line.endswith('}'):
                                try:
                                    recovered = json.loads(line)
                                    break
                                except Exception:
                                    continue
                except Exception:
                    recovered = None

                if isinstance(recovered, dict):
                    t.result = recovered
                    t.status = "done"
                    t.progress = 100
                    try:
                        _persist_task_outputs(t)
                    except Exception as e:
                        t.add_log(f"PERSIST ERROR: {e}")
                    _asr_rescan()
                    _nllb_rescan()
                    try:
                        self.system_log(f"Task finished: {kind} (task_id={task_id}) -> DONE (recovered JSON)")
                    except Exception:
                        pass
                else:
                    t.status = "error"
                    t.error = f"Invalid JSON from worker: {e}"
                    # Keep a tail for debugging (avoid huge payloads)
                    t.add_log("STDOUT (tail): " + out[-2000:])
                    try:
                        self.system_log(f"Task finished: {kind} (task_id={task_id}) -> ERROR (invalid JSON)")
                    except Exception:
                        pass

        threading.Thread(target=finalize, daemon=True).start()
        return t

    def start_python_fn(self, kind: str, project_id: str, fn, *args, task_id_override: 'Optional[str]' = None, **kwargs) -> TaskState:
        task_id = str(task_id_override or uuid.uuid4().hex)
        t = None
        if task_id_override:
            try:
                t = self.get(task_id)
            except Exception:
                t = None
        if t is None:
            t = TaskState(task_id=task_id, kind=kind, project_id=project_id, status="running", progress=0)
            self._set(t)
        else:
            # Reuse existing task object (best-effort)
            t.kind = kind
            t.project_id = project_id
            t.status = "running"
            t.progress = 0
            t.error = None
            t.result = None
            t.started_at = now_iso()
            t.finished_at = None

        # Task header
        t.add_log(f"Task started: kind={kind}, project_id={project_id}, task_id={task_id}")
        try:
            self.system_log(f"Task started: {kind} (project_id={project_id}, task_id={task_id})")
        except Exception:
            pass

        def run() -> None:
            try:
                def log_cb(msg: str) -> None:
                    t.add_log(str(msg))
                def progress_cb(pct: int) -> None:
                    t.progress = max(0, min(100, int(pct)))

                # Only pass callbacks if the function supports them.
                import inspect

                call_kwargs = dict(kwargs)
                try:
                    sig = inspect.signature(fn)
                    params = sig.parameters
                    accepts_varkw = any(p.kind == inspect.Parameter.VAR_KEYWORD for p in params.values())
                    if accepts_varkw or "log_cb" in params:
                        call_kwargs.setdefault("log_cb", log_cb)
                    if accepts_varkw or "progress_cb" in params:
                        call_kwargs.setdefault("progress_cb", progress_cb)
                except Exception:
                    # If introspection fails, try best-effort with callbacks.
                    call_kwargs.setdefault("log_cb", log_cb)
                    call_kwargs.setdefault("progress_cb", progress_cb)

                try:
                    res = fn(*args, **call_kwargs)
                except TypeError:
                    # Fallback: some callables reject unexpected kwargs.
                    res = fn(*args)
                t.result = res if isinstance(res, dict) else {"result": res}
                t.status = "done"
                t.progress = 100
                try:
                    self.system_log(f"Task finished: {kind} (task_id={task_id}) -> DONE")
                except Exception:
                    pass
            except Exception as e:
                import traceback
                t.status = "error"
                t.error = str(e)
                t.add_log(traceback.format_exc())
                try:
                    self.system_log(f"Task finished: {kind} (task_id={task_id}) -> ERROR: {e}")
                except Exception:
                    pass
            finally:
                t.finished_at = now_iso()

        threading.Thread(target=run, daemon=True).start()
        return t


TASKS = TaskManager()


def app_log(msg: str) -> None:
    """Server-side app log (English only). Visible in Logs tab as "system" task."""
    try:
        TASKS.system_log(msg)
    except Exception:
        pass



# ---------------------------
# GPU Resource Manager (Admin)
# ---------------------------

@dataclass
class _GPUJob:
    # Common fields
    task_id: str
    kind: str
    project_id: str
    created_at: str
    priority: int = 0  # higher = more important

    # Subprocess job fields
    cmd: Optional[List[str]] = None
    cwd: Optional[Path] = None
    env: Dict[str, str] = field(default_factory=dict)

    # Python function job fields
    fn: Any = None
    args: List[Any] = field(default_factory=list)
    kwargs: Dict[str, Any] = field(default_factory=dict)

    assigned_device: Optional[str] = None  # "gpu:0" / "cpu"
    status: str = "queued"  # queued|running|done|error


class GPUResourceManager:
    """
    Lightweight in-process scheduler that queues GPU-heavy jobs and starts them
    only when a slot is available. This prevents VRAM spikes and keeps the UI responsive.

    Design goals:
    - single GPU (desktop) works out of the box
    - scalable: multi-GPU servers (CUDA_VISIBLE_DEVICES per job)
    - multi-user: queue + concurrency limits (slots per GPU / CPU slots)
    """

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._jobs: List[_GPUJob] = []
        self._running: Dict[str, int] = {}  # device -> count
        self._gpus: List[Dict[str, Any]] = []
        self._cuda_available: bool = False

        # runtime config (persisted in global settings via _get/_save_gpu_rm_settings)
        self.gpu_mem_fraction: float = 0.85
        self.gpu_slots_per_gpu: int = 1
        self.cpu_slots: int = 1
        self.enabled: bool = True

        # Job priorities are configured per feature area (admin-facing).
        # Higher value = higher priority.
        self.category_priorities: Dict[str, int] = {
            "transcription": 300,
            "diarization": 200,
            "translation": 180,
            "analysis_quick": 140,
            "analysis": 120,
            "sound_detection": 100,
            "tts": 80,
            "chat": 60,
        }

        self._stop = False
        self._thread = threading.Thread(target=self._loop, daemon=True)
        self._thread.start()

    def apply_config(self, cfg: Dict[str, Any]) -> None:
        with self._lock:
            self.gpu_mem_fraction = float(cfg.get("gpu_mem_fraction", self.gpu_mem_fraction))
            self.gpu_slots_per_gpu = int(cfg.get("gpu_slots_per_gpu", self.gpu_slots_per_gpu))
            self.cpu_slots = int(cfg.get("cpu_slots", self.cpu_slots))
            self.enabled = bool(cfg.get("enabled", self.enabled))

            # Optional per-category priorities (persisted in global settings)
            pr = cfg.get("priorities")
            if isinstance(pr, dict):
                merged = dict(self.category_priorities)
                for k, v in pr.items():
                    ks = str(k)
                    if ks not in merged:
                        continue
                    try:
                        merged[ks] = int(v)
                    except Exception:
                        continue
                self.category_priorities = merged

    def _detect_gpus(self) -> None:
        # Best-effort; do not crash if torch is missing.
        gpus: List[Dict[str, Any]] = []
        cuda_ok = False
        try:
            import torch  # type: ignore

            cuda_ok = bool(torch.cuda.is_available())
            if cuda_ok:
                n = int(torch.cuda.device_count())
                for i in range(n):
                    props = torch.cuda.get_device_properties(i)
                    gpus.append({
                        "id": i,
                        "name": getattr(props, "name", f"cuda:{i}"),
                        "total_vram_bytes": int(getattr(props, "total_memory", 0)),
                    })
        except Exception:
            cuda_ok = False
            gpus = []

        with self._lock:
            self._cuda_available = cuda_ok
            self._gpus = gpus

    def _available_devices(self) -> List[str]:
        # Return device identifiers that still have free slots.
        # Note: we always include CPU if it has free slots, even when GPUs exist.
        # Actual dispatch policy is enforced in _loop (some job kinds may be GPU-only).
        with self._lock:
            if not self.enabled:
                return []

            out: List[str] = []

            # GPU slots (if available)
            if self._cuda_available and self._gpus:
                for g in self._gpus:
                    dev = f"gpu:{g['id']}"
                    running = int(self._running.get(dev, 0))
                    if running < max(1, self.gpu_slots_per_gpu):
                        out.append(dev)

            # CPU slot (trackable even when GPU exists)
            running_cpu = int(self._running.get("cpu", 0))
            if running_cpu < max(1, self.cpu_slots):
                out.append("cpu")

            return out

    def _prio(self, kind: str) -> int:
        """Resolve scheduling priority from job kind using category mapping."""
        k = str(kind or "")
        cat = ""
        if k.startswith("transcribe"):
            cat = "transcription"
        elif k.startswith("diarize"):
            cat = "diarization"
        elif k.startswith("translate_"):
            cat = "translation"
        elif k.startswith("analysis_quick"):
            cat = "analysis_quick"
        elif k.startswith("analysis"):
            cat = "analysis"
        elif k.startswith("sound_detection"):
            cat = "sound_detection"
        elif k.startswith("tts"):
            cat = "tts"
        elif k.startswith("chat"):
            cat = "chat"

        try:
            return int(self.category_priorities.get(cat, 50))
        except Exception:
            return 50

    def enqueue_subprocess(
        self,
        kind: str,
        project_id: str,
        cmd: List[str],
        cwd: Path,
        env: Optional[Dict[str, str]] = None,
        task_id_override: Optional[str] = None,
    ) -> TaskState:
        # Always create a queued placeholder so the UI can poll immediately.
        t = TASKS.create_queued_task(kind=kind, project_id=project_id, task_id=task_id_override)
        job = _GPUJob(
            task_id=t.task_id,
            kind=kind,
            project_id=project_id,
            created_at=now_iso(),
            priority=self._prio(kind),
            cmd=cmd,
            cwd=cwd,
            env=dict(env or {}),
        )
        with self._lock:
            self._jobs.append(job)
        t.add_log(f"Queued by GPU Resource Manager (prio={job.priority}).")
        return t

    def enqueue_python_fn(
        self,
        kind: str,
        project_id: str,
        fn,
        *args,
        task_id_override: Optional[str] = None,
        **kwargs,
    ) -> TaskState:
        # Queue an in-process python callable (executed through TaskManager).
        t = TASKS.create_queued_task(kind=kind, project_id=project_id, task_id=task_id_override)
        job = _GPUJob(
            task_id=t.task_id,
            kind=kind,
            project_id=project_id,
            created_at=now_iso(),
            priority=self._prio(kind),
            fn=fn,
            args=list(args),
            kwargs=dict(kwargs),
        )
        with self._lock:
            self._jobs.append(job)
        t.add_log(f"Queued by GPU Resource Manager (python job, prio={job.priority}).")
        return t

    def cancel(self, task_id: str) -> bool:
        # Cancel queued job if present, otherwise attempt to terminate running task.
        with self._lock:
            for i, j in enumerate(self._jobs):
                if j.task_id == task_id and j.status == "queued":
                    self._jobs.pop(i)
                    TASKS.cancel_task(task_id, reason="Canceled (removed from GPU queue)")
                    return True
        # running: best-effort terminate
        return TASKS.cancel_task(task_id, reason="Canceled by admin")

    def status_snapshot(self) -> Dict[str, Any]:
        with self._lock:
            return {
                "cuda_available": self._cuda_available,
                "gpus": list(self._gpus),
                "queue_size": len([j for j in self._jobs if j.status == "queued"]),
                "running": dict(self._running),
                "config": {
                    "gpu_mem_fraction": self.gpu_mem_fraction,
                    "gpu_slots_per_gpu": self.gpu_slots_per_gpu,
                    "cpu_slots": self.cpu_slots,
                    "enabled": self.enabled,
                    "priorities": dict(self.category_priorities),
                },
            }

    def jobs_snapshot(self) -> Dict[str, Any]:
        rows: List[Dict[str, Any]] = []
        with self._lock:
            for j in self._jobs:
                # reflect current task status (done/error) if it finished
                try:
                    st = TASKS.get(j.task_id).status
                except Exception:
                    st = j.status
                device = j.assigned_device or "-"
                job_type = "python" if j.fn is not None else "subprocess"
                rows.append({
                    "task_id": j.task_id,
                    "kind": j.kind,
                    "project_id": j.project_id,
                    "created_at": j.created_at,
                    "priority": j.priority,
                    "job_type": job_type,
                    "status": st,
                    "device": device,
                })
        return {"jobs": rows}

    def _mark_running(self, device: str, delta: int) -> None:
        with self._lock:
            self._running[device] = int(self._running.get(device, 0)) + int(delta)
            if self._running[device] <= 0:
                self._running.pop(device, None)

    def _start_job(self, job: _GPUJob, device: str) -> None:
        # Start either a queued subprocess job or a queued python callable.
        # We keep slot accounting independent from job type.
        job.status = "running"
        job.assigned_device = device
        self._mark_running(device, +1)

        try:
            TASKS.get(job.task_id).add_log(f"Assigned device: {device} (prio={job.priority})")
        except Exception:
            pass

        try:
            if job.fn is not None:
                # Python job (e.g., quick/deep LLM) - run via TaskManager thread
                # IMPORTANT: pass the first positional parameters positionally.
                # Using kind=... plus *args can re-fill the same slot and causes:
                #   TaskManager.start_python_fn() got multiple values for argument 'kind'
                TASKS.start_python_fn(
                    job.kind,
                    job.project_id,
                    job.fn,
                    *list(job.args or []),
                    task_id_override=job.task_id,
                    **(job.kwargs or {}),
                )
            else:
                # Subprocess job (e.g., Whisper / Pyannote)
                if not job.cmd or not job.cwd:
                    raise RuntimeError("Invalid subprocess job: missing cmd/cwd")

                env = dict(job.env or {})
                env["AISTATE_GPU_MEM_FRACTION"] = str(self.gpu_mem_fraction)
                env["AISTATE_GPU_DEVICE"] = device

                if device.startswith("gpu:"):
                    gpu_id = device.split(":")[1]
                    # isolate GPU
                    env["CUDA_VISIBLE_DEVICES"] = str(gpu_id)
                    env["AISTATE_GPU_ID"] = str(gpu_id)
                elif device == "cpu":
                    # Explicitly hide GPUs so workers don't accidentally use CUDA
                    env["CUDA_VISIBLE_DEVICES"] = ""

                TASKS.start_subprocess_with_id(
                    task_id=job.task_id,
                    kind=job.kind,
                    project_id=job.project_id,
                    cmd=job.cmd,
                    cwd=job.cwd,
                    env=env,
                )
        except Exception as e:
            # Mark task as failed and release slot immediately
            try:
                t = TASKS.get(job.task_id)
                t.status = "error"
                t.error = str(e)
                t.finished_at = now_iso()
                t.add_log(f"GPU RM start failed: {e}")
            except Exception:
                pass
            job.status = "error"
            self._mark_running(device, -1)
            return

        def watch_finish() -> None:
            # Wait until task finishes, then free slot.
            finished = False
            for _ in range(60 * 60 * 24):  # up to 24h
                try:
                    st = TASKS.get(job.task_id).status
                    if st in ("done", "error"):
                        finished = True
                        break
                except Exception:
                    break
                time.sleep(1.0)
            if not finished:
                try:
                    t = TASKS.get(job.task_id)
                    t.status = "error"
                    t.error = "Task stuck — forcibly freed after 24h timeout"
                    t.finished_at = now_iso()
                    t.add_log("GPU RM: 24h watchdog timeout — slot released")
                    app_log(f"GPU RM watchdog: task {job.task_id} timed out after 24h on {device}")
                except Exception:
                    pass
            self._mark_running(device, -1)

        threading.Thread(target=watch_finish, daemon=True).start()

    def _loop(self) -> None:

        # periodic GPU refresh + dispatcher
        self._detect_gpus()
        last_gpu_refresh = time.time()

        while not self._stop:
            try:
                # refresh GPU list every 10s
                if time.time() - last_gpu_refresh > 10:
                    self._detect_gpus()
                    last_gpu_refresh = time.time()

                devs = self._available_devices()
                if not devs:
                    time.sleep(0.5)
                    continue

                dispatched = 0
                for dev in devs:
                    # Find next queued job by priority (higher value wins). FIFO within same priority.
                    with self._lock:
                        nxt = None
                        for j in self._jobs:
                            if j.status != "queued":
                                continue

                            # When GPUs exist, reserve CPU slots for lightweight jobs
                            # (translation). GPU-heavy jobs (ASR, diarization) stay on GPU.
                            # When NO GPU exists, allow everything on CPU.
                            if dev == "cpu" and self._cuda_available and self._gpus:
                                k = str(j.kind)
                                if not k.startswith("translate_"):
                                    continue

                            if (nxt is None) or (j.priority > nxt.priority):
                                nxt = j
                    if not nxt:
                        break

                    self._start_job(nxt, dev)
                    dispatched += 1

                if dispatched == 0:
                    time.sleep(0.5)
            except Exception as e:
                try:
                    app_log(f"GPUResourceManager loop error: {e}")
                except Exception:
                    pass
                time.sleep(1.0)


def project_path(project_id: str) -> Path:
    p = PROJECTS_DIR / project_id
    p.mkdir(parents=True, exist_ok=True)
    return p


def project_meta_path(project_id: str) -> Path:
    return project_path(project_id) / "project.json"


def read_project_meta(project_id: str) -> Dict[str, Any]:
    path = project_meta_path(project_id)
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def write_project_meta(project_id: str, meta: Dict[str, Any]) -> None:
    path = project_meta_path(project_id)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)


def ensure_project(project_id: Optional[str]) -> str:
    if project_id:
        project_path(project_id)
        meta = read_project_meta(project_id)
        if not meta:
            write_project_meta(project_id, {"project_id": project_id, "created_at": now_iso(), "name": "projekt"})
        return project_id
    new_id = uuid.uuid4().hex
    project_path(new_id)
    write_project_meta(new_id, {"project_id": new_id, "created_at": now_iso(), "name": "projekt"})
    return new_id


# ---------- Translation draft persistence (per-project) ----------

def translation_draft_path(project_id: str) -> Path:
    """Return path to translation draft file for a project."""
    return project_path(project_id) / "translation_draft.json"


def read_translation_draft(project_id: str) -> Dict[str, Any]:
    path = translation_draft_path(project_id)
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}
    return {}


def write_translation_draft(project_id: str, draft: Dict[str, Any]) -> None:
    path = translation_draft_path(project_id)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(draft, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)


def save_upload(project_id: str, upload: UploadFile) -> Path:
    pdir = project_path(project_id)
    fname = safe_filename(upload.filename or "audio")
    dst = pdir / fname
    with dst.open("wb") as f:
        shutil.copyfileobj(upload.file, f)
    meta = read_project_meta(project_id)
    meta["audio_file"] = fname
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    # Pre-generate waveform peaks in background so they're ready when user opens transcription
    def _gen_peaks():
        try:
            _generate_waveform_peaks(project_id)
        except Exception:
            pass
    threading.Thread(target=_gen_peaks, daemon=True).start()
    return dst


def require_existing_file(path: Path, msg: str) -> None:
    if not path.exists() or not path.is_file():
        raise HTTPException(status_code=400, detail=msg)


app = FastAPI(title=f"{APP_NAME} Web", version=APP_VERSION)
app.mount("/static", StaticFiles(directory=str(Path(__file__).resolve().parent / "static")), name="static")

# --- Mount auth/setup/users routers ---
auth_router.init(
    user_store=USER_STORE,
    session_store=SESSION_STORE,
    deployment_store=DEPLOYMENT_STORE,
    app_log_fn=app_log,
    get_session_timeout=_get_session_timeout,
    message_store=MESSAGE_STORE,
    audit_store=AUDIT_STORE,
    get_settings=load_settings,
)
app.include_router(auth_router.router)

users_router.init(
    user_store=USER_STORE,
    session_store=SESSION_STORE,
    app_log_fn=app_log,
    audit_store=AUDIT_STORE,
    get_settings=load_settings,
)
app.include_router(users_router.router)

messages_router.init(
    message_store=MESSAGE_STORE,
    app_log_fn=app_log,
)
app.include_router(messages_router.router)

setup_router.init(
    user_store=USER_STORE,
    deployment_store=DEPLOYMENT_STORE,
    app_log_fn=app_log,
    projects_dir=PROJECTS_DIR,
)
app.include_router(setup_router.router)

# --- Mount routers (refactored modules) ---
# Chat: inject ollama + log now; GPU_RM + TASKS injected later (after GPU_RM creation)
chat_router.init(ollama_client=OLLAMA, app_log_fn=app_log, tasks=TASKS)
app.include_router(chat_router.router)

tasks_router.init(tasks_manager=TASKS)
app.include_router(tasks_router.router)

# AML/DB router (SQL-backed project management + AML analysis)
app.include_router(aml_router.router)

# Workspaces router (project workspaces, subprojects, collaboration)
workspaces_router.init(workspace_store=WORKSPACE_STORE, user_store=USER_STORE)
app.include_router(workspaces_router.router)

# Initialize SQLite database on startup
try:
    from backend.db.engine import init_db
    init_db()
    # Migrate legacy JSON auth data → SQLite (runs once, renames old files to .bak)
    try:
        DEPLOYMENT_STORE.migrate_from_json()
        USER_STORE.migrate_from_json()
        SESSION_STORE.migrate_from_json()
        AUDIT_STORE.migrate_from_json()
        MESSAGE_STORE.migrate_from_json()
    except Exception as _mig_err:
        import logging as _mig_lg
        _mig_lg.getLogger("aistate").warning("Auth JSON→SQLite migration: %s", _mig_err)
    # Migrate file-based projects → workspaces (one-time, idempotent)
    try:
        # In multiuser mode, prefer a real superadmin over creating a phantom 'admin' user
        _default_uid = None
        if DEPLOYMENT_STORE.is_multiuser():
            for _u in USER_STORE.list_users():
                if getattr(_u, "is_superadmin", False):
                    _default_uid = _u.user_id
                    break
        if not _default_uid and DEPLOYMENT_STORE.is_configured():
            from backend.db.engine import get_default_user_id
            _default_uid = get_default_user_id()
        if _default_uid:
            WORKSPACE_STORE.migrate_file_projects(PROJECTS_DIR, _default_uid)
    except Exception as _ws_err:
        import logging as _ws_lg
        _ws_lg.getLogger("aistate").warning("Workspace migration: %s", _ws_err)
except Exception as _db_err:
    import logging as _lg
    _lg.getLogger("aistate").warning("DB init deferred: %s", _db_err)

# Admin router is initialised later (after GPU_RM and helper functions are defined).
# See: _mount_admin_router() below.


# --- Security headers middleware ---
@app.middleware("http")
async def _security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["X-XSS-Protection"] = "1; mode=block"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    return response


# --- Ensure UTF-8 across the app (templates + JSON + JS/CSS) ---
@app.middleware("http")
async def _force_utf8_charset(request: Request, call_next):
    response = await call_next(request)
    try:
        ct = response.headers.get("content-type") or ""
        ctl = ct.lower()
        # Only add charset when missing and it's a text-like response
        if ct and ("charset=" not in ctl) and (
            ctl.startswith("text/")
            or ctl.startswith("application/json")
            or ctl.startswith("application/javascript")
            or ctl.startswith("application/x-javascript")
        ):
            response.headers["content-type"] = f"{ct}; charset=utf-8"
    except Exception:
        pass
    return response


# --- Multi-user auth + authorization middleware ---
@app.middleware("http")
async def _auth_middleware(request: Request, call_next):
    """Session check + route authorization for multi-user mode."""
    path = request.url.path
    request.state.user = None
    request.state.multiuser = DEPLOYMENT_STORE.is_multiuser()

    # In single-user mode, skip all auth
    if not request.state.multiuser:
        # If not configured at all, redirect to /setup
        if not DEPLOYMENT_STORE.is_configured() and path not in ("/setup", "/static/app.css", "/static/logo.png") and not path.startswith("/static/") and not path.startswith("/api/setup/"):
            if path.startswith("/api/"):
                return JSONResponse({"status": "error", "message": "Setup required"}, status_code=503)
            return RedirectResponse(url="/setup", status_code=302)
        return await call_next(request)

    # Multi-user mode — check public routes first
    if path in PUBLIC_ROUTES:
        return await call_next(request)
    for prefix in PUBLIC_PREFIXES:
        if path.startswith(prefix):
            return await call_next(request)

    # Setup page is always public
    if path == "/setup":
        return await call_next(request)

    # If no users exist yet (setup incomplete), redirect to /setup
    if not USER_STORE.has_users():
        if path.startswith("/api/"):
            return JSONResponse({"status": "error", "message": "Setup required"}, status_code=503)
        return RedirectResponse(url="/setup", status_code=302)

    # Check session cookie
    token = request.cookies.get(SessionStore.COOKIE_NAME)
    if not token:
        if path.startswith("/api/"):
            return JSONResponse({"status": "error", "message": "Not authenticated"}, status_code=401)
        return RedirectResponse(url="/login", status_code=302)

    session = SESSION_STORE.get_session(token)
    if session is None:
        response = RedirectResponse(url="/login", status_code=302) if not path.startswith("/api/") else JSONResponse({"status": "error", "message": "Session expired"}, status_code=401)
        response.delete_cookie(key=SessionStore.COOKIE_NAME, path="/")
        return response

    # Load user
    user = USER_STORE.get_user(session["user_id"])
    if user is None:
        SESSION_STORE.delete_session(token)
        response = RedirectResponse(url="/login", status_code=302) if not path.startswith("/api/") else JSONResponse({"status": "error", "message": "User not found"}, status_code=401)
        response.delete_cookie(key=SessionStore.COOKIE_NAME, path="/")
        return response

    # Check ban
    if user.banned:
        if user.banned_until:
            from datetime import datetime as _dt
            try:
                if _dt.now() > _dt.fromisoformat(user.banned_until):
                    USER_STORE.update_user(user.user_id, {"banned": False, "banned_until": None, "ban_reason": None, "show_ban_expiry": True})
                else:
                    SESSION_STORE.delete_session(token)
                    if path.startswith("/api/"):
                        return JSONResponse({"status": "error", "message": "Account banned"}, status_code=403)
                    from urllib.parse import urlencode as _ue
                    _bp: dict = {}
                    if user.ban_reason:
                        _bp["reason"] = user.ban_reason
                    if getattr(user, "show_ban_expiry", True) and user.banned_until:
                        _bp["until"] = user.banned_until
                    return RedirectResponse(url="/banned" + ("?" + _ue(_bp) if _bp else ""), status_code=302)
            except ValueError:
                pass
        else:
            SESSION_STORE.delete_session(token)
            if path.startswith("/api/"):
                return JSONResponse({"status": "error", "message": "Account banned"}, status_code=403)
            from urllib.parse import urlencode as _ue2
            _bp2: dict = {}
            if user.ban_reason:
                _bp2["reason"] = user.ban_reason
            return RedirectResponse(url="/banned" + ("?" + _ue2(_bp2) if _bp2 else ""), status_code=302)

    request.state.user = user

    # ---- Must-change-password enforcement ----
    # If the user has never changed their password (first login) or the
    # password has expired, only allow the change-password flow and
    # essential routes.  This closes a security gap where refreshing the
    # page after login could bypass the client-side overlay.
    _pw_changed = getattr(user, "password_changed_at", None)
    _force_pw_change = not _pw_changed  # first login: password_changed_at is NULL
    if _pw_changed and not _force_pw_change:
        try:
            from datetime import datetime as _dt2, timedelta as _td2
            _s = load_settings()
            _expiry_days = getattr(_s, "password_expiry_days", 0)
            if _expiry_days and _expiry_days > 0:
                _pw_dt = _dt2.fromisoformat(_pw_changed)
                if _dt2.now() > _pw_dt + _td2(days=_expiry_days):
                    _force_pw_change = True
        except (ValueError, TypeError, AttributeError):
            pass

    if _force_pw_change:
        _pw_allowed_exact = {
            "/api/auth/change-password",
            "/api/auth/me",
            "/api/auth/logout",
            "/api/auth/password-policy",
            "/change-password",
        }
        if path not in _pw_allowed_exact and not path.startswith("/static/"):
            if path.startswith("/api/"):
                return JSONResponse(
                    {"status": "error", "message": "Password change required", "code": "must_change_password"},
                    status_code=403,
                )
            return RedirectResponse(url="/change-password", status_code=302)

    # Authorization: check route access
    user_modules = get_user_modules(user.role, user.is_admin, user.admin_roles, user.is_superadmin)

    # Common routes: /, /info, /api/auth/me, etc — allowed for any logged-in user
    # Module-specific routes: checked against user's modules
    if not is_route_allowed(path, user_modules):
        if path.startswith("/api/"):
            return JSONResponse({"status": "error", "message": "Access denied"}, status_code=403)
        return RedirectResponse(url="/", status_code=302)

    return await call_next(request)


# --- Startup: autoscan ASR caches so the UI can label models as installed/uninstalled ---
@app.on_event("startup")
async def _startup_asr_autoscan() -> None:
    def _run() -> None:
        try:
            scan_and_persist_asr_model_cache()
        except Exception as e:
            try:
                app_log(f"ASR autoscan failed: {e}")
            except Exception:
                pass

    threading.Thread(target=_run, daemon=True).start()


@app.on_event("startup")
async def _startup_nllb_autoscan() -> None:
    """Autoscan NLLB model cache for the NLLB Settings UI."""
    def _run() -> None:
        try:
            scan_and_persist_nllb_model_cache()
        except Exception as e:
            try:
                app_log(f"NLLB autoscan failed: {e}")
            except Exception:
                pass

    threading.Thread(target=_run, daemon=True).start()

@app.get("/", response_class=HTMLResponse)
def home(request: Request) -> Any:
    # Multi-user mode: skip Intro animation, go directly to projects
    if getattr(request.state, "multiuser", False):
        return RedirectResponse(url="/projects", status_code=302)

    # Single-user mode: play Intro (once per browser session) then go to the app.
    # We keep it client-side via sessionStorage so it doesn't require cookies.
    html = """<!doctype html>
<html lang=\"pl\"><head>
  <meta charset=\"utf-8\"/>
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1\"/>
  <title>AI S.T.A.T.E Web</title>
  <meta http-equiv=\"cache-control\" content=\"no-cache\"/>
  <meta http-equiv=\"pragma\" content=\"no-cache\"/>
  <meta http-equiv=\"expires\" content=\"0\"/>
</head><body>
<script>
  (function(){
    // Default start page: Projects
    var NEXT = '/projects';
    try {
      if (sessionStorage.getItem('aistate_intro_seen') === '1') {
        location.replace(NEXT);
      } else {
        location.replace('/static/Intro.html?next=' + encodeURIComponent(NEXT));
      }
    } catch (e) {
      // If storage is blocked, just show Intro and then continue.
      location.replace('/static/Intro.html?next=' + encodeURIComponent(NEXT));
    }
  })();
</script>
</body></html>"""
    return HTMLResponse(html)


@app.get("/login", response_class=HTMLResponse)
def page_login(request: Request) -> Any:
    return TEMPLATES.TemplateResponse("login.html", {
        "request": request, "app_name": APP_NAME, "app_version": APP_VERSION,
    })


@app.get("/change-password", response_class=HTMLResponse)
def page_change_password(request: Request) -> Any:
    """Standalone page for forced password change (first login / expired)."""
    return TEMPLATES.TemplateResponse("change_password.html", {
        "request": request, "app_name": APP_NAME, "app_version": APP_VERSION,
    })


@app.get("/setup", response_class=HTMLResponse)
def page_setup(request: Request) -> Any:
    return TEMPLATES.TemplateResponse("setup.html", {
        "request": request, "app_name": APP_NAME, "app_version": APP_VERSION,
    })


@app.get("/banned", response_class=HTMLResponse)
def page_banned(request: Request) -> Any:
    reason = request.query_params.get("reason", "")
    until = request.query_params.get("until", "")
    return TEMPLATES.TemplateResponse("banned.html", {
        "request": request, "app_name": APP_NAME, "app_version": APP_VERSION,
        "reason": reason, "until": until,
    })


@app.get("/register", response_class=HTMLResponse)
def page_register(request: Request) -> Any:
    return TEMPLATES.TemplateResponse("register.html", {
        "request": request, "app_name": APP_NAME, "app_version": APP_VERSION,
    })


@app.get("/pending", response_class=HTMLResponse)
def page_pending(request: Request) -> Any:
    return TEMPLATES.TemplateResponse("pending.html", {
        "request": request, "app_name": APP_NAME, "app_version": APP_VERSION,
    })


@app.get("/users", response_class=HTMLResponse)
def page_users(request: Request) -> Any:
    return render_page(request, "users.html", "Zarządzanie użytkownikami", "users")


_BREADCRUMBS = {
    "projects":       [{"label": "Projekty", "href": "/projects"}],
    "new_project":    [{"label": "Projekty", "href": "/projects"}, {"label": "Nowy (legacy)"}],
    "transcription":  [{"label": "Projekty", "href": "/projects"}, {"label": "Transkrypcja"}],
    "diarization":    [{"label": "Projekty", "href": "/projects"}, {"label": "Diaryzacja"}],
    "analysis":       [{"label": "Projekty", "href": "/projects"}, {"label": "Analiza"}],
    "chat":           [{"label": "Chat LLM"}],
    "translation":    [{"label": "Tłumaczenie"}],
    "settings":       [{"label": "Ustawienia"}],
    "admin":          [{"label": "Admin", "href": "/admin"}, {"label": "GPU"}],
    "llm_settings":   [{"label": "Admin", "href": "/admin"}, {"label": "LLM"}],
    "asr_settings":   [{"label": "Admin", "href": "/admin"}, {"label": "ASR"}],
    "nllb_settings":  [{"label": "Admin", "href": "/admin"}, {"label": "NLLB"}],
    "tts_settings":   [{"label": "Admin", "href": "/admin"}, {"label": "TTS"}],
    "logs":           [{"label": "Admin", "href": "/admin"}, {"label": "Logi"}],
    "users":          [{"label": "Admin", "href": "/admin"}, {"label": "Użytkownicy"}],
    "info":           [{"label": "Info"}],
}

def render_page(request: Request, tpl: str, title: str, active: str, current_project: Optional[str] = None, **ctx: Any):
    settings = load_settings()
    # Multi-user context
    multiuser = getattr(request.state, "multiuser", False)
    user = getattr(request.state, "user", None)
    user_modules = get_user_modules(user.role, user.is_admin, user.admin_roles, user.is_superadmin) if user else []
    breadcrumbs = _BREADCRUMBS.get(active, [])
    resp = TEMPLATES.TemplateResponse(
        tpl,
        {
            "request": request,
            "title": title,
            "active": active,
            "app_name": APP_NAME,
            "app_fullname": "Artificial Intelligence Speech‑To‑Analysis‑Translation‑Engine",
            "app_version": APP_VERSION,
            "static_ts": int(time.time()),
            "whisper_models": WHISPER_MODELS,
            "default_whisper_model": getattr(settings, "whisper_model", "large-v3") or "large-v3",
            "nemo_models": NEMO_MODELS,
            "nemo_diarization_models": NEMO_DIARIZATION_MODELS,
            "pyannote_pipelines": PYANNOTE_PIPELINES,
            "nllb_fast_models": NLLB_FAST_MODELS,
            "nllb_accurate_models": NLLB_ACCURATE_MODELS,
            "current_project": current_project,
            "multiuser": multiuser,
            "user": user,
            "user_modules": user_modules,
            "breadcrumbs": breadcrumbs,
            **ctx,
        },
    )
    # Prevent browser from caching HTML pages (ensures fresh static_ts on each load)
    resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    resp.headers["Pragma"] = "no-cache"
    resp.headers["Expires"] = "0"
    return resp




# --- Legacy Polish routes (compat) ---
@app.get("/transkrypcja", include_in_schema=False)
def legacy_transkrypcja():
    return RedirectResponse(url="/transcription")

@app.get("/nowy-projekt", include_in_schema=False)
def legacy_nowy_projekt():
    return RedirectResponse(url="/projects")

@app.get("/diaryzacja", include_in_schema=False)
def legacy_diaryzacja():
    return RedirectResponse(url="/diarization")

@app.get("/ustawienia", include_in_schema=False)
def legacy_ustawienia():
    return RedirectResponse(url="/settings")


@app.get("/ustawienia-llm", include_in_schema=False)
def legacy_ustawienia_llm():
    return RedirectResponse(url="/llm-settings")

@app.get("/logi", include_in_schema=False)
def legacy_logi():
    return RedirectResponse(url="/logs")

@app.get("/zapis", include_in_schema=False)
def legacy_zapis():
    return RedirectResponse(url="/save")

@app.get("/transcription", response_class=HTMLResponse)
def page_transcribe(request: Request) -> Any:
    return render_page(request, "transcription.html", "Transkrypcja", "transcription")


@app.get("/new-project", response_class=HTMLResponse)
def page_new_project(request: Request) -> Any:
    return render_page(request, "new_project.html", "Nowy projekt", "new_project")


@app.get("/projects", response_class=HTMLResponse)
def page_projects(request: Request) -> Any:
    return render_page(request, "projects.html", "Projekty", "projects")


@app.get("/projects/{workspace_id}", response_class=HTMLResponse)
def page_workspace(request: Request, workspace_id: str) -> Any:
    return render_page(request, "projects.html", "Projekt", "projects")


@app.get("/diarization", response_class=HTMLResponse)
def page_diarize(request: Request) -> Any:
    return render_page(request, "diarization.html", "Diaryzacja", "diarization")



@app.get("/chat", response_class=HTMLResponse)
def page_chat(request: Request) -> Any:
    return render_page(request, "chat.html", "Chat LLM", "chat")


@app.get("/analysis", response_class=HTMLResponse)
def page_analysis(request: Request) -> Any:
    return render_page(request, "analysis.html", "Analiza", "analysis")


@app.get("/analiza", response_class=HTMLResponse)
def page_analiza(request: Request) -> Any:
    return page_analysis(request)

@app.get("/translation", response_class=HTMLResponse)
def page_translation(request: Request) -> Any:
    """Translation page"""
    return render_page(request, "translation.html", "Tłumaczenie", "translation")

@app.get("/settings", response_class=HTMLResponse)
def page_settings(request: Request) -> Any:
    s = load_settings()
    return render_page(request, "settings.html", "Ustawienia", "settings", settings=s, data_dir=str(DATA_DIR))


@app.get("/llm-settings", response_class=HTMLResponse)
def page_llm_settings(request: Request) -> Any:
    # Dedicated LLM/Ollama settings page (UI-only split from the general Settings tab)
    return render_page(request, "llm_settings.html", "Ustawienia LLM", "llm_settings")


@app.get("/asr-settings", response_class=HTMLResponse)
def page_asr_settings(request: Request) -> Any:
    # Admin panel: ASR engines (Whisper / NeMo / pyannote) management
    s = load_settings()
    return render_page(request, "asr_settings.html", "Ustawienia ASR", "asr_settings", settings=s)


@app.get("/nllb-settings", response_class=HTMLResponse)
def page_nllb_settings(request: Request) -> Any:
    # Admin panel: NLLB translation models management
    return render_page(request, "nllb_settings.html", "Ustawienia NLLB", "nllb_settings")


@app.get("/tts-settings", response_class=HTMLResponse)
def page_tts_settings(request: Request) -> Any:
    # Admin panel: TTS engine management
    return render_page(request, "tts_settings.html", "Ustawienia TTS", "tts_settings")


@app.get("/admin", response_class=HTMLResponse)
def page_admin(request: Request) -> Any:
    return render_page(request, "admin.html", "Ustawienia GPU", "admin")


@app.get("/logs", response_class=HTMLResponse)
def page_logs(request: Request) -> Any:
    return render_page(request, "logs.html", "Logi", "logs")


@app.get("/save", response_class=HTMLResponse)
def page_save(request: Request) -> Any:
    return render_page(request, "save.html", "Zapis", "save")


@app.get("/info", response_class=HTMLResponse)
def page_info(request: Request) -> Any:
    # language priority:
    # 1) explicit ?lang=en|pl
    # 2) UI language from settings
    lang = (request.query_params.get("lang") or "").lower().strip()
    if not lang:
        try:
            s = load_settings()
            lang = (getattr(s, "ui_language", "") or "").lower().strip()
        except Exception:
            lang = ""
    if not lang:
        lang = "pl"

    static_root = ROOT / "webapp" / "static"
    md_path = static_root / ("info_en.md" if lang.startswith("en") else "info_pl.md")

    # safety fallback if file missing
    if not md_path.exists():
        md_path = static_root / "info_pl.md"

    source = str(md_path.relative_to(ROOT))
    text = md_path.read_text(encoding="utf-8", errors="ignore")

    # optional placeholders
    text = (text
        .replace("{APP_NAME}", APP_NAME)
        .replace("{APP_VERSION}", APP_VERSION)
    )

    if md_to_html:
        html = md_to_html(text, extensions=["fenced_code", "tables"])
    else:
        html = "<pre>" + text.replace("&", "&amp;").replace("<", "&lt;") + "</pre>"

    return render_page(request, "info.html", "Info", "info", content=html, source=source)


# ---------- API: translation (NLLB) ----------

TRANSLATION_TMP_DIR = (DATA_DIR / "_tmp" / "translation").resolve()
TRANSLATION_UPLOAD_DIR = (DATA_DIR / "_tmp" / "translation_uploads").resolve()
TRANSLATION_TMP_DIR.mkdir(parents=True, exist_ok=True)
TRANSLATION_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

TRANSLATION_DRAFTS_DIR = (DATA_DIR / "_tmp" / "translation_drafts").resolve()
TRANSLATION_DRAFTS_DIR.mkdir(parents=True, exist_ok=True)


def _safe_client_id(client_id: str) -> str:
    cid = str(client_id or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{8,64}", cid):
        raise HTTPException(status_code=400, detail="Invalid client_id")
    return cid


def translation_global_draft_path(client_id: str) -> Path:
    return TRANSLATION_DRAFTS_DIR / f"{_safe_client_id(client_id)}.json"


def read_translation_global_draft(client_id: str) -> Dict[str, Any]:
    path = translation_global_draft_path(client_id)
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}
    return {}


def write_translation_global_draft(client_id: str, draft: Dict[str, Any]) -> None:
    path = translation_global_draft_path(client_id)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(draft, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)



def _nllb_model_installed(mode: str, model_id: str) -> bool:
    mode = str(mode or "fast").lower().strip()
    if mode not in ("fast", "accurate"):
        mode = "fast"
    model_id = str(model_id or "").strip()
    if not model_id:
        return False
    try:
        reg = _read_nllb_registry()
        if not reg:
            reg = scan_and_persist_nllb_model_cache()
        return bool((reg.get(mode) or {}).get(model_id))
    except Exception:
        return False




@app.get("/api/translation/draft/{client_id}")
def api_get_translation_global_draft(client_id: str) -> Any:
    """Load translation draft not tied to a project (fallback when localStorage quota is exceeded)."""
    draft = read_translation_global_draft(client_id)
    return {"draft": draft or None}


@app.post("/api/translation/draft/{client_id}")
async def api_save_translation_global_draft(client_id: str, request: Request) -> Any:
    """Save translation draft not tied to a project (fallback when localStorage quota is exceeded)."""
    payload = await request.json()
    draft = payload.get("draft") if isinstance(payload, dict) else None
    if draft is None:
        draft = payload
    if not isinstance(draft, dict):
        raise HTTPException(status_code=400, detail="draft must be a JSON object (dict).")

    if "saved_at" not in draft:
        draft["saved_at"] = int(time.time() * 1000)

    write_translation_global_draft(client_id, draft)
    return {"ok": True}

@app.post("/api/translation/upload")
async def api_translation_upload(
    file: UploadFile = File(...),
    source_lang: str = Form("auto"),
    target_langs: str = Form(""),
    mode: str = Form("fast"),
) -> Any:
    """Upload document and extract text for translation input."""
    if not file:
        raise HTTPException(status_code=400, detail="Brak pliku")

    # Save to temp
    name = str(file.filename or "upload")
    ext = (Path(name).suffix or "").lower()
    tmp_path = (TRANSLATION_UPLOAD_DIR / f"{uuid.uuid4().hex}{ext}").resolve()
    try:
        content = await file.read()
        tmp_path.write_bytes(content)

        # Extract using translation handlers (supports TXT/DOCX/PDF/SRT)
        from backend.translation.document_handlers import extract_text_from_file  # type: ignore

        text = await run_in_threadpool(extract_text_from_file, tmp_path)
        return {"ok": True, "text": text, "filename": name, "ext": ext, "size": len(content)}
    except Exception as e:
        app_log(f"Translation upload error: {e}")
        return {"ok": False, "error": "Nie udało się przetworzyć pliku."}
    finally:
        try:
            tmp_path.unlink(missing_ok=True)  # type: ignore[arg-type]
        except Exception:
            pass


@app.post("/api/translation/translate")
async def api_translation_translate(
    text: str = Form(...),
    source_lang: str = Form("auto"),
    target_langs: str = Form(""),
    mode: str = Form("fast"),
    nllb_model: str = Form(""),
    generate_summary: str = Form("0"),
    summary_detail: str = Form("5"),
    use_glossary: str = Form("0"),
    preserve_formatting: str = Form("1"),
) -> Any:
    """Start translation task (subprocess) and return a task id."""
    t = str(text or "").strip()
    if not t:
        raise HTTPException(status_code=400, detail="Brak tekstu do przetłumaczenia")

    mode = str(mode or "fast").lower().strip()
    if mode not in ("fast", "accurate"):
        mode = "fast"

    model = str(nllb_model or "").strip()
    if not model:
        raise HTTPException(status_code=400, detail="Wybierz model NLLB")

    if not _nllb_model_installed(mode, model):
        raise HTTPException(status_code=400, detail="Wybrany model NLLB nie jest zainstalowany (Ustawienia NLLB)")

    targets = [x.strip() for x in str(target_langs or "").split(",") if x.strip()]
    if not targets:
        raise HTTPException(status_code=400, detail="Wybierz przynajmniej jeden język docelowy")

    gs = str(generate_summary or "0").lower().strip() in ("1", "true", "yes", "on")
    ug = str(use_glossary or "0").lower().strip() in ("1", "true", "yes", "on")
    pf = str(preserve_formatting or "1").lower().strip() in ("1", "true", "yes", "on")
    try:
        sd = int(summary_detail)
    except Exception:
        sd = 5
    sd = max(1, min(10, sd))

    payload_path = (TRANSLATION_TMP_DIR / f"translate_{uuid.uuid4().hex}.json").resolve()
    payload = {
        "text": t,
        "source_lang": str(source_lang or "auto").lower().strip(),
        "target_langs": targets,
        "mode": mode,
        "nllb_model": model,
        "generate_summary": gs,
        "summary_detail": sd,
        "use_glossary": ug,
        "preserve_formatting": pf,
    }
    payload_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    worker = ROOT / "backend" / "translation_worker.py"
    require_existing_file(worker, "Brak translation_worker.py")

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--input", str(payload_path),
    ]

    env = {
        # Prefer offline behaviour: if model is not cached, worker will fail fast.
        "TRANSFORMERS_OFFLINE": os.environ.get("TRANSFORMERS_OFFLINE", "1"),
        "HF_HUB_DISABLE_TELEMETRY": os.environ.get("HF_HUB_DISABLE_TELEMETRY", "1"),
        # Ensure local imports (backend.*) work in subprocess even without PYTHONPATH set.
        "PYTHONPATH": os.environ.get("PYTHONPATH", str(ROOT)),
    }

    app_log(f"Translation requested: mode={mode}, model={model}, targets={','.join(targets)}")
    if GPU_RM.enabled:
        task = GPU_RM.enqueue_subprocess(kind=f"translate_{mode}", project_id="-", cmd=cmd, cwd=ROOT, env=env)
    else:
        task = TASKS.start_subprocess(kind=f"translate_{mode}", project_id="-", cmd=cmd, cwd=ROOT, env=env)

    # Best-effort cleanup of the temp payload file after task ends.
    def _cleanup() -> None:
        try:
            for _ in range(3600):  # up to ~1h
                try:
                    tt = TASKS.get(task.task_id)
                    if tt.status in ("done", "error"):
                        break
                except Exception:
                    break
                time.sleep(1)
        finally:
            try:
                payload_path.unlink(missing_ok=True)  # type: ignore[arg-type]
            except Exception:
                pass

    threading.Thread(target=_cleanup, daemon=True).start()
    return {"task_id": task.task_id}


@app.get("/api/translation/progress/{task_id}")
def api_translation_progress(task_id: str) -> Any:
    """Return translation task status in the format expected by translation.js."""
    tid = str(task_id or "").strip()
    if not tid:
        raise HTTPException(status_code=400, detail="task_id required")
    try:
        t = TASKS.get(tid)
    except Exception:
        return {"status": "failed", "progress": 0, "error": "Unknown task"}

    if t.status == "done":
        result = t.result or {}
        return {
            "status": "completed",
            "progress": 100,
            "results": result.get("results") or {},
            "summary": result.get("summary"),
            "meta": {
                "mode": result.get("mode"),
                "nllb_model": result.get("nllb_model"),
                "source_lang": result.get("source_lang"),
                "detected_source_lang": result.get("detected_source_lang"),
            },
        }

    if t.status == "error":
        return {"status": "failed", "progress": int(t.progress or 0), "error": t.error or "Translation failed"}

    # running
    return {"status": "processing", "progress": int(t.progress or 0)}


@app.post("/api/translation/export")
async def api_translation_export(
    text: str = Form(...),
    format: str = Form("txt"),
    filename: str = Form("translation"),
) -> Any:
    """Export translated text as TXT/HTML/DOCX (returned as a downloadable file)."""
    fmt = str(format or "txt").lower().strip()
    if fmt == "doc":
        fmt = "docx"
    name = (str(filename or "translation").strip() or "translation").replace("/", "_")

    from html import escape as _html_escape
    import io as _io

    if fmt == "txt":
        data = (text or "").encode("utf-8")
        return StreamingResponse(
            _io.BytesIO(data),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename=\"{name}.txt\""},
        )

    if fmt == "html":
        body = _html_escape(text or "")
        html = f"<!doctype html><html><head><meta charset='utf-8'></head><body><pre>{body}</pre></body></html>"
        data = html.encode("utf-8")
        return StreamingResponse(
            _io.BytesIO(data),
            media_type="text/html; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename=\"{name}.html\""},
        )

    if fmt == "docx":
        try:
            from docx import Document  # type: ignore
        except Exception:
            raise HTTPException(status_code=500, detail="python-docx nie jest zainstalowany")

        doc = Document()
        for line in (text or "").splitlines() or [text or ""]:
            doc.add_paragraph(line)
        buf = _io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename=\"{name}.docx\""},
        )

    raise HTTPException(status_code=400, detail="Unsupported format")

# ---------- API: settings ----------

@app.get("/api/settings")
def api_get_settings() -> Any:
    s = load_settings()
    return {"hf_token_present": bool(getattr(s, "hf_token", "")), "whisper_model": getattr(s, "whisper_model", "large-v3")}


@app.post("/api/settings")
def api_save_settings(payload: Dict[str, Any]) -> Any:
    s = load_settings()
    if "hf_token" in payload:
        s.hf_token = str(payload.get("hf_token") or "")
    if "whisper_model" in payload:
        s.whisper_model = str(payload.get("whisper_model") or "large-v3")
    if "ui_language" in payload:
        s.ui_language = str(payload.get("ui_language") or "pl")
    # Security settings
    if "account_lockout_threshold" in payload:
        s.account_lockout_threshold = max(0, int(payload["account_lockout_threshold"]))
    if "account_lockout_duration" in payload:
        s.account_lockout_duration = max(1, int(payload["account_lockout_duration"]))
    if "password_policy" in payload:
        pp = str(payload["password_policy"])
        if pp in ("none", "basic", "medium", "strong"):
            s.password_policy = pp
    if "password_expiry_days" in payload:
        s.password_expiry_days = max(0, int(payload["password_expiry_days"]))
    save_settings(s)
    return {"ok": True}


@app.get("/api/settings/security")
def api_get_security_settings(request: Request) -> Any:
    """Return security policy settings (admin only)."""
    user = getattr(request.state, "user", None)
    if user and not user.is_admin and not user.is_superadmin:
        return JSONResponse({"status": "error", "message": "Admin access required"}, status_code=403)
    s = load_settings()
    return {
        "status": "ok",
        "account_lockout_threshold": getattr(s, "account_lockout_threshold", 5),
        "account_lockout_duration": getattr(s, "account_lockout_duration", 15),
        "password_policy": getattr(s, "password_policy", "basic"),
        "password_expiry_days": getattr(s, "password_expiry_days", 0),
        "session_timeout_hours": getattr(s, "session_timeout_hours", 8),
    }


# ---------- API: ASR engines (Whisper / NeMo / pyannote) ----------

def _pkg_info(module_name: str) -> Dict[str, Any]:
    import importlib
    import importlib.util

    try:
        spec = importlib.util.find_spec(module_name)
        if spec is None:
            return {"installed": False, "version": None}
    except Exception:
        return {"installed": False, "version": None}

    # Try to import for version
    version = None
    try:
        mod = importlib.import_module(module_name)
        version = getattr(mod, "__version__", None)
    except Exception:
        version = None
    return {"installed": True, "version": version}




# --- ASR model cache registry (installed/cached models) ---
# We persist cache scan results under DATA_DIR/projects/_global/asr_models.json
# so the UI can show “(not installed)” per-model and hide Install buttons when cached.
ASR_REGISTRY_REL = Path("_global") / "asr_models.json"

# NLLB model cache registry (installed/cached NLLB models)
NLLB_REGISTRY_REL = Path("_global") / "nllb_models.json"


def _asr_registry_path() -> Path:
    p = (PROJECTS_DIR / ASR_REGISTRY_REL).resolve()
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _read_asr_registry() -> Dict[str, Any]:
    p = _asr_registry_path()
    if not p.exists():
        return {}
    try:
        obj = json.loads(p.read_text(encoding="utf-8"))
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def _write_asr_registry(obj: Dict[str, Any]) -> None:
    p = _asr_registry_path()
    try:
        p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        # best effort
        pass


def _nllb_registry_path() -> Path:
    p = (PROJECTS_DIR / NLLB_REGISTRY_REL).resolve()
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _read_nllb_registry() -> Dict[str, Any]:
    p = _nllb_registry_path()
    if not p.exists():
        return {}
    try:
        obj = json.loads(p.read_text(encoding="utf-8"))
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def _write_nllb_registry(obj: Dict[str, Any]) -> None:
    p = _nllb_registry_path()
    try:
        p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        # best effort
        pass


def _cache_base_dir() -> Path:
    # XDG compatible cache base
    home = Path.home()
    xdg = os.environ.get("XDG_CACHE_HOME")
    return Path(xdg).expanduser().resolve() if xdg else (home / ".cache")


def _scan_whisper_cache() -> Dict[str, bool]:
    base = _cache_base_dir() / "whisper"
    stems = set()
    if base.exists() and base.is_dir():
        for p in base.glob("*.pt"):
            stems.add(p.stem.lower())
            stems.add(p.name.lower())
    out: Dict[str, bool] = {}
    for m in WHISPER_MODELS:
        key = str(m).lower().strip()
        if key == "turbo":
            out[m] = ("large-v3-turbo" in stems) or ("turbo" in stems) or ("large-v3-turbo.pt" in stems) or ("turbo.pt" in stems)
        else:
            out[m] = (key in stems) or (f"{key}.pt" in stems)
    return out


def _scan_nemo_cache(model_ids: List[str]) -> Dict[str, bool]:
    """Best-effort scan for cached NeMo models.

    Why this is heuristic:
    - NeMo historically cached pretrained models under ~/.cache/torch/NeMo/... as a *directory*
      (often named after the model id, e.g. "diar_msdd_telephonic"), not always leaving a *.nemo file.
    - Some environments override torch cache locations (TORCH_HOME).

    We therefore scan for both:
    - filenames with typical extensions (".nemo", ".ckpt", ".pt", ...)
    - directory names that include the model short id
    """


    # Deterministic marker-based detection (written by backend/asr_worker.py)
    # This avoids relying only on NeMo's internal cache structure.
    marker_dir = (ROOT / "backend" / "models_cache" / "nemo").resolve()
    marker_models = set()
    if marker_dir.exists() and marker_dir.is_dir():
        for mp in marker_dir.glob("*.json"):
            try:
                obj = json.loads(mp.read_text(encoding="utf-8"))
                m = str(obj.get("model") or "").strip().lower()
                if m:
                    marker_models.add(m)
                    marker_models.add(m.split("/")[-1])
            except Exception:
                # Fallback: use filename stem
                marker_models.add(mp.stem.lower())

    base = _cache_base_dir()

    cand: List[Path] = []
    torch_home = os.environ.get("TORCH_HOME")
    if torch_home:
        # torch hub style override
        cand.append((Path(torch_home).expanduser().resolve() / "NeMo").resolve())

    # Common NeMo locations
    cand.extend([
        base / "torch" / "NeMo",
        base / "torch" / "nemo",
        base / "nemo",
        base / "NeMo",
    ])

    names: List[str] = []
    exts = {".nemo", ".ckpt", ".pt", ".pth", ".onnx", ".yaml", ".yml", ".json"}

    for d in cand:
        if not (d.exists() and d.is_dir()):
            continue
        try:
            for p in d.rglob("*"):
                try:
                    n = p.name.lower()
                except Exception:
                    continue

                # Directory-based caches are common in NeMo.
                if p.is_dir():
                    names.append(n)
                    continue

                # File-based caches
                suf = p.suffix.lower()
                if suf in exts:
                    names.append(n)
                    try:
                        names.append(p.stem.lower())
                    except Exception:
                        pass
        except Exception:
            continue

    name_set = set(names)

    out: Dict[str, bool] = {}
    for mid in model_ids:
        mid_s = str(mid)
        short = mid_s.split("/")[-1].lower()
        # Heuristic: NeMo caches usually include the short model name in filename
        hit = (mid_s.lower() in marker_models) or (short in marker_models)
        for n in name_set:
            if short and short in n:
                hit = True
                break
            # diarization model ids are already short
            if mid_s.lower() in n:
                hit = True
                break
        out[mid] = hit
    return out


def _hf_hub_dir() -> Path:
    """Return HuggingFace Hub cache directory.

    We prefer the resolved path used by huggingface_hub itself (HF_HUB_CACHE),
    then fall back to env vars and common defaults.
    """
    # Preferred: ask huggingface_hub for its resolved cache path.
    try:
        from huggingface_hub.constants import HF_HUB_CACHE  # type: ignore
        return Path(str(HF_HUB_CACHE)).expanduser().resolve()
    except Exception:
        pass

    # Fallback: common env vars (older/newer naming).
    for env in ("HF_HUB_CACHE", "HUGGINGFACE_HUB_CACHE"):
        v = os.environ.get(env)
        if v:
            return Path(v).expanduser().resolve()

    hf_home = os.environ.get("HF_HOME")
    if hf_home:
        return (Path(hf_home).expanduser() / "hub").resolve()

    # Final fallback: typical default on Linux.
    return (_cache_base_dir() / "huggingface" / "hub").resolve()


def _scan_hf_pipelines(pipelines: List[str]) -> Dict[str, bool]:
    """Best-effort scan for cached HuggingFace pipelines.

    We combine two signals:
    1) HF hub cache directory existence (fast heuristic).
    2) Deterministic marker files written by backend/asr_worker.py after a successful
       Pipeline.from_pretrained() call.

    Markers make the UI robust when HF cache is moved via env vars or when the hub cache
    layout differs between users.
    """
    hub = _hf_hub_dir()

    # Marker-based detection (written by backend/asr_worker.py)
    marker_dir = (ROOT / "backend" / "models_cache" / "pyannote").resolve()
    marker_pipelines = set()
    if marker_dir.exists() and marker_dir.is_dir():
        for mp in marker_dir.glob("*.json"):
            try:
                obj = json.loads(mp.read_text(encoding="utf-8"))
                pid = str(obj.get("pipeline") or obj.get("model") or obj.get("id") or "").strip()
                if pid:
                    marker_pipelines.add(pid)
            except Exception:
                continue

    out: Dict[str, bool] = {}
    for pid in pipelines:
        s = str(pid)
        ok = s in marker_pipelines

        if "/" in s:
            org, name = s.split("/", 1)
            d = hub / f"models--{org}--{name}"
            if d.exists() and d.is_dir():
                ok = True

        out[pid] = ok

    return out


def _scan_hf_models(model_ids: List[str]) -> Dict[str, bool]:
    """Best-effort scan for cached HuggingFace models.

    Signals:
    1) Marker files written by backend/nllb_worker.py
    2) HF hub cache directory existence
    """
    hub = _hf_hub_dir()

    marker_dir = (ROOT / "backend" / "models_cache" / "nllb").resolve()
    marker_models = set()
    if marker_dir.exists() and marker_dir.is_dir():
        for mp in marker_dir.glob("*.json"):
            try:
                obj = json.loads(mp.read_text(encoding="utf-8"))
                mid = str(obj.get("model") or obj.get("id") or "").strip()
                if mid:
                    marker_models.add(mid)
            except Exception:
                continue

    out: Dict[str, bool] = {}
    for mid in model_ids:
        s = str(mid)
        ok = s in marker_models
        if "/" in s:
            org, name = s.split("/", 1)
            d = hub / f"models--{org}--{name}"
            if d.exists() and d.is_dir():
                ok = True
        out[mid] = ok
    return out


def scan_and_persist_nllb_model_cache() -> Dict[str, Any]:
    # Scan all configured models and split into groups
    all_ids = list(dict.fromkeys(list(NLLB_FAST_MODELS) + list(NLLB_ACCURATE_MODELS)))
    scanned = _scan_hf_models(all_ids)
    state = {
        "fast": {m: bool(scanned.get(m)) for m in NLLB_FAST_MODELS},
        "accurate": {m: bool(scanned.get(m)) for m in NLLB_ACCURATE_MODELS},
        "last_scan": now_iso(),
    }
    _write_nllb_registry(state)
    return state


def scan_and_persist_asr_model_cache() -> Dict[str, Any]:
    # Build full state map
    state = {
        "whisper": _scan_whisper_cache(),
        "nemo": _scan_nemo_cache(NEMO_MODELS),
        "nemo_diar": _scan_nemo_cache(NEMO_DIARIZATION_MODELS),
        "pyannote": _scan_hf_pipelines(PYANNOTE_PIPELINES),
        "last_scan": now_iso(),
    }
    _write_asr_registry(state)
    return state

@app.get("/api/asr/status")
def api_asr_status() -> Any:
    s = load_settings()
    hf_token = getattr(s, "hf_token", "") or ""
    return {
        "whisper": _pkg_info("whisper"),
        "nemo": _pkg_info("nemo"),
        # NeMo diarization shares the same Python package as NeMo ASR
        "nemo_diar": _pkg_info("nemo"),
        "pyannote": _pkg_info("pyannote.audio"),
        "hf_token_present": bool(hf_token),
    }




@app.get("/api/asr/models_state")
def api_asr_models_state(refresh: int = 0) -> Any:
    """Return per-model cached/installed state for ASR.

    refresh=1 forces a filesystem scan (Whisper/NeMo/HF cache)."""
    if refresh:
        return scan_and_persist_asr_model_cache()
    reg = _read_asr_registry()
    if not reg:
        reg = scan_and_persist_asr_model_cache()
    return reg

@app.get("/api/asr/installed/diarization")
def api_asr_installed_diarization(refresh: int = 0) -> Any:
    """Return *installed* diarization + ASR options for the UI.

    This endpoint is used by the Diarization tab to avoid any on-the-fly downloads.
    Users must install engines/models only in ASR Settings.

    refresh=1 forces a filesystem cache scan (Whisper/NeMo/HF hub).
    """
    reg = api_asr_models_state(refresh=refresh)  # includes cached model state
    status = api_asr_status()

    def _installed_dict(d: Any) -> dict:
        return d if isinstance(d, dict) else {}

    reg_whisper = _installed_dict(reg.get('whisper'))
    reg_nemo = _installed_dict(reg.get('nemo'))
    reg_nemo_diar = _installed_dict(reg.get('nemo_diar'))
    reg_pyannote = _installed_dict(reg.get('pyannote'))

    # --- diarization engines ---
    diar_engines = []

    # Text/simple diarization is always available (built-in)
    diar_engines.append({
        'id': 'text',
        'label': 'Text (simple)',
        'installed': True,
        'supports_language': False,
        'models': [
            {'id': 'simple', 'label': 'simple', 'installed': True, 'supports_language': False},
        ],
    })

    # pyannote diarization (requires pyannote.audio). We keep the engine visible even
    # when no pipelines are cached yet (so the UI can direct the user to ASR Settings).
    py_pkg_ok = bool((status.get('pyannote') or {}).get('installed'))
    py_models = [mid for mid, ok in reg_pyannote.items() if ok]
    if py_pkg_ok:
        diar_engines.append({
            'id': 'pyannote',
            'label': 'pyannote',
            'installed': bool(py_models),
            'supports_language': False,
            'models': [
                {'id': mid, 'label': mid, 'installed': True, 'supports_language': False}
                for mid in py_models
            ],
        })

    # NeMo diarization (requires nemo + diarization model cache)
    nemo_pkg_ok = bool((status.get('nemo') or {}).get('installed'))
    nd_models = [mid for mid, ok in reg_nemo_diar.items() if ok]
    if nemo_pkg_ok and nd_models:
        diar_engines.append({
            'id': 'nemo_diar',
            'label': 'NeMo diarization',
            'installed': True,
            'supports_language': False,
            'models': [
                {'id': mid, 'label': mid, 'installed': True, 'supports_language': False}
                for mid in nd_models
            ],
        })

    # --- ASR engines/models (installed only) ---
    asr_engines = []
    asr_models = {'whisper': [], 'nemo': []}

    if bool((status.get('whisper') or {}).get('installed')):
        asr_engines.append({'id': 'whisper', 'label': 'Whisper', 'installed': True, 'supports_language': True})
        asr_models['whisper'] = [m for m, ok in reg_whisper.items() if ok]

    if bool((status.get('nemo') or {}).get('installed')):
        asr_engines.append({'id': 'nemo', 'label': 'NeMo', 'installed': True, 'supports_language': False})
        asr_models['nemo'] = [m for m, ok in reg_nemo.items() if ok]

    return {
        'diarization_engines': diar_engines,
        'asr_engines': asr_engines,
        'asr_models': asr_models,
        'hf_token_present': bool(status.get('hf_token_present')),
        'last_scan': reg.get('last_scan'),
    }

@app.post("/api/asr/install")
def api_asr_install(payload: Dict[str, Any] = Body(...)) -> Any:
    comp = str(payload.get("component") or "").strip().lower()
    if comp not in ("whisper", "nemo", "pyannote"):
        raise HTTPException(status_code=400, detail="Invalid component")

    worker = ROOT / "backend" / "asr_worker.py"
    require_existing_file(worker, "Brak asr_worker.py")

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--action", "install",
        "--component", comp,
    ]

    app_log(f"ASR install requested: component={comp}")
    t = TASKS.start_subprocess(kind=f"asr_install_{comp}", project_id="-", cmd=cmd, cwd=ROOT)
    return {"task_id": t.task_id}


@app.post("/api/asr/predownload")
def api_asr_predownload(payload: Dict[str, Any] = Body(...)) -> Any:
    engine = str(payload.get("engine") or "").strip().lower()
    if engine not in ("whisper", "nemo", "nemo_diar", "pyannote"):
        raise HTTPException(status_code=400, detail="Invalid engine")

    model = str(payload.get("model") or "").strip()
    pipeline = str(payload.get("pipeline") or "").strip()

    worker = ROOT / "backend" / "asr_worker.py"
    require_existing_file(worker, "Brak asr_worker.py")

    s = load_settings()
    hf_token = getattr(s, "hf_token", "") or ""

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--action", "predownload",
        "--engine", engine,
    ]

    if engine in ("whisper", "nemo", "nemo_diar"):
        if not model:
            raise HTTPException(status_code=400, detail="Missing model")
        cmd += ["--model", model]
    else:
        # pyannote
        pid = pipeline or (PYANNOTE_PIPELINES[0] if PYANNOTE_PIPELINES else "")
        if not pid:
            raise HTTPException(status_code=400, detail="Missing pipeline")
        if not hf_token:
            raise HTTPException(status_code=400, detail="Brak tokena HF. Ustaw go w Ustawieniach.")
        cmd += ["--pipeline", pid, "--hf_token", hf_token]

    app_log(f"ASR predownload requested: engine={engine}, model={model or '-'}, pipeline={pipeline or '-'}")
    t = TASKS.start_subprocess(kind=f"asr_predownload_{engine}", project_id="-", cmd=cmd, cwd=ROOT)
    return {"task_id": t.task_id}


# ---------- API: Sound Detection models ----------

SOUND_DETECTION_MODELS = {
    "yamnet": {
        "name": "YAMNet",
        "size_mb": 14,
        "classes": 521,
        "framework": "tensorflow",
        "speed": "fast",
        "accuracy": "good",
    },
    "panns_cnn14": {
        "name": "PANNs CNN14",
        "size_mb": 300,
        "classes": 527,
        "framework": "pytorch",
        "speed": "medium",
        "accuracy": "high",
    },
    "panns_cnn6": {
        "name": "PANNs CNN6",
        "size_mb": 20,
        "classes": 527,
        "framework": "pytorch",
        "speed": "fast",
        "accuracy": "good",
    },
    "beats": {
        "name": "BEATs",
        "size_mb": 90,
        "classes": 527,
        "framework": "pytorch",
        "speed": "slow",
        "accuracy": "highest",
    },
}


def _sound_detection_registry_path() -> Path:
    return PROJECTS_DIR / "_global" / "sound_detection_models.json"


def _read_sound_detection_registry() -> Dict[str, Any]:
    p = _sound_detection_registry_path()
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def _write_sound_detection_registry(data: Dict[str, Any]) -> None:
    p = _sound_detection_registry_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def scan_sound_detection_models() -> Dict[str, bool]:
    """Scan for installed sound detection models via marker files."""
    cache_dir = ROOT / "backend" / "models_cache" / "sound_detection"
    result = {}

    for model_id in SOUND_DETECTION_MODELS:
        marker = cache_dir / f"{model_id}.json"
        result[model_id] = marker.exists()

    # Persist to registry
    reg = {"models": result, "last_scan": now_iso()}
    _write_sound_detection_registry(reg)

    return result


@app.get("/api/sound-detection/status")
def api_sound_detection_status() -> Any:
    """Return dependency status for sound detection frameworks."""
    return {
        "tensorflow": _pkg_info("tensorflow"),
        "tensorflow_hub": _pkg_info("tensorflow_hub"),
        "panns_inference": _pkg_info("panns_inference"),
        "transformers": _pkg_info("transformers"),
    }


@app.get("/api/sound-detection/models")
def api_sound_detection_models() -> Any:
    """Return available sound detection models with metadata."""
    return SOUND_DETECTION_MODELS


@app.get("/api/sound-detection/models_state")
def api_sound_detection_models_state(refresh: int = 0) -> Any:
    """Return cached model-state map for sound detection models."""
    if refresh:
        return scan_sound_detection_models()
    reg = _read_sound_detection_registry()
    if not reg or not reg.get("models"):
        try:
            return scan_sound_detection_models()
        except Exception:
            return {}
    return reg.get("models", {})


@app.post("/api/sound-detection/install")
def api_sound_detection_install(payload: Dict[str, Any] = Body(...)) -> Any:
    """Install dependencies for a sound detection model."""
    try:
        model_id = str(payload.get("model") or "").strip()
        app_log(f"Sound detection install requested: model='{model_id}'")

        if not model_id:
            raise HTTPException(status_code=400, detail="Nie podano modelu. Wybierz model z listy.")

        if model_id not in SOUND_DETECTION_MODELS:
            raise HTTPException(status_code=400, detail=f"Nieznany model: {model_id}")

        worker = ROOT / "backend" / "sound_detection_worker.py"
        if not worker.exists():
            app_log(f"Worker not found: {worker}")
            raise HTTPException(status_code=500, detail="Brak pliku sound_detection_worker.py")

        cmd = [
            os.environ.get("PYTHON", sys.executable),
            str(worker),
            "--action", "install",
            "--model", model_id,
        ]

        app_log(f"Starting sound detection install task: cmd={cmd}")
        t = TASKS.start_subprocess(kind=f"sound_detection_install_{model_id}", project_id="-", cmd=cmd, cwd=ROOT)
        app_log(f"Sound detection install task started: task_id={t.task_id}")
        return {"task_id": t.task_id}

    except HTTPException:
        raise
    except Exception as e:
        app_log(f"Sound detection install error: {e}")
        raise HTTPException(status_code=500, detail=f"Błąd instalacji: {e}")


@app.post("/api/sound-detection/predownload")
def api_sound_detection_predownload(payload: Dict[str, Any] = Body(...)) -> Any:
    """Download/cache a sound detection model."""
    try:
        model_id = str(payload.get("model") or "").strip()
        app_log(f"Sound detection predownload requested: model='{model_id}'")

        if not model_id:
            raise HTTPException(status_code=400, detail="Nie podano modelu. Wybierz model z listy.")

        if model_id not in SOUND_DETECTION_MODELS:
            raise HTTPException(status_code=400, detail=f"Nieznany model: {model_id}")

        worker = ROOT / "backend" / "sound_detection_worker.py"
        if not worker.exists():
            app_log(f"Worker not found: {worker}")
            raise HTTPException(status_code=500, detail="Brak pliku sound_detection_worker.py")

        cmd = [
            os.environ.get("PYTHON", sys.executable),
            str(worker),
            "--action", "predownload",
            "--model", model_id,
        ]

        app_log(f"Starting sound detection predownload task: cmd={cmd}")
        t = TASKS.start_subprocess(kind=f"sound_detection_predownload_{model_id}", project_id="-", cmd=cmd, cwd=ROOT)
        app_log(f"Sound detection predownload task started: task_id={t.task_id}")
        return {"task_id": t.task_id}

    except HTTPException:
        raise
    except Exception as e:
        app_log(f"Sound detection predownload error: {e}")
        raise HTTPException(status_code=500, detail=f"Błąd pobierania modelu: {e}")


# ---------- API: TTS (Text-to-Speech) ----------

TTS_ENGINES = {
    "piper": {
        "name": "Piper TTS",
        "packages": ["piper-tts", "pathvalidate"],
        "pip_check": "piper",
        "size_mb": 30,
        "languages": "~50",
        "quality": "good",
        "speed": "very_fast",
        "license": "MIT",
    },
    "mms": {
        "name": "MMS-TTS (Meta)",
        "packages": ["transformers", "torch", "scipy"],
        "pip_check": "transformers",
        "size_mb": 30,
        "languages": "1100+",
        "quality": "ok",
        "speed": "fast",
        "license": "CC-BY-NC-4.0",
    },
    "kokoro": {
        "name": "Kokoro TTS",
        "packages": ["spacy>=3.7,<4.0", "kokoro>=0.3", "soundfile"],
        "pip_check": "kokoro",
        "size_mb": 82,
        "languages": "9",
        "quality": "very_good",
        "speed": "very_fast",
        "license": "Apache 2.0",
    },
}

TTS_REGISTRY_REL = Path("_global") / "tts_models.json"


def _tts_registry_path() -> Path:
    return PROJECTS_DIR / TTS_REGISTRY_REL


def _read_tts_registry() -> Dict[str, Any]:
    p = _tts_registry_path()
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def _write_tts_registry(data: Dict[str, Any]) -> None:
    p = _tts_registry_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def scan_tts_models() -> Dict[str, Any]:
    """Scan cache dir for downloaded TTS models and update registry."""
    cache_dir = ROOT / "backend" / "models_cache" / "tts"
    reg: Dict[str, Any] = {}

    if cache_dir.exists():
        for marker in cache_dir.glob("*.json"):
            try:
                info = json.loads(marker.read_text(encoding="utf-8"))
                engine = info.get("engine", "")
                key = marker.stem  # e.g. "piper_en_US-amy-medium", "mms_pol", "kokoro"
                reg[key] = {
                    "engine": engine,
                    "downloaded": info.get("status") == "ready",
                    "downloaded_at": info.get("downloaded_at", ""),
                    **{k: v for k, v in info.items() if k not in ("status",)},
                }
            except Exception:
                continue

    _write_tts_registry(reg)
    return reg


@app.get("/api/tts/status")
def api_tts_status() -> Any:
    """Return dependency status for TTS engines."""
    return {
        "piper": _pkg_info("piper"),
        "mms": {
            "installed": _pkg_info("transformers").get("installed", False)
            and _pkg_info("torch").get("installed", False),
            "transformers": _pkg_info("transformers"),
            "torch": _pkg_info("torch"),
        },
        "kokoro": _pkg_info("kokoro"),
    }


@app.get("/api/tts/engines")
def api_tts_engines() -> Any:
    """Return TTS engine definitions."""
    return TTS_ENGINES


@app.get("/api/tts/models_state")
def api_tts_models_state(refresh: int = 0) -> Any:
    """Return cached model-state map for TTS voices."""
    if refresh:
        return scan_tts_models()
    reg = _read_tts_registry()
    if not reg:
        try:
            reg = scan_tts_models()
        except Exception:
            reg = {}
    return reg


@app.post("/api/tts/install")
def api_tts_install(payload: Dict[str, Any] = Body(...)) -> Any:
    """Install dependencies for a TTS engine."""
    try:
        engine = str(payload.get("engine") or "").strip()
        app_log(f"TTS install requested: engine='{engine}'")

        if engine not in TTS_ENGINES:
            raise HTTPException(status_code=400, detail=f"Nieznany silnik TTS: {engine}")

        worker = ROOT / "backend" / "tts_worker.py"
        require_existing_file(worker, "Brak pliku tts_worker.py")

        cmd = [
            os.environ.get("PYTHON", sys.executable),
            str(worker),
            "--action", "install",
            "--engine", engine,
        ]

        t = TASKS.start_subprocess(kind=f"tts_install_{engine}", project_id="-", cmd=cmd, cwd=ROOT)
        return {"task_id": t.task_id}

    except HTTPException:
        raise
    except Exception as e:
        app_log(f"TTS install error: {e}")
        raise HTTPException(status_code=500, detail=f"Błąd instalacji TTS: {e}")


@app.post("/api/tts/predownload")
def api_tts_predownload(payload: Dict[str, Any] = Body(...)) -> Any:
    """Download/cache a TTS voice or model."""
    try:
        engine = str(payload.get("engine") or "").strip()
        voice = str(payload.get("voice") or "").strip()
        app_log(f"TTS predownload requested: engine='{engine}', voice='{voice}'")

        if engine not in TTS_ENGINES:
            raise HTTPException(status_code=400, detail=f"Nieznany silnik TTS: {engine}")

        worker = ROOT / "backend" / "tts_worker.py"
        require_existing_file(worker, "Brak pliku tts_worker.py")

        cmd = [
            os.environ.get("PYTHON", sys.executable),
            str(worker),
            "--action", "predownload",
            "--engine", engine,
        ]
        if voice:
            cmd += ["--voice", voice]

        t = TASKS.start_subprocess(kind=f"tts_predownload_{engine}", project_id="-", cmd=cmd, cwd=ROOT)
        return {"task_id": t.task_id}

    except HTTPException:
        raise
    except Exception as e:
        app_log(f"TTS predownload error: {e}")
        raise HTTPException(status_code=500, detail=f"Błąd pobierania modelu TTS: {e}")


@app.post("/api/tts/synthesize")
def api_tts_synthesize(payload: Dict[str, Any] = Body(...)) -> Any:
    """Synthesize speech from text. Returns task_id; audio file served via /api/tts/audio/."""
    try:
        engine = str(payload.get("engine") or "piper").strip()
        text = str(payload.get("text") or "").strip()
        voice = str(payload.get("voice") or "").strip()
        lang = str(payload.get("lang") or "").strip()
        project_id = str(payload.get("project_id") or "-").strip()

        if not text:
            raise HTTPException(status_code=400, detail="Brak tekstu do syntezowania.")

        if engine not in TTS_ENGINES:
            raise HTTPException(status_code=400, detail=f"Nieznany silnik TTS: {engine}")

        worker = ROOT / "backend" / "tts_worker.py"
        require_existing_file(worker, "Brak pliku tts_worker.py")

        # Generate unique output filename
        import hashlib
        text_hash = hashlib.md5(f"{engine}:{voice}:{text[:200]}".encode()).hexdigest()[:12]
        tts_dir = ROOT / "backend" / "models_cache" / "tts" / "audio_cache"
        tts_dir.mkdir(parents=True, exist_ok=True)
        output_file = tts_dir / f"tts_{text_hash}.wav"

        # If cached audio exists and is non-empty, return immediately
        if output_file.exists() and output_file.stat().st_size > 44:
            return {"status": "cached", "audio_url": f"/api/tts/audio/{output_file.name}"}

        cmd = [
            os.environ.get("PYTHON", sys.executable),
            str(worker),
            "--action", "synthesize",
            "--engine", engine,
            "--voice", voice,
            "--lang", lang,
            "--text", text,
            "--output", str(output_file),
        ]

        if GPU_RM.enabled:
            t = GPU_RM.enqueue_subprocess(kind="tts_synthesize", project_id=project_id, cmd=cmd, cwd=ROOT)
        else:
            t = TASKS.start_subprocess(kind="tts_synthesize", project_id=project_id, cmd=cmd, cwd=ROOT)

        return {"task_id": t.task_id, "audio_url": f"/api/tts/audio/{output_file.name}"}

    except HTTPException:
        raise
    except Exception as e:
        app_log(f"TTS synthesize error: {e}")
        raise HTTPException(status_code=500, detail=f"Błąd syntezy mowy: {e}")


@app.get("/api/tts/audio/{filename}")
def api_tts_audio(filename: str) -> Any:
    """Serve generated TTS audio file."""
    from fastapi.responses import FileResponse

    # Sanitize filename
    safe = Path(filename).name
    audio_dir = ROOT / "backend" / "models_cache" / "tts" / "audio_cache"
    audio_path = audio_dir / safe

    if not audio_path.exists():
        raise HTTPException(status_code=404, detail="Audio file not found")

    return FileResponse(str(audio_path), media_type="audio/wav", filename=safe)


@app.get("/api/tts/voices")
def api_tts_voices() -> Any:
    """Return language-to-voice mapping for all TTS engines."""
    # Import from worker to keep single source of truth
    try:
        import importlib.util
        spec = importlib.util.spec_from_file_location(
            "tts_worker", str(ROOT / "backend" / "tts_worker.py")
        )
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return getattr(mod, "LANG_VOICE_MAP", {})
    except Exception:
        return {}


# ---------- API: NLLB translation models ----------


@app.get("/api/nllb/status")
def api_nllb_status() -> Any:
    """Return dependency status for NLLB (Transformers)."""
    return {
        "torch": _pkg_info("torch"),
        "transformers": _pkg_info("transformers"),
        "sentencepiece": _pkg_info("sentencepiece"),
        "sacremoses": _pkg_info("sacremoses"),
    }


@app.get("/api/nllb/models_state")
def api_nllb_models_state(refresh: int = 0) -> Any:
    """Return cached model-state map for NLLB presets."""
    if refresh:
        return scan_and_persist_nllb_model_cache()
    reg = _read_nllb_registry()
    # If registry empty (first run), auto-scan.
    if not reg:
        try:
            reg = scan_and_persist_nllb_model_cache()
        except Exception:
            reg = {}
    return reg


@app.post("/api/nllb/install_deps")
def api_nllb_install_deps() -> Any:
    worker = ROOT / "backend" / "nllb_worker.py"
    require_existing_file(worker, "Brak nllb_worker.py")

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--action", "install_deps",
    ]
    app_log("NLLB deps install requested")
    t = TASKS.start_subprocess(kind="nllb_install_deps", project_id="-", cmd=cmd, cwd=ROOT)
    return {"task_id": t.task_id}


@app.post("/api/nllb/predownload")
def api_nllb_predownload(payload: Dict[str, Any] = Body(...)) -> Any:
    mode = str(payload.get("mode") or "fast").strip().lower()
    if mode not in ("fast", "accurate"):
        raise HTTPException(status_code=400, detail="Invalid mode")

    model = str(payload.get("model") or "").strip()
    if not model:
        raise HTTPException(status_code=400, detail="Missing model")

    allowed = set(list(NLLB_FAST_MODELS) + list(NLLB_ACCURATE_MODELS))
    if model not in allowed:
        raise HTTPException(status_code=400, detail="Unknown model")

    worker = ROOT / "backend" / "nllb_worker.py"
    require_existing_file(worker, "Brak nllb_worker.py")

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--action", "predownload",
        "--mode", mode,
        "--model", model,
    ]
    app_log(f"NLLB predownload requested: mode={mode}, model={model}")
    t = TASKS.start_subprocess(kind=f"nllb_predownload_{mode}", project_id="-", cmd=cmd, cwd=ROOT)
    return {"task_id": t.task_id}


# ---------- API: global model settings (Ollama LLM) ----------

GLOBAL_SETTINGS_REL = Path("_global") / "settings.json"


def _global_settings_path() -> Path:
    # Stored under DATA_DIR/projects/_global/settings.json
    g = (PROJECTS_DIR / GLOBAL_SETTINGS_REL).resolve()
    g.parent.mkdir(parents=True, exist_ok=True)
    return g


def _read_global_settings() -> Dict[str, Any]:
    fp = _global_settings_path()
    if not fp.exists():
        return {}
    try:
        obj = json.loads(fp.read_text(encoding="utf-8"))
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def _write_global_settings(obj: Dict[str, Any]) -> None:
    fp = _global_settings_path()
    fp.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def _get_analysis_settings() -> Dict[str, Any]:
    """Return analysis-related flags from global settings."""
    obj = _read_global_settings()
    raw = obj.get("analysis") if isinstance(obj.get("analysis"), dict) else {}
    quick_enabled = raw.get("quick_enabled")
    if quick_enabled is None:
        quick_enabled = True
    return {"quick_enabled": bool(quick_enabled)}


def _save_analysis_settings(patch: Dict[str, Any]) -> Dict[str, Any]:
    """Persist analysis-related flags into global settings."""
    obj = _read_global_settings()
    raw = obj.get("analysis") if isinstance(obj.get("analysis"), dict) else {}
    raw = dict(raw or {})

    if "quick_enabled" in patch:
        raw["quick_enabled"] = bool(patch.get("quick_enabled"))

    obj["analysis"] = raw
    _write_global_settings(obj)
    return _get_analysis_settings()



def _get_gpu_rm_settings() -> Dict[str, Any]:
    """Load persisted GPU Resource Manager settings from global settings file."""
    obj = _read_global_settings()
    raw = obj.get("gpu_rm") if isinstance(obj, dict) else {}
    if not isinstance(raw, dict):
        raw = {}

    # Priorities are configured per *feature area* (admin-facing), not per internal kind.
    # Higher value means earlier dispatch.
    # Default order: transcription > diarization > translation > analysis_quick > analysis > chat
    default_prio = {
        "transcription": 300,
        "diarization": 200,
        "translation": 180,
        "analysis_quick": 140,
        "analysis": 120,
        "chat": 60,
    }

    # Backward compatible mapping from older per-kind keys.
    pr = raw.get("priorities")
    merged_prio = dict(default_prio)
    if isinstance(pr, dict):
        def _as_int(x, dv):
            try:
                return int(x)
            except Exception:
                return dv

                # New keys (preferred)
        if "transcription" in pr:
            merged_prio["transcription"] = _as_int(pr.get("transcription"), merged_prio["transcription"])
        if "diarization" in pr:
            merged_prio["diarization"] = _as_int(pr.get("diarization"), merged_prio["diarization"])
        if "translation" in pr:
            merged_prio["translation"] = _as_int(pr.get("translation"), merged_prio["translation"])
        if "analysis_quick" in pr:
            merged_prio["analysis_quick"] = _as_int(pr.get("analysis_quick"), merged_prio["analysis_quick"])
        if "analysis" in pr:
            merged_prio["analysis"] = _as_int(pr.get("analysis"), merged_prio["analysis"])
        if "chat" in pr:
            merged_prio["chat"] = _as_int(pr.get("chat"), merged_prio["chat"])

        # Old keys (migrate best-effort)
        if "transcribe" in pr and "transcription" not in pr:
            merged_prio["transcription"] = _as_int(pr.get("transcribe"), merged_prio["transcription"])
        if "diarize_voice" in pr and "diarization" not in pr:
            merged_prio["diarization"] = _as_int(pr.get("diarize_voice"), merged_prio["diarization"])
        if ("translate_fast" in pr or "translate_accurate" in pr) and "translation" not in pr:
            merged_prio["translation"] = _as_int(pr.get("translate_fast", pr.get("translate_accurate")), merged_prio["translation"])

    return {
        "enabled": bool(raw.get("enabled", True)),
        "gpu_mem_fraction": float(raw.get("gpu_mem_fraction", 0.85)),
        "gpu_slots_per_gpu": int(raw.get("gpu_slots_per_gpu", 1)),
        "cpu_slots": int(raw.get("cpu_slots", 1)),
        "priorities": merged_prio,
    }


def _save_gpu_rm_settings(cfg: Dict[str, Any]) -> None:
    obj = _read_global_settings()
    if not isinstance(obj, dict):
        obj = {}

    pr = cfg.get("priorities")
    pr_out = None
    if isinstance(pr, dict):
        # Persist only admin-facing keys.
        allow = {"transcription", "diarization", "translation", "analysis_quick", "analysis", "chat"}
        pr_out = {}
        for k in allow:
            if k not in pr:
                continue
            try:
                pr_out[k] = int(pr.get(k))
            except Exception:
                continue

    payload = {
        "enabled": bool(cfg.get("enabled", True)),
        "gpu_mem_fraction": float(cfg.get("gpu_mem_fraction", 0.85)),
        "gpu_slots_per_gpu": int(cfg.get("gpu_slots_per_gpu", 1)),
        "cpu_slots": int(cfg.get("cpu_slots", 1)),
    }
    if pr_out is not None:
        payload["priorities"] = pr_out

    obj["gpu_rm"] = payload
    _write_global_settings(obj)


# Global singleton (used by transcription/diarization + Admin panel)
GPU_RM = GPUResourceManager()
GPU_RM.apply_config(_get_gpu_rm_settings())

# Inject GPU_RM into chat router (deferred — GPU_RM created after initial mount)
chat_router.init(ollama_client=OLLAMA, app_log_fn=app_log, gpu_rm=GPU_RM, tasks=TASKS)

# Mount admin router now that GPU_RM is ready
admin_router.init(
    gpu_rm=GPU_RM,
    get_gpu_rm_settings=_get_gpu_rm_settings,
    save_gpu_rm_settings=_save_gpu_rm_settings,
    app_log_fn=app_log,
)
app.include_router(admin_router.router)

def _read_custom_models() -> Dict[str, List[str]]:
    """Read user-added custom model ids grouped by category from global settings."""
    obj = _read_global_settings()
    raw = obj.get("custom_models") if isinstance(obj.get("custom_models"), dict) else {}
    out: Dict[str, List[str]] = {}
    for g in (DEFAULT_MODELS or {}).keys():
        arr = raw.get(g)
        if isinstance(arr, list):
            cleaned = []
            for x in arr:
                s = str(x or "").strip()
                if s and s not in cleaned:
                    cleaned.append(s)
            out[g] = cleaned
        else:
            out[g] = []
    return out


def _write_custom_models(custom_models: Dict[str, List[str]]) -> None:
    """Persist custom model ids into global settings."""
    obj = _read_global_settings()
    obj["custom_models"] = custom_models
    _write_global_settings(obj)


def _get_model_settings() -> Dict[str, str]:
    obj = _read_global_settings()
    models = obj.get("models") if isinstance(obj.get("models"), dict) else {}
    out: Dict[str, str] = {}
    # Keep backward compatibility: always include quick/deep and ignore unknown keys.
    for k, dv in (DEFAULT_MODELS or {}).items():
        v = (models or {}).get(k)
        v = str(v).strip() if v is not None else ""
        out[k] = v or dv
    # If older global settings contained only quick/deep, above will fill the rest.
    return out





# Admin GPU endpoints moved to webapp/routers/admin.py




@app.get("/api/settings/analysis")
async def api_get_analysis_settings() -> Any:
    """Return analysis settings (feature flags)."""
    return _get_analysis_settings()


@app.post("/api/settings/analysis")
async def api_save_analysis_settings(payload: Dict[str, Any] = Body(...)) -> Any:
    """Persist analysis settings (feature flags)."""
    saved = _save_analysis_settings(payload or {})
    app_log("Analysis settings saved: " + ", ".join([f"{k}={saved.get(k)}" for k in saved.keys()]))
    return {"status": "saved", **saved}

@app.get("/api/settings/models")
async def api_get_model_settings() -> Any:
    """Return saved global model selection.

    Returns keys for all supported groups:
      quick, deep, vision, translation, financial, specialized
    """
    models = _get_model_settings()
    # Keep a predictable place for defaults while preserving top-level keys.
    return {**models, "defaults": DEFAULT_MODELS}


@app.post("/api/settings/models")
async def api_save_model_settings(payload: Dict[str, Any] = Body(...)) -> Any:
    """Persist global model selection.

    Request can contain any subset of groups, e.g.:
      {"quick":"...","deep":"..."}
      {"vision":"..."}
      {"quick":"...","deep":"...","translation":"...",...}
    """
    obj = _read_global_settings()
    models = obj.get("models") if isinstance(obj.get("models"), dict) else {}
    models = dict(models or {})

    # Apply incoming changes (partial updates supported)
    for k in (DEFAULT_MODELS or {}).keys():
        if k not in payload:
            continue
        raw = payload.get(k)
        val = "" if raw is None else str(raw).strip()
        if not val:
            # Treat empty value as "use default" (remove persisted override)
            models.pop(k, None)
        else:
            models[k] = val

    # Normalize: ensure every known key resolves to a non-empty model id
    resolved = {}
    for k, dv in (DEFAULT_MODELS or {}).items():
        v = str(models.get(k) or "").strip() or dv
        resolved[k] = v
        models[k] = v

    obj["models"] = models
    _write_global_settings(obj)
    app_log("Model settings saved: " + ", ".join([f"{k}={resolved.get(k)}" for k in resolved.keys()]))
    return {"status": "saved", "models": resolved, **resolved, "defaults": DEFAULT_MODELS}


@app.get("/api/models/info/{model_name}")
async def api_model_info(model_name: str) -> Any:
    """Return info for a specific model (catalog or custom)."""
    name = str(model_name or "").strip()
    info = MODELS_INFO.get(name)

    out: Dict[str, Any] = {"id": name}
    if info:
        out.update(info)
    else:
        out.update({
            "display_name": name,
            "category": "custom",
            "hardware": {},
            "performance": {},
            "use_cases": ["Custom model (added by user) – not in built-in catalog"],
            "recommendation": (
                "Model dodany ręcznie. Aplikacja nie ma wbudowanych metadanych (VRAM/szybkość/use-cases). \nJeśli chcesz widzieć szczegóły w panelu info, dopisz go do backend/models_info.py."
            ),
            "warning": "Custom model – brak metadanych w katalogu",
            "capabilities": ["custom"],
        })

    # best-effort installed flag
    try:
        installed = await OLLAMA.list_model_names()
    except Exception:
        installed = []
    out["installed"] = name in installed
    return out
@app.get("/api/models/custom")
async def api_models_custom() -> Any:
    """Return user-added custom models grouped by category."""
    return {"custom_models": _read_custom_models()}


@app.post("/api/models/custom")
async def api_models_custom_add(payload: Dict[str, Any] = Body(...)) -> Any:
    """Add a custom model id to a selected group."""
    group = str(payload.get("group") or "deep").strip()
    model_id = str(payload.get("model_id") or payload.get("model") or "").strip()

    if not model_id:
        raise HTTPException(status_code=400, detail="Missing model_id")
    if group not in (DEFAULT_MODELS or {}):
        raise HTTPException(status_code=400, detail="Unknown group")

    custom = _read_custom_models()
    arr = list(custom.get(group) or [])
    if model_id not in arr:
        arr.append(model_id)
    custom[group] = arr
    _write_custom_models(custom)

    return {"status": "ok", "custom_models": custom}


@app.post("/api/models/custom/remove")
async def api_models_custom_remove(payload: Dict[str, Any] = Body(...)) -> Any:
    """Remove a custom model id from a group."""
    group = str(payload.get("group") or "deep").strip()
    model_id = str(payload.get("model_id") or payload.get("model") or "").strip()

    if not model_id:
        raise HTTPException(status_code=400, detail="Missing model_id")
    if group not in (DEFAULT_MODELS or {}):
        raise HTTPException(status_code=400, detail="Unknown group")

    custom = _read_custom_models()
    arr = [x for x in (custom.get(group) or []) if str(x) != model_id]
    custom[group] = arr
    _write_custom_models(custom)

    return {"status": "ok", "custom_models": custom}



@app.get("/api/models/list")
async def api_models_list() -> Any:
    """Return recommended models grouped by category with install status.

    Includes:
    - built-in catalog models (MODELS_INFO / MODELS_GROUPS)
    - user-added custom models (stored in global settings -> custom_models)
    """
    st = await OLLAMA.status()
    installed: List[str] = []
    if st.status == "online":
        installed = st.models or []

    custom = _read_custom_models()

    def _entry(mid: str, category: str) -> Dict[str, Any]:
        info = MODELS_INFO.get(mid) or {}
        perf = info.get("performance") or {}
        at = None
        try:
            at = (perf.get("analysis_time") or {}).get(category)
        except Exception:
            at = None

        warning = str(info.get("warning") or "")
        if not info:
            warning = warning or "Custom model – brak metadanych w katalogu"

        return {
            "id": mid,
            "display_name": str(info.get("display_name") or mid),
            "installed": mid in installed,
            "default": bool((info.get("defaults") or {}).get(category)) if info else False,
            "vram": str(((info.get("hardware") or {}).get("vram")) or ""),
            "speed": str(at or ""),
            "requires_multi_gpu": bool(info.get("requires_multi_gpu") or False) if info else False,
            "warning": warning,
        }

    out: Dict[str, Any] = {}
    for category, ids in (MODELS_GROUPS or {}).items():
        all_ids: List[str] = []
        for mid in (ids or []):
            mid = str(mid or "").strip()
            if mid and mid not in all_ids:
                all_ids.append(mid)
        for mid in (custom.get(category) or []):
            mid = str(mid or "").strip()
            if mid and mid not in all_ids:
                all_ids.append(mid)
        out[category] = [_entry(mid, category) for mid in all_ids]

    # Preserve legacy keys (quick/deep) even if groups definition changes
    out.setdefault("quick", [])
    out.setdefault("deep", [])

    out.update({
        "ollama_status": st.status,
        "ollama_version": st.version,
        "ollama_url": getattr(OLLAMA, "base_url", "http://127.0.0.1:11434"),
        "models_count": len(installed),
        "custom_models": custom,
    })
    return out

# ---------- API: projects ----------

@app.post("/api/projects/create")
async def api_create_project(
    request: Request,
    name: str = Form(...),
    audio: Optional[UploadFile] = File(None),
) -> Any:
    # Create a new project with a user-provided name.
    # Audio is optional: if provided, it becomes the project's source audio.
    pid = ensure_project(None)
    pname = str(name or "").strip() or default_project_name()
    up_name = ""
    try:
        up_name = (audio.filename or "") if audio else ""
    except Exception:
        up_name = ""
    app_log(f"Project create: project_id={pid}, name='{pname}', upload='{up_name}'")
    meta = read_project_meta(pid)
    meta["name"] = pname
    meta["created_at"] = meta.get("created_at") or now_iso()
    meta["updated_at"] = now_iso()
    # Set owner in multi-user mode
    user = getattr(request.state, "user", None)
    if user and getattr(request.state, "multiuser", False):
        meta["owner_id"] = user.user_id
    write_project_meta(pid, meta)

    audio_path: Optional[Path] = None
    if audio and getattr(audio, "filename", None):
        audio_path = save_upload(pid, audio)
        size_b = 0
        try:
            size_b = audio_path.stat().st_size
        except Exception:
            size_b = 0
        app_log(f"Project upload saved: project_id={pid}, file='{audio_path.name}', size_bytes={size_b}")

    meta = read_project_meta(pid)
    meta["updated_at"] = now_iso()
    write_project_meta(pid, meta)

    return {
        "project_id": pid,
        "name": meta.get("name"),
        "audio_file": meta.get("audio_file") or "",
        "audio_path": str(audio_path.name) if audio_path else "",
    }

@app.post("/api/projects/{project_id}/upload_audio")
async def api_upload_audio(project_id: str, audio: UploadFile = File(...)) -> Any:
    """Upload an audio file to an existing project."""
    meta = read_project_meta(project_id)
    if not meta:
        raise HTTPException(status_code=404, detail="Project not found")
    audio_path = save_upload(project_id, audio)
    # Re-read meta since save_upload updates audio_file
    meta = read_project_meta(project_id)
    app_log(f"Audio upload: project_id={project_id}, file='{audio_path.name}'")
    return {
        "ok": True,
        "audio_file": meta.get("audio_file", ""),
        "project_id": project_id,
    }


@app.post("/api/projects/new")
def api_new_project(request: Request) -> Any:
    pid = ensure_project(None)
    # Set owner in multi-user mode
    user = getattr(request.state, "user", None)
    if user and getattr(request.state, "multiuser", False):
        meta = read_project_meta(pid)
        meta["owner_id"] = user.user_id
        write_project_meta(pid, meta)
    return {"project_id": pid}


@app.get("/api/admin/user-projects")
def api_admin_user_projects(request: Request) -> Any:
    """Return all projects and workspaces grouped by user. Superadmin only."""
    user = getattr(request.state, "user", None)
    if not user or not getattr(user, "is_superadmin", False):
        raise HTTPException(status_code=403, detail="Superadmin required")

    # Build username lookup
    all_users = USER_STORE.list_users()
    uid_map: Dict[str, Dict[str, Any]] = {}
    for u in all_users:
        uid_map[u.user_id] = {
            "user_id": u.user_id,
            "username": u.username,
            "display_name": u.display_name or u.username,
            "role": u.role or "",
            "is_admin": u.is_admin,
            "is_superadmin": getattr(u, "is_superadmin", False),
        }

    # 1) File-based projects from disk
    file_projects: Dict[str, List[Dict[str, Any]]] = {}  # owner_id -> [project...]
    orphan_projects: List[Dict[str, Any]] = []
    for p in PROJECTS_DIR.glob("*"):
        if not p.is_dir() or p.name.startswith("_"):
            continue
        pid = p.name
        meta = read_project_meta(pid)
        # Calculate directory size
        dir_size = 0
        try:
            for f in p.rglob("*"):
                if f.is_file():
                    dir_size += f.stat().st_size
        except Exception:
            pass
        audio_file = meta.get("audio_file", "")
        audio_size = 0
        if audio_file:
            audio_path = p / audio_file
            try:
                audio_size = audio_path.stat().st_size if audio_path.exists() else 0
            except Exception:
                pass
        entry = {
            "project_id": pid,
            "name": meta.get("name", ""),
            "created_at": meta.get("created_at", ""),
            "updated_at": meta.get("updated_at", ""),
            "audio_file": audio_file,
            "audio_size": audio_size,
            "has_transcript": bool(meta.get("has_transcript")),
            "has_diarized": bool(meta.get("has_diarized")),
            "dir_path": str(p),
            "dir_size": dir_size,
            "shares": meta.get("shares", []),
            "owner_id": meta.get("owner_id", ""),
        }
        owner = meta.get("owner_id", "")
        if owner:
            file_projects.setdefault(owner, []).append(entry)
        else:
            orphan_projects.append(entry)

    # 2) Workspaces from database
    user_workspaces: Dict[str, List[Dict[str, Any]]] = {}  # owner_id -> [ws...]
    try:
        from backend.db.engine import get_conn
        with get_conn() as conn:
            ws_rows = conn.execute(
                "SELECT * FROM project_workspaces WHERE status != 'deleted' ORDER BY updated_at DESC"
            ).fetchall()
            for r in ws_rows:
                ws = dict(r)
                ws_id = ws["id"]
                owner_id = ws.get("owner_id", "")
                # Members
                members = []
                mem_rows = conn.execute(
                    "SELECT user_id, role, status FROM project_members WHERE workspace_id = ? AND status='accepted'",
                    (ws_id,),
                ).fetchall()
                for mr in mem_rows:
                    m = dict(mr)
                    u_info = uid_map.get(m["user_id"])
                    m["username"] = u_info["username"] if u_info else m["user_id"][:8]
                    m["display_name"] = u_info["display_name"] if u_info else ""
                    members.append(m)
                # Subprojects
                subprojects = []
                sp_rows = conn.execute(
                    "SELECT * FROM subprojects WHERE workspace_id = ? ORDER BY created_at",
                    (ws_id,),
                ).fetchall()
                for sp in sp_rows:
                    sp_dict = dict(sp)
                    # Calculate subproject dir size
                    data_dir = sp_dict.get("data_dir", "")
                    sp_dir_size = 0
                    sp_dir_path = ""
                    if data_dir:
                        sp_path = (PROJECTS_DIR.parent / data_dir).resolve()
                        sp_dir_path = str(sp_path)
                        try:
                            for f in sp_path.rglob("*"):
                                if f.is_file():
                                    sp_dir_size += f.stat().st_size
                        except Exception:
                            pass
                    sp_dict["dir_path"] = sp_dir_path
                    sp_dict["dir_size"] = sp_dir_size
                    subprojects.append(sp_dict)
                ws_entry = {
                    "id": ws_id,
                    "name": ws.get("name", ""),
                    "description": ws.get("description", ""),
                    "color": ws.get("color", "#4a6cf7"),
                    "status": ws.get("status", "active"),
                    "created_at": ws.get("created_at", ""),
                    "updated_at": ws.get("updated_at", ""),
                    "owner_id": owner_id,
                    "members": members,
                    "subprojects": subprojects,
                }
                user_workspaces.setdefault(owner_id, []).append(ws_entry)
    except Exception as e:
        app_log(f"admin user-projects: workspace query failed: {e}")

    # 3) Assemble per-user result
    all_user_ids = set(file_projects.keys()) | set(user_workspaces.keys())
    users_result = []
    for uid in all_user_ids:
        u_info = uid_map.get(uid, {
            "user_id": uid, "username": uid[:8] + "...", "display_name": "",
            "role": "", "is_admin": False, "is_superadmin": False,
        })
        fp = file_projects.get(uid, [])
        ws = user_workspaces.get(uid, [])
        total_size = sum(p["dir_size"] for p in fp)
        total_size += sum(
            sp["dir_size"]
            for w in ws for sp in w.get("subprojects", [])
        )
        users_result.append({
            "user": u_info,
            "file_projects": fp,
            "workspaces": ws,
            "total_size": total_size,
            "project_count": len(fp),
            "workspace_count": len(ws),
        })
    users_result.sort(key=lambda x: x["total_size"], reverse=True)

    return {
        "status": "ok",
        "users": users_result,
        "orphan_projects": orphan_projects,
    }


@app.post("/api/admin/delete-workspace/{workspace_id}")
async def api_admin_delete_workspace(workspace_id: str, request: Request) -> Any:
    """Delete a workspace and all its subproject data. Superadmin only."""
    user = getattr(request.state, "user", None)
    if not user or not getattr(user, "is_superadmin", False):
        raise HTTPException(status_code=403, detail="Superadmin required")

    body = await request.json()
    wipe_method = (body.get("wipe_method") or "none").strip().lower()

    ws = WORKSPACE_STORE.get_workspace(workspace_id)
    if ws is None:
        raise HTTPException(status_code=404, detail="Workspace not found")

    # Delete subproject data directories
    subs = WORKSPACE_STORE.list_subprojects(workspace_id)
    deleted_dirs = []
    for sp in subs:
        data_dir = sp.get("data_dir", "")
        if data_dir:
            sp_path = (PROJECTS_DIR.parent / data_dir).resolve()
            if sp_path.exists() and sp_path.is_dir() and PROJECTS_DIR.resolve() in sp_path.parents:
                try:
                    secure_delete_project_dir(sp_path, wipe_method)
                    deleted_dirs.append(str(sp_path))
                except Exception as e:
                    app_log(f"admin delete-workspace: wipe failed for {sp_path}: {e}")

    # Hard-delete the workspace + all subprojects/members/invitations/activity from DB
    WORKSPACE_STORE.hard_delete_workspace(workspace_id)

    if AUDIT_STORE:
        AUDIT_STORE.log_event(
            "workspace_deleted_admin",
            user_id=user.user_id, username=user.username,
            detail=f"workspace={workspace_id} name={ws.get('name','')} wipe={wipe_method}",
            actor_id=user.user_id, actor_name=user.username,
        )
    app_log(f"Admin '{user.username}' deleted workspace '{ws.get('name','')}' (id={workspace_id}, wipe={wipe_method})")

    return {"status": "ok", "deleted_dirs": deleted_dirs}


@app.post("/api/admin/delete-subproject/{subproject_id}")
async def api_admin_delete_subproject(subproject_id: str, request: Request) -> Any:
    """Delete a single subproject and its data. Superadmin only."""
    user = getattr(request.state, "user", None)
    if not user or not getattr(user, "is_superadmin", False):
        raise HTTPException(status_code=403, detail="Superadmin required")

    body = await request.json()
    wipe_method = (body.get("wipe_method") or "none").strip().lower()

    sp = WORKSPACE_STORE.get_subproject(subproject_id)
    if sp is None:
        raise HTTPException(status_code=404, detail="Subproject not found")

    # Delete data directory
    data_dir = sp.get("data_dir", "")
    if data_dir:
        sp_path = (PROJECTS_DIR.parent / data_dir).resolve()
        if sp_path.exists() and sp_path.is_dir() and PROJECTS_DIR.resolve() in sp_path.parents:
            try:
                secure_delete_project_dir(sp_path, wipe_method)
            except Exception as e:
                app_log(f"admin delete-subproject: wipe failed for {sp_path}: {e}")

    # Hard-delete from DB
    WORKSPACE_STORE.delete_subproject(subproject_id)

    if AUDIT_STORE:
        AUDIT_STORE.log_event(
            "subproject_deleted_admin",
            user_id=user.user_id, username=user.username,
            detail=f"subproject={subproject_id} name={sp.get('name','')} wipe={wipe_method}",
            actor_id=user.user_id, actor_name=user.username,
        )
    app_log(f"Admin '{user.username}' deleted subproject '{sp.get('name','')}' (id={subproject_id}, wipe={wipe_method})")

    return {"status": "ok"}


@app.post("/api/admin/delete-file-project/{project_id}")
async def api_admin_delete_file_project(project_id: str, request: Request) -> Any:
    """Delete a file-based (orphan/legacy) project. Superadmin only."""
    user = getattr(request.state, "user", None)
    if not user or not getattr(user, "is_superadmin", False):
        raise HTTPException(status_code=403, detail="Superadmin required")

    body = await request.json()
    wipe_method = (body.get("wipe_method") or "none").strip().lower()

    pdir = (PROJECTS_DIR / project_id).resolve()
    root = PROJECTS_DIR.resolve()
    if root not in pdir.parents or not pdir.exists() or not pdir.is_dir():
        raise HTTPException(status_code=404, detail="Project not found")

    try:
        secure_delete_project_dir(pdir, wipe_method)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Delete error: {e}")

    # Clean DB references: subprojects pointing to this data_dir + projects table
    try:
        from backend.db.engine import get_conn
        data_dir_rel = f"projects/{project_id}"
        with get_conn() as conn:
            conn.execute("DELETE FROM subprojects WHERE data_dir = ?", (data_dir_rel,))
            conn.execute("DELETE FROM projects WHERE id = ?", (project_id,))
    except Exception:
        pass  # DB tables may not exist for pure file-based projects

    if AUDIT_STORE:
        AUDIT_STORE.log_event(
            "project_deleted_admin",
            user_id=user.user_id, username=user.username,
            detail=f"file_project={project_id} wipe={wipe_method}",
            actor_id=user.user_id, actor_name=user.username,
        )
    app_log(f"Admin '{user.username}' deleted file-project '{project_id}' (wipe={wipe_method})")

    return {"status": "ok"}


@app.get("/api/projects")
def api_list_projects(request: Request) -> Any:
    multiuser = getattr(request.state, "multiuser", False)
    user = getattr(request.state, "user", None)

    projects: List[Dict[str, Any]] = []
    for p in PROJECTS_DIR.glob("*"):
        if p.is_dir():
            pid = p.name
            meta = read_project_meta(pid)

            # In multi-user mode, filter by ownership/sharing
            if multiuser and user and not user.is_admin:
                owner = meta.get("owner_id", "")
                shares = meta.get("shares") or []
                shared_user_ids = [s.get("user_id") for s in shares if isinstance(s, dict)]
                if owner != user.user_id and user.user_id not in shared_user_ids:
                    continue

            entry = {
                "project_id": pid,
                "created_at": meta.get("created_at"),
                "name": meta.get("name"),
                "updated_at": meta.get("updated_at"),
            }
            if multiuser:
                entry["owner_id"] = meta.get("owner_id", "")
                entry["shares"] = meta.get("shares", [])
            projects.append(entry)
    projects.sort(key=lambda x: x.get("updated_at") or x.get("created_at") or "", reverse=True)
    return {"projects": projects}


# --- Project sharing endpoints ---

@app.post("/api/projects/{project_id}/share")
async def api_share_project(project_id: str, request: Request) -> Any:
    """Share a project with another user. Only owner or admin can share."""
    user = getattr(request.state, "user", None)
    meta = read_project_meta(project_id)
    if not meta:
        raise HTTPException(status_code=404, detail="Project not found")

    multiuser = getattr(request.state, "multiuser", False)
    if not multiuser:
        raise HTTPException(status_code=400, detail="Sharing only available in multi-user mode")

    # Check ownership
    if user and not user.is_admin and meta.get("owner_id") != user.user_id:
        raise HTTPException(status_code=403, detail="Only owner or admin can share")

    try:
        body = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid request")

    target_username = (body.get("username") or "").strip()
    permission = body.get("permission", "read")  # "read" or "edit"
    if permission not in ("read", "edit"):
        raise HTTPException(status_code=400, detail="Permission must be 'read' or 'edit'")

    target_user = USER_STORE.get_by_username(target_username)
    if target_user is None:
        raise HTTPException(status_code=404, detail="User not found")

    shares = meta.get("shares") or []
    # Remove existing share for this user if any
    shares = [s for s in shares if s.get("user_id") != target_user.user_id]
    shares.append({
        "user_id": target_user.user_id,
        "username": target_user.username,
        "permission": permission,
        "granted_at": now_iso(),
        "granted_by": user.user_id if user else "system",
    })
    meta["shares"] = shares
    write_project_meta(project_id, meta)
    return {"status": "ok", "shares": shares}


@app.post("/api/projects/{project_id}/unshare")
async def api_unshare_project(project_id: str, request: Request) -> Any:
    """Remove sharing for a user."""
    user = getattr(request.state, "user", None)
    meta = read_project_meta(project_id)
    if not meta:
        raise HTTPException(status_code=404, detail="Project not found")

    multiuser = getattr(request.state, "multiuser", False)
    if not multiuser:
        raise HTTPException(status_code=400, detail="Sharing only available in multi-user mode")

    if user and not user.is_admin and meta.get("owner_id") != user.user_id:
        raise HTTPException(status_code=403, detail="Only owner or admin can unshare")

    try:
        body = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid request")

    target_user_id = body.get("user_id", "").strip()
    shares = meta.get("shares") or []
    meta["shares"] = [s for s in shares if s.get("user_id") != target_user_id]
    write_project_meta(project_id, meta)
    return {"status": "ok", "shares": meta["shares"]}


@app.get("/api/projects/{project_id}/meta")
def api_project_meta(project_id: str) -> Any:
    # Ensure the project exists and return its current metadata
    project_path(project_id)
    meta = read_project_meta(project_id)
    # Return only JSON-serializable basic fields
    return {
        "project_id": project_id,
        "created_at": meta.get("created_at"),
        "updated_at": meta.get("updated_at"),
        "name": meta.get("name"),
        "audio_file": meta.get("audio_file") or "",
        "has_transcript": bool(meta.get("has_transcript")),
        "has_diarized": bool(meta.get("has_diarized")),
    }


@app.get("/api/projects/{project_id}/waveform")
def api_waveform(project_id: str) -> Any:
    """Return cached waveform peaks.  Generate on first request if missing."""
    pdir = project_path(project_id)
    peaks_path = pdir / "peaks.json"
    if not peaks_path.exists():
        ok = _generate_waveform_peaks(project_id)
        if not ok:
            raise HTTPException(status_code=404, detail="Nie można wygenerować fali dźwiękowej")
    try:
        data = json.loads(peaks_path.read_text(encoding="utf-8"))
        return JSONResponse(data)
    except Exception:
        raise HTTPException(status_code=500, detail="Błąd odczytu peaks.json")


@app.get("/api/projects/{project_id}/download/{filename}")
def api_download(project_id: str, filename: str) -> Any:
    path = _safe_child_path(project_path(project_id), filename)
    require_existing_file(path, "Plik nie istnieje.")
    return FileResponse(str(path), filename=path.name)


@app.post("/api/projects/{project_id}/save_transcript")
def api_save_transcript(project_id: str, payload: Dict[str, Any]) -> Any:
    text = str(payload.get("text") or "")
    app_log(f"Project save transcript: project_id={project_id}, chars={len(text)}")
    path = project_path(project_id) / "transcript.txt"
    path.write_text(text, encoding="utf-8")
    meta = read_project_meta(project_id)
    meta["has_transcript"] = True
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    return {"ok": True}


@app.post("/api/projects/{project_id}/save_diarized")
def api_save_diarized(project_id: str, payload: Dict[str, Any]) -> Any:
    text = str(payload.get("text") or "")
    app_log(f"Project save diarized: project_id={project_id}, chars={len(text)}")
    path = project_path(project_id) / "diarized.txt"
    path.write_text(text, encoding="utf-8")
    meta = read_project_meta(project_id)
    meta["has_diarized"] = True
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    return {"ok": True}


@app.get("/api/projects/{project_id}/transcript_segments")
def api_get_transcript_segments(project_id: str) -> Any:
    """Get transcription segments with confidence scores."""
    pdir = project_path(project_id)
    segments_file = pdir / "transcript_segments.json"

    if not segments_file.exists():
        return {"segments": []}

    try:
        data = json.loads(segments_file.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return {"segments": data}
        return {"segments": []}
    except Exception as e:
        app_log(f"Error reading transcript segments: {e}")
        return {"segments": []}


@app.get("/api/projects/{project_id}/diarized_segments")
def api_get_diarized_segments(project_id: str) -> Any:
    """Get diarization segments with confidence scores."""
    pdir = project_path(project_id)
    segments_file = pdir / "diarized_segments.json"

    if not segments_file.exists():
        return {"segments": []}

    try:
        data = json.loads(segments_file.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return {"segments": data}
        return {"segments": []}
    except Exception as e:
        app_log(f"Error reading diarized segments: {e}")
        return {"segments": []}


@app.get("/api/projects/{project_id}/translation/draft")
def api_get_translation_draft(project_id: str) -> Any:
    """Load translation draft for a project.

    Returns:
      {"draft": {...}, "detected_lang": "pl"|null} or {"draft": null}
    """
    # Ensure the project exists
    project_path(project_id)
    draft = read_translation_draft(project_id)
    # Include Whisper's detected language (if transcription ran) for auto source-lang
    detected_lang = None
    try:
        meta = read_project_meta(project_id)
        detected_lang = meta.get("detected_lang") or None
    except Exception:
        pass
    return {"draft": draft or None, "detected_lang": detected_lang}


@app.post("/api/projects/{project_id}/translation/draft")
async def api_save_translation_draft(project_id: str, request: Request) -> Any:
    """Save translation draft for a project.

    Expects JSON:
      {"draft": {...}}  (or directly a dict draft)
    """
    payload = await request.json()
    draft = payload.get("draft") if isinstance(payload, dict) else None
    if draft is None:
        draft = payload
    if not isinstance(draft, dict):
        raise HTTPException(status_code=400, detail="draft must be a JSON object (dict).")

    # Ensure the project exists
    project_path(project_id)
    # Add/update timestamp (non-breaking if frontend also sets it)
    if "saved_at" not in draft:
        draft["saved_at"] = int(time.time() * 1000)

    write_translation_draft(project_id, draft)

    meta = read_project_meta(project_id)
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)

    return {"ok": True}



@app.post("/api/projects/{project_id}/speaker_map")
def api_save_speaker_map(project_id: str, payload: Dict[str, Any]) -> Any:
    mapping = payload.get("mapping") or {}
    if not isinstance(mapping, dict):
        raise HTTPException(status_code=400, detail="mapping musi być obiektem JSON (dict).")
    # keep only str->str
    clean: Dict[str, str] = {}
    for k, v in mapping.items():
        if isinstance(k, str) and isinstance(v, str) and k.strip():
            clean[k.strip()] = v.strip()
    meta = read_project_meta(project_id)
    meta["speaker_map"] = clean
    app_log(f"Project save speaker map: project_id={project_id}, entries={len(clean)}")
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    return {"ok": True, "count": len(clean)}


# ---------- API: notes (two-level: global + per-block) ----------

@app.post("/api/projects/{project_id}/notes")
def api_save_notes(project_id: str, payload: Dict[str, Any]) -> Any:
    """Save notes to project (global + per-block).
    
    Structure:
    {
      "global": "Global note for entire conversation...",
      "blocks": {
        "0": "Note for block 0",
        "3": "Note for block 3"
      }
    }
    """
    notes = payload.get("notes") or payload
    
    if not isinstance(notes, dict):
        raise HTTPException(status_code=400, detail="notes must be a JSON object (dict).")
    
    # Validate structure
    global_note = notes.get("global", "")
    blocks = notes.get("blocks", {})
    
    if not isinstance(global_note, str):
        global_note = ""
    
    if not isinstance(blocks, dict):
        blocks = {}
    
    # Clean blocks: only str->str with non-empty values
    clean_blocks: Dict[str, str] = {}
    for k, v in blocks.items():
        if isinstance(k, str) and isinstance(v, str) and v.strip():
            clean_blocks[k.strip()] = v.strip()
    
    clean_notes = {
        "global": global_note.strip(),
        "blocks": clean_blocks
    }
    
    # Save to project.json
    meta = read_project_meta(project_id)
    meta["notes"] = clean_notes
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    
    app_log(f"Project save notes: project_id={project_id}, global_len={len(global_note)}, blocks={len(clean_blocks)}")
    
    return {
        "ok": True,
        "notes": clean_notes,
        "blocks_count": len(clean_blocks)
    }


@app.get("/api/projects/{project_id}/notes")
def api_get_notes(project_id: str) -> Any:
    """Get notes from project.
    
    Returns:
    {
      "global": "...",
      "blocks": {"0": "...", "3": "..."}
    }
    """
    meta = read_project_meta(project_id)
    notes = meta.get("notes", {})
    
    # Ensure proper structure
    if not isinstance(notes, dict):
        notes = {"global": "", "blocks": {}}
    
    notes.setdefault("global", "")
    notes.setdefault("blocks", {})
    
    return notes


@app.get("/api/projects/{project_id}/export.aistate")
@app.get("/api/projects/{project_id}/export.zip")
def api_export_project(project_id: str) -> Any:
    """Export the whole project folder as a portable package.

    NOTE: .aistate is a ZIP container with a custom extension.
    """
    pdir = (PROJECTS_DIR / project_id).resolve()
    root = PROJECTS_DIR.resolve()
    if root not in pdir.parents or not pdir.exists() or not pdir.is_dir():
        raise HTTPException(status_code=404, detail="Nie ma takiego projektu.")

    _skip_suffixes = {".tmp", ".bak", ".pyc"}
    _skip_dirs = {"__pycache__", ".cache"}

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for fp in pdir.rglob("*"):
            if not fp.is_file():
                continue
            if fp.suffix.lower() in _skip_suffixes:
                continue
            if any(part in _skip_dirs for part in fp.relative_to(pdir).parts):
                continue
            z.write(fp, arcname=str(fp.relative_to(pdir)))
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="project-{project_id}.aistate"'},
    )


def _safe_extract_zip(zf: zipfile.ZipFile, dst: Path) -> None:
    """Protect against Zip Slip (path traversal)."""
    dst = dst.resolve()
    for member in zf.infolist():
        name = member.filename
        # Disallow absolute paths
        if name.startswith("/") or name.startswith("\\"):
            raise HTTPException(status_code=400, detail="Nieprawidłowa paczka (ścieżka absolutna).")
        target = (dst / name).resolve()
        if dst != target and dst not in target.parents:
            raise HTTPException(status_code=400, detail="Nieprawidłowa paczka (path traversal).")
    zf.extractall(dst)


def _pick_extracted_root(tmp_dir: Path) -> Path:
    # If archive contains a single top-level folder, use it; otherwise use tmp_dir.
    entries = [p for p in tmp_dir.iterdir() if p.name not in ("__MACOSX", ".DS_Store")]
    if len(entries) == 1 and entries[0].is_dir():
        return entries[0]
    return tmp_dir

@app.post("/api/projects/{project_id}/notes/transcription")
def api_save_notes_transcription(project_id: str, request: Request) -> Any:
    """Save transcription-specific notes to project (global + per-block).
    
    Structure:
    {
      "global": "Global note for entire transcription...",
      "blocks": {
        "0": "Note for block 0",
        "3": "Note for block 3"
      }
    }
    """
    import asyncio
    payload = asyncio.run(request.json())
    notes = payload.get("notes") or payload
    
    if not isinstance(notes, dict):
        raise HTTPException(status_code=400, detail="notes must be a JSON object (dict).")
    
    # Validate structure
    global_note = notes.get("global", "")
    blocks = notes.get("blocks", {})
    
    if not isinstance(global_note, str):
        global_note = ""
    
    if not isinstance(blocks, dict):
        blocks = {}
    
    # Clean blocks: only str->str with non-empty values
    clean_blocks: Dict[str, str] = {}
    for k, v in blocks.items():
        if isinstance(k, str) and isinstance(v, str) and v.strip():
            clean_blocks[k.strip()] = v.strip()
    
    clean_notes = {
        "global": global_note.strip(),
        "blocks": clean_blocks
    }
    
    # Save to project.json under "notes_transcription"
    meta = read_project_meta(project_id)
    meta["notes_transcription"] = clean_notes
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    
    app_log(f"Project save transcription notes: project_id={project_id}, global_len={len(global_note)}, blocks={len(clean_blocks)}")
    
    return {
        "ok": True,
        "notes": clean_notes,
        "blocks_count": len(clean_blocks)
    }


@app.get("/api/projects/{project_id}/notes/transcription")
def api_get_notes_transcription(project_id: str) -> Any:
    """Get transcription notes from project.
    
    Returns:
    {
      "global": "...",
      "blocks": {"0": "...", "3": "..."}
    }
    """
    meta = read_project_meta(project_id)
    notes = meta.get("notes_transcription", {})
    
    # Ensure proper structure
    if not isinstance(notes, dict):
        notes = {"global": "", "blocks": {}}
    
    notes.setdefault("global", "")
    notes.setdefault("blocks", {})
    
    return notes


@app.get("/api/projects/{project_id}/sound_events")
def api_get_sound_events(project_id: str) -> Any:
    """Get detected sound events from project.

    Returns:
    {
      "events": [
        {"start": 12.5, "end": 13.1, "type": "dog", "label": "Dog bark", "confidence": 0.87},
        ...
      ],
      "model": "yamnet",
      "detected_at": "2024-01-15 10:30:00"
    }
    """
    proj_dir = project_path(project_id)
    events_file = proj_dir / "sound_events.json"

    if not events_file.exists():
        return {"events": [], "model": None, "detected_at": None}

    try:
        data = json.loads(events_file.read_text(encoding="utf-8"))
        if isinstance(data, list):
            # Old format: just list of events
            return {"events": data, "model": None, "detected_at": None}
        return data
    except Exception as e:
        app_log(f"Error reading sound events: {e}")
        return {"events": [], "model": None, "detected_at": None}


@app.post("/api/projects/{project_id}/notes/diarization")
def api_save_notes_diarization(project_id: str, request: Request) -> Any:
    """Save diarization-specific notes to project (global + per-block).
    
    Same structure as transcription notes.
    """
    import asyncio
    payload = asyncio.run(request.json())
    notes = payload.get("notes") or payload
    
    if not isinstance(notes, dict):
        raise HTTPException(status_code=400, detail="notes must be a JSON object (dict).")
    
    # Validate structure
    global_note = notes.get("global", "")
    blocks = notes.get("blocks", {})
    
    if not isinstance(global_note, str):
        global_note = ""
    
    if not isinstance(blocks, dict):
        blocks = {}
    
    # Clean blocks: only str->str with non-empty values
    clean_blocks: Dict[str, str] = {}
    for k, v in blocks.items():
        if isinstance(k, str) and isinstance(v, str) and v.strip():
            clean_blocks[k.strip()] = v.strip()
    
    clean_notes = {
        "global": global_note.strip(),
        "blocks": clean_blocks
    }
    
    # Save to project.json under "notes_diarization"
    meta = read_project_meta(project_id)
    meta["notes_diarization"] = clean_notes
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    
    app_log(f"Project save diarization notes: project_id={project_id}, global_len={len(global_note)}, blocks={len(clean_blocks)}")
    
    return {
        "ok": True,
        "notes": clean_notes,
        "blocks_count": len(clean_blocks)
    }


@app.get("/api/projects/{project_id}/notes/diarization")
def api_get_notes_diarization(project_id: str) -> Any:
    """Get diarization notes from project.
    
    Returns:
    {
      "global": "...",
      "blocks": {"0": "...", "3": "..."}
    }
    """
    meta = read_project_meta(project_id)
    notes = meta.get("notes_diarization", {})
    
    # Ensure proper structure
    if not isinstance(notes, dict):
        notes = {"global": "", "blocks": {}}
    
    notes.setdefault("global", "")
    notes.setdefault("blocks", {})
    
    return notes


# Dodaj endpoint do zapisu transkrypcji (text/plain)
# Ten endpoint pozwala na wysyłanie zwykłego tekstu zamiast JSON
@app.post("/api/projects/{project_id}/save/transcript")
async def api_save_transcript_text(project_id: str, request: Request) -> Any:
    """Save transcript text (accepts text/plain).
    
    This is an alternative to /api/projects/{project_id}/save_transcript
    that accepts plain text instead of JSON payload.
    """
    text = (await request.body()).decode("utf-8")
    app_log(f"Project save transcript (text): project_id={project_id}, chars={len(text)}")
    path = project_path(project_id) / "transcript.txt"
    path.write_text(text, encoding="utf-8")
    meta = read_project_meta(project_id)
    meta["has_transcript"] = True
    meta["updated_at"] = now_iso()
    write_project_meta(project_id, meta)
    return {"ok": True}


# Dodaj endpoint do generowania raportów specyficznych dla transkrypcji
@app.get("/api/projects/{project_id}/report/transcription")
def api_generate_transcription_report(project_id: str, format: str = "pdf", include_logs: int = 0, include_notes: int = 0) -> Any:
    """Generate report specifically for transcription (without diarization data).

    Similar to /api/projects/{project_id}/report but focuses only on transcription.
    """
    project_path(project_id)  # ensure exists
    fmt = (format or "").lower()
    if fmt not in ("txt", "html", "pdf", "doc"):
        raise HTTPException(status_code=400, detail="format must be txt|html|pdf|doc")

    pdir = project_path(project_id)
    ts = time.strftime("%Y%m%d_%H%M%S")
    out_name = f"transcription_report_{ts}.{fmt}"
    out_path = pdir / out_name

    app_log(f"Transcription report requested: project_id={project_id}, format={fmt}, include_logs={bool(include_logs)}, include_notes={bool(include_notes)}")

    # Collect data similar to _collect_report_data but transcription-focused
    data = _collect_transcription_report_data(project_id, export_formats=[fmt], include_logs=bool(include_logs), include_notes=bool(include_notes))

    if fmt == "txt":
        generate_txt_report(data, logs=bool(include_logs), output_path=str(out_path))
        return FileResponse(str(out_path), filename=out_name)
    if fmt == "html":
        generate_html_report(data, logs=bool(include_logs), output_path=str(out_path))
        return FileResponse(str(out_path), filename=out_name)
    if fmt == "doc":
        # Word-compatible HTML saved with .doc extension (opens in Word/LibreOffice).
        generate_html_report(data, logs=bool(include_logs), output_path=str(out_path))
        try:
            html = out_path.read_text(encoding="utf-8", errors="ignore")
            if 'xmlns:w="urn:schemas-microsoft-com:office:word"' not in html:
                html = html.replace(
                    '<html',
                    '<html xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:w="urn:schemas-microsoft-com:office:word"',
                    1,
                )
                out_path.write_text(html, encoding="utf-8")
        except Exception:
            pass
        return FileResponse(str(out_path), filename=out_name)
    generate_pdf_report(data, logs=bool(include_logs), output_path=str(out_path))

    _cleanup_old_reports(pdir, "transcription_report_")
    return FileResponse(str(out_path), filename=out_name)


def _gather_report_notes(meta: Dict[str, Any], notes_key: str) -> Dict[str, Any] | None:
    """Read notes from project meta for inclusion in reports.

    Returns {"global": str, "blocks": {idx_str: str}} or None.
    """
    notes = meta.get(notes_key) or meta.get("notes") or {}
    if not isinstance(notes, dict):
        return None
    global_note = str(notes.get("global") or "").strip()
    blocks_raw = notes.get("blocks") or {}
    blocks: Dict[str, str] = {}
    if isinstance(blocks_raw, dict):
        for k, v in blocks_raw.items():
            txt = str(v or "").strip()
            if txt:
                blocks[str(k)] = txt
    if not global_note and not blocks:
        return None
    return {"global": global_note, "blocks": blocks}


def _collect_transcription_report_data(project_id: str, export_formats: List[str], include_logs: bool, include_notes: bool = False) -> Dict[str, Any]:
    """Collect data for transcription-specific report."""
    meta = read_project_meta(project_id)
    pdir = project_path(project_id)

    audio_file = meta.get("audio_file") or ""
    audio_duration = ""
    audio_specs = ""
    if audio_file:
        ap = pdir / audio_file
        if ap.exists():
            audio_duration, audio_specs = _probe_audio_basic(ap)

    transcript_lines: List[str] = []
    if (pdir / "transcript.txt").exists():
        transcript_lines = [ln.rstrip() for ln in (pdir / "transcript.txt").read_text(encoding="utf-8", errors="ignore").splitlines() if ln.strip()]

    s = load_settings()
    logs_text = ""
    if include_logs:
        parts = []
        for t in TASKS.list_tasks():
            if t.project_id == project_id and t.kind == "tr":
                parts.append(f"=== TASK {t.task_id} (transcription) ===")
                parts.extend(t.logs[-400:])
                parts.append("")
        logs_text = "\n".join(parts).strip()

    # Gather notes if requested
    notes_data = None
    if include_notes:
        notes_data = _gather_report_notes(meta, "notes_transcription")

    data = {
        "program_name": APP_NAME,
        "program_version": APP_VERSION,
        "author_email": AUTHOR_EMAIL,
        "processed_at": now_iso(),
        "audio_file": audio_file,
        "audio_duration": audio_duration,
        "audio_specs": audio_specs,
        "whisper_model": getattr(s, "whisper_model", "") or "",
        "language": meta.get("language", "auto"),
        "segments_count": len(transcript_lines),
        "transcript": transcript_lines,
        "raw_transcript": transcript_lines,
        "export_formats": export_formats,
        "logs": logs_text,
        "ui_language": "pl",
        "section_title": "Transkrypcja",
        "notes": notes_data,
    }
    return data

@app.post("/api/projects/import")
async def api_import_project(file: UploadFile = File(...)) -> Any:
    fname = (file.filename or "").lower()
    if not fname.endswith(".aistate"):
        raise HTTPException(status_code=400, detail="Plik musi mieć rozszerzenie .aistate")

    tmp_dir = Path(tempfile.mkdtemp(prefix="aistate_import_"))
    tmp_file = tmp_dir / "upload.aistate"
    try:
        # Stream upload to disk (avoid keeping whole audio in RAM)
        with tmp_file.open("wb") as out:
            shutil.copyfileobj(file.file, out)

        extract_dir = tmp_dir / "extracted"
        extract_dir.mkdir(parents=True, exist_ok=True)

        with zipfile.ZipFile(str(tmp_file), "r") as zf:
            _safe_extract_zip(zf, extract_dir)

        extracted_root = _pick_extracted_root(extract_dir)

        # Determine preferred project id (if present and unused)
        preferred_id = None
        pj = extracted_root / "project.json"
        if pj.exists():
            try:
                meta = json.loads(pj.read_text(encoding="utf-8"))
                if isinstance(meta, dict) and isinstance(meta.get("project_id"), str):
                    preferred_id = meta.get("project_id").strip() or None
            except Exception:
                preferred_id = None

        final_id = preferred_id
        if not final_id or (PROJECTS_DIR / final_id).exists():
            final_id = uuid.uuid4().hex

        dest = (PROJECTS_DIR / final_id).resolve()
        root = PROJECTS_DIR.resolve()
        if root not in dest.parents:
            raise HTTPException(status_code=400, detail="Nieprawidłowy docelowy katalog projektu.")
        if dest.exists():
            raise HTTPException(status_code=409, detail="Projekt o takim ID już istnieje.")

        shutil.copytree(extracted_root, dest)

        # Normalize project.json
        meta = read_project_meta(final_id)
        meta["project_id"] = final_id
        meta["created_at"] = meta.get("created_at") or now_iso()
        meta["updated_at"] = now_iso()
        write_project_meta(final_id, meta)

        return {"ok": True, "project_id": final_id}

    except HTTPException:
        raise
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="Nieprawidłowa paczka .aistate (uszkodzony ZIP).")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Błąd importu: {e}")
    finally:
        try:
            shutil.rmtree(tmp_dir, ignore_errors=True)
        except Exception:
            pass


@app.delete("/api/projects/{project_id}")
def api_delete_project(project_id: str, wipe_method: str = "none") -> Any:
    pdir = (PROJECTS_DIR / project_id).resolve()
    root = PROJECTS_DIR.resolve()
    if root not in pdir.parents or not pdir.exists() or not pdir.is_dir():
        raise HTTPException(status_code=404, detail="Nie ma takiego projektu.")

    try:
        secure_delete_project_dir(pdir, wipe_method)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Błąd usuwania projektu: {e}")

    return {"ok": True}




# ---------- API: reports ----------

def _probe_audio_basic(path: Path) -> tuple[str, str]:
    """Best-effort: duration/specs for WAV via stdlib wave; otherwise file size only."""
    import wave
    size_b = 0
    try:
        size_b = path.stat().st_size
    except Exception:
        pass
    size_mb = f"{int(round(size_b / (1024*1024)))}MB" if size_b else ""

    duration = ""
    specs = size_mb

    try:
        with wave.open(str(path), "rb") as wf:
            fr = wf.getframerate()
            n = wf.getnframes()
            ch = wf.getnchannels()
            dur_s = (n / float(fr)) if fr else 0.0
            duration = f"{dur_s:.1f}s"
            specs2 = f"{ch}ch {fr}Hz"
            specs = (specs2 + ((" • " + size_mb) if size_mb else "")).strip(" •")
    except Exception:
        pass
    return duration, specs


WAVEFORM_NUM_PEAKS = 800  # Match canvas width in seg_tools.js


def _generate_waveform_peaks(project_id: str) -> bool:
    """Generate waveform peak amplitudes from project audio and save as peaks.json.

    Uses soundfile (fast C library) for reading.  Falls back to wave stdlib.
    Returns True if peaks.json was written successfully.
    """
    pdir = project_path(project_id)
    meta = read_project_meta(project_id)
    audio_file = meta.get("audio_file") or ""
    if not audio_file:
        return False
    audio_path = pdir / audio_file
    if not audio_path.exists():
        return False

    peaks_path = pdir / "peaks.json"
    num_peaks = WAVEFORM_NUM_PEAKS

    try:
        # Try soundfile first (handles WAV, FLAC, OGG, etc.)
        import soundfile as sf
        import numpy as np
        data, sr = sf.read(str(audio_path), dtype="float32", always_2d=True)
        # Mix to mono
        mono = np.mean(data, axis=1) if data.shape[1] > 1 else data[:, 0]
        total = len(mono)
        block = max(1, total // num_peaks)
        peaks = []
        for i in range(num_peaks):
            start = i * block
            end = min(start + block, total)
            if start >= total:
                peaks.append(0.0)
            else:
                peaks.append(float(np.max(np.abs(mono[start:end]))))
        duration = total / sr if sr else 0.0
    except Exception:
        try:
            # Fallback: stdlib wave (WAV only)
            import wave
            import struct
            with wave.open(str(audio_path), "rb") as wf:
                sr = wf.getframerate()
                n = wf.getnframes()
                ch = wf.getnchannels()
                sw = wf.getsampwidth()
                raw = wf.readframes(n)

            # Decode to float samples
            if sw == 2:
                fmt = "<" + "h" * (n * ch)
                samples = struct.unpack(fmt, raw)
                scale = 32768.0
            elif sw == 1:
                fmt = "B" * (n * ch)
                samples = struct.unpack(fmt, raw)
                samples = [s - 128 for s in samples]
                scale = 128.0
            else:
                return False

            # Mix to mono
            if ch > 1:
                mono = [sum(samples[i:i + ch]) / ch for i in range(0, len(samples), ch)]
            else:
                mono = list(samples)

            total = len(mono)
            block = max(1, total // num_peaks)
            peaks = []
            for i in range(num_peaks):
                start = i * block
                end = min(start + block, total)
                if start >= total:
                    peaks.append(0.0)
                else:
                    mx = max(abs(mono[j]) for j in range(start, end))
                    peaks.append(mx / scale)
            duration = total / sr if sr else 0.0
        except Exception:
            return False

    # Normalize peaks to 0..1
    max_peak = max(peaks) if peaks else 1.0
    if max_peak > 0:
        peaks = [round(p / max_peak, 4) for p in peaks]

    payload = {"peaks": peaks, "duration": round(duration, 3), "num_peaks": num_peaks}
    try:
        tmp = peaks_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(payload), encoding="utf-8")
        tmp.replace(peaks_path)
    except Exception:
        return False

    return True


def _collect_report_data(project_id: str, export_formats: List[str], include_logs: bool, include_notes: bool = False) -> Dict[str, Any]:
    meta = read_project_meta(project_id)
    pdir = project_path(project_id)

    audio_file = meta.get("audio_file") or ""
    audio_duration = ""
    audio_specs = ""
    if audio_file:
        ap = pdir / audio_file
        if ap.exists():
            audio_duration, audio_specs = _probe_audio_basic(ap)

    transcript_lines: List[str] = []
    diarized_lines: List[str] = []
    if (pdir / "transcript.txt").exists():
        transcript_lines = [ln.rstrip() for ln in (pdir / "transcript.txt").read_text(encoding="utf-8", errors="ignore").splitlines() if ln.strip()]
    if (pdir / "diarized.txt").exists():
        diarized_lines = [ln.rstrip() for ln in (pdir / "diarized.txt").read_text(encoding="utf-8", errors="ignore").splitlines() if ln.strip()]

    s = load_settings()
    logs_text = ""
    if include_logs:
        parts = []
        for t in TASKS.list_tasks():
            if t.project_id == project_id:
                parts.append(f"=== TASK {t.task_id} ({t.kind}) ===")
                parts.extend(t.logs[-400:])
                parts.append("")
        logs_text = "\n".join(parts).strip()

    # Gather notes if requested
    notes_data = None
    if include_notes:
        notes_data = _gather_report_notes(meta, "notes_diarization")

    data = {
        "program_name": APP_NAME,
        "program_version": APP_VERSION,
        "author_email": AUTHOR_EMAIL,
        "processed_at": now_iso(),
        "audio_file": audio_file,
        "audio_duration": audio_duration,
        "audio_specs": audio_specs,
        "whisper_model": getattr(s, "whisper_model", "") or "",
        "language": "auto",
        "pyannote_model": "pyannote.audio",
        "speakers_count": "",
        "segments_count": len(diarized_lines) if diarized_lines else len(transcript_lines),
        "speaker_times": {},
        "transcript": diarized_lines or transcript_lines,
        "raw_transcript": transcript_lines,
        "non_verbal": [],
        "export_formats": export_formats,
        "logs": logs_text,
        "ui_language": "pl",
        "theme": "",
        "speaker_name_map": (meta.get("speaker_map") or {}),
        "section_title": "Transkrypcja / Diaryzacja",
        "notes": notes_data,
    }
    return data


def _cleanup_old_reports(pdir: Path, prefix: str = "report_", keep: int = 10) -> None:
    """Remove old report files, keeping the *keep* most recent per prefix."""
    try:
        candidates = sorted(
            [f for f in pdir.iterdir() if f.is_file() and f.name.startswith(prefix)],
            key=lambda f: f.stat().st_mtime,
            reverse=True,
        )
        for old in candidates[keep:]:
            try:
                old.unlink()
            except Exception:
                pass
    except Exception:
        pass


@app.get("/api/projects/{project_id}/report")
def api_generate_report(project_id: str, format: str = "pdf", include_logs: int = 0, include_notes: int = 0) -> Any:
    project_path(project_id)  # ensure exists
    fmt = (format or "").lower()
    if fmt not in ("txt", "html", "pdf", "doc"):
        raise HTTPException(status_code=400, detail="format must be txt|html|pdf|doc")

    pdir = project_path(project_id)
    ts = time.strftime("%Y%m%d_%H%M%S")
    out_name = f"report_{ts}.{fmt}"
    out_path = pdir / out_name

    app_log(f"Report requested: project_id={project_id}, format={fmt}, include_logs={bool(include_logs)}, include_notes={bool(include_notes)}")

    data = _collect_report_data(project_id, export_formats=[fmt], include_logs=bool(include_logs), include_notes=bool(include_notes))

    if fmt == "txt":
        generate_txt_report(data, logs=bool(include_logs), output_path=str(out_path))
        return FileResponse(str(out_path), filename=out_name)
    if fmt == "html":
        generate_html_report(data, logs=bool(include_logs), output_path=str(out_path))
        return FileResponse(str(out_path), filename=out_name)
    if fmt == "doc":
        # Word-compatible: HTML saved as .doc (lightweight, opens in Word/LibreOffice).
        generate_html_report(data, logs=bool(include_logs), output_path=str(out_path))
        try:
            html = out_path.read_text(encoding="utf-8", errors="ignore")
            if 'xmlns:w="urn:schemas-microsoft-com:office:word"' not in html and '<html' in html:
                html = html.replace(
                    '<html',
                    '<html xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:w="urn:schemas-microsoft-com:office:word"',
                    1,
                )
                out_path.write_text(html, encoding="utf-8")
        except Exception:
            pass
        return FileResponse(str(out_path), filename=out_name)
    generate_pdf_report(data, logs=bool(include_logs), output_path=str(out_path))

    # Trim old reports (keep last 10 per prefix)
    _cleanup_old_reports(pdir, "report_")
    return FileResponse(str(out_path), filename=out_name)


# Tasks endpoints moved to webapp/routers/tasks.py



# ---------- API: prompts (system + user) ----------

@app.get("/api/prompts/list")
async def api_prompts_list() -> Any:
    """Return all prompts (system + user)."""
    return PROMPTS.list_all()


@app.post("/api/prompts/create")
async def api_prompts_create(data: Dict[str, Any] = Body(...)) -> Any:
    """Create a new user prompt (stored in projects/_global/prompts)."""
    try:
        pid = PROMPTS.create_user_prompt(data)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"status": "created", "id": pid}


@app.put("/api/prompts/{prompt_id}")
async def api_prompts_update(prompt_id: str, updates: Dict[str, Any] = Body(...)) -> Any:
    """Update an existing user prompt."""
    try:
        PROMPTS.update_user_prompt(prompt_id, updates)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Prompt not found.")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"status": "updated", "id": prompt_id}


@app.delete("/api/prompts/{prompt_id}")
async def api_prompts_delete(prompt_id: str) -> Any:
    """Delete a user prompt."""
    try:
        PROMPTS.delete_user_prompt(prompt_id)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Prompt not found.")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"status": "deleted", "id": prompt_id}


@app.get("/api/prompts/{prompt_id}/export")
async def api_prompts_export(prompt_id: str) -> Any:
    """Export a user prompt as a JSON file."""
    try:
        fp = PROMPTS.export_user_prompt_path(prompt_id)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Prompt not found.")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return FileResponse(str(fp), media_type="application/json", filename=fp.name)


@app.post("/api/prompts/import")
async def api_prompts_import(file: UploadFile = File(...)) -> Any:
    """Import a prompt from a JSON file (creates a new user prompt)."""
    try:
        raw = await file.read()
        obj = json.loads(raw.decode("utf-8", errors="strict"))
        if not isinstance(obj, dict):
            raise ValueError("JSON must be an object.")
        pid = PROMPTS.import_user_prompt(obj)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="Invalid JSON file.")
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="File is not valid UTF-8.")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"status": "imported", "id": pid}



# ---------- API: documents (project-level attachments) ----------

def _documents_dir(project_id: str) -> Path:
    pdir = project_path(project_id)
    ddir = pdir / "documents"
    ddir.mkdir(parents=True, exist_ok=True)
    (ddir / ".cache").mkdir(parents=True, exist_ok=True)
    return ddir


def _doc_cache_paths(doc_path: Path) -> tuple[Path, Path]:
    cache_dir = doc_path.parent / ".cache"
    cache_txt = cache_dir / f"{doc_path.name}.txt"
    cache_meta = cache_dir / f"{doc_path.name}.meta.json"
    return cache_txt, cache_meta


async def _extract_and_cache_document(doc_path: Path) -> Dict[str, Any]:
    """Extract text and store cache files. Returns extraction summary."""
    cache_txt, cache_meta = _doc_cache_paths(doc_path)

    # cache hit if cache newer than doc and non-empty
    if cache_txt.exists() and cache_txt.stat().st_mtime >= doc_path.stat().st_mtime and cache_txt.stat().st_size > 0:
        try:
            meta = json.loads(cache_meta.read_text(encoding="utf-8")) if cache_meta.exists() else {}
        except Exception:
            meta = {}
        return {
            "cached": True,
            "cache_txt": str(cache_txt.name),
            "chars": int(cache_txt.stat().st_size),
            "metadata": meta,
        }

    # do extraction in threadpool (can be heavy)
    try:
        extracted = await run_in_threadpool(extract_text, doc_path)
        cache_txt.write_text(extracted.text or "", encoding="utf-8")
        cache_meta.write_text(json.dumps(extracted.metadata or {}, ensure_ascii=False, indent=2), encoding="utf-8")
        return {
            "cached": False,
            "cache_txt": str(cache_txt.name),
            "chars": len(extracted.text or ""),
            "metadata": extracted.metadata or {},
        }
    except DocumentProcessingError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Document processing failed: {e}")


@app.post("/api/documents/upload")
async def api_documents_upload(project_id: str = Form(""), file: UploadFile = File(...)) -> Any:
    """Upload a document to a project and cache extracted text."""
    project_id = ensure_project(project_id or None)
    ddir = _documents_dir(project_id)

    original_name = safe_filename(file.filename or "document")
    ext = Path(original_name).suffix.lower()
    if ext and ext not in SUPPORTED_EXTS:
        raise HTTPException(status_code=400, detail=f"Unsupported format: {ext}")

    dst = ddir / original_name
    raw = await file.read()
    dst.write_bytes(raw)
    app_log(f"Document upload: project_id={project_id}, file='{dst.name}', bytes={len(raw)}")

    extraction = await _extract_and_cache_document(dst)
    return {
        "status": "uploaded",
        "project_id": project_id,
        "filename": dst.name,
        "size": dst.stat().st_size,
        "type": ext.lstrip("."),
        "extraction": extraction,
    }


@app.get("/api/documents/{project_id}/list")
async def api_documents_list(project_id: str) -> Any:
    """List documents in a project."""
    ddir = _documents_dir(project_id)
    out: List[Dict[str, Any]] = []
    for fp in sorted(ddir.iterdir()):
        if fp.is_dir():
            continue
        if fp.name.startswith("."):
            continue
        cache_txt, _ = _doc_cache_paths(fp)
        out.append(
            {
                "filename": fp.name,
                "size": fp.stat().st_size,
                "type": fp.suffix.lower().lstrip("."),
                "cached": cache_txt.exists() and cache_txt.stat().st_size > 0,
                "download_url": f"/api/documents/{project_id}/download/{fp.name}",
            }
        )
    return out


@app.get("/api/documents/{project_id}/download/{filename}")
async def api_documents_download(project_id: str, filename: str) -> Any:
    ddir = _documents_dir(project_id)
    fname = safe_filename(filename)
    path = (ddir / fname).resolve()
    # ensure inside documents dir
    if ddir.resolve() not in path.parents:
        raise HTTPException(status_code=400, detail="Invalid filename")
    require_existing_file(path, "Plik nie istnieje.")
    return FileResponse(str(path), filename=fname)


@app.delete("/api/documents/{project_id}/{filename}")
async def api_documents_delete(project_id: str, filename: str) -> Any:
    ddir = _documents_dir(project_id)
    fname = safe_filename(filename)
    path = (ddir / fname).resolve()
    if ddir.resolve() not in path.parents:
        raise HTTPException(status_code=400, detail="Invalid filename")
    if not path.exists():
        raise HTTPException(status_code=404, detail="Document not found")

    # delete caches
    cache_txt, cache_meta = _doc_cache_paths(path)
    cache_txt.unlink(missing_ok=True)  # type: ignore[arg-type]
    cache_meta.unlink(missing_ok=True)  # type: ignore[arg-type]
    path.unlink(missing_ok=True)  # type: ignore[arg-type]
    app_log(f"Document delete: project_id={project_id}, file='{fname}'")
    return {"status": "deleted", "filename": fname}



# ---------- API: Ollama (status/models) ----------

@app.get("/api/ollama/status")
async def api_ollama_status() -> Any:
    st = await OLLAMA.status()
    d = st.to_dict()
    d["url"] = getattr(OLLAMA, "base_url", "http://127.0.0.1:11434")
    d["models_count"] = len(d.get("models") or [])
    return d


@app.get("/api/ollama/models")
async def api_ollama_models() -> Any:
    """Return available Ollama models (best-effort raw list)."""
    try:
        models = await OLLAMA.list_models()
        return models
    except OllamaError as e:
        raise HTTPException(status_code=503, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Ollama unavailable: {e}")

@app.post("/api/ollama/install")
async def api_ollama_install(payload: Dict[str, Any] = Body(...)) -> Any:
    """Install (pull) an Ollama model in the background so the UI can offer an 'Install' button."""
    model = str((payload or {}).get("model") or "").strip()
    scope = str((payload or {}).get("scope") or "unknown").strip() or "unknown"
    if not model:
        raise HTTPException(status_code=400, detail="model required")

    st = await OLLAMA.status()
    if st.status != "online":
        raise HTTPException(status_code=503, detail="Ollama offline. Start: ollama serve")

    # If already installed, return ok.
    try:
        installed = await OLLAMA.list_model_names()
    except Exception:
        installed = []
    if model in installed:
        _set_install_state(model, status="done", progress=100, stage="already installed", scope=scope, error=None, started_at=now_iso())
        app_log(f"Ollama install skipped (already installed): model={model}, scope={scope}")
        return {"status": "done", "model": model, "scope": scope}

    # If task is already running, do not start again.
    if OLLAMA_INSTALL_TASKS.get(model, {}).get("status") == "running":
        return {"status": "running", "model": model, "scope": scope}

    # Start background pull
    t = threading.Thread(target=_ollama_install_worker, args=(model, scope), daemon=True)
    t.start()
    return {"status": "started", "model": model, "scope": scope}


@app.get("/api/ollama/install/status")
async def api_ollama_install_status(model: str) -> Any:
    """Return current install status for a model."""
    mid = str(model or "").strip()
    if not mid:
        raise HTTPException(status_code=400, detail="model required")
    return OLLAMA_INSTALL_TASKS.get(mid, {"status": "idle", "progress": 0, "stage": "", "scope": "unknown", "error": None})




# ---------- API: analysis (quick + deep) ----------

def _analysis_dir(project_id: str) -> Path:
    adir = project_path(project_id) / "analysis"
    adir.mkdir(parents=True, exist_ok=True)
    (adir / "reports").mkdir(parents=True, exist_ok=True)
    (adir / "prompts").mkdir(parents=True, exist_ok=True)
    return adir


def _analysis_reports_dir(project_id: str) -> Path:
    return _analysis_dir(project_id) / "reports"


def _analysis_ui_state_path(project_id: str) -> Path:
    """Per-project UI state for Analysis tab (custom prompt, source checkboxes, etc.)."""
    return _analysis_dir(project_id) / "ui_state.json"


def _analysis_deep_task_state_path(project_id: str) -> Path:
    """Per-project deep analysis task state (running task id + metadata)."""
    return _analysis_dir(project_id) / "deep_task.json"


def _analysis_deep_task_output_path(project_id: str) -> Path:
    """Per-project streaming output for deep analysis (grows while task runs)."""
    return _analysis_dir(project_id) / "deep_running.md"


def _analysis_deep_latest_path(project_id: str) -> Path:
    """Final (latest) deep analysis payload."""
    return _analysis_dir(project_id) / "deep_latest.json"


def _read_json_file(path: Path, default: Any) -> Any:
    try:
        if not path.exists():
            return default
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def _write_json_file(path: Path, obj: Any) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def _safe_child_path(parent: Path, filename: str) -> Path:
    """Resolve a child file path and prevent path traversal."""
    fname = safe_filename(filename)
    candidate = (parent / fname).resolve()
    parent_res = parent.resolve()
    if parent_res not in candidate.parents and candidate != parent_res:
        raise HTTPException(status_code=400, detail="Invalid filename")
    return candidate


def _load_text_file(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""


async def _gather_analysis_sources(project_id: str, include_sources: Dict[str, Any]) -> str:
    """Load selected sources and return a combined string (best-effort)."""
    pdir = project_path(project_id)
    parts: List[str] = []

    # transcript / diarization
    if include_sources.get("transcript", True):
        tp = pdir / "transcript.txt"
        if tp.exists():
            parts.append("## Transkrypcja\n" + _load_text_file(tp).strip())
    if include_sources.get("diarization", True):
        dp = pdir / "diarized.txt"
        if dp.exists():
            parts.append("## Diaryzacja\n" + _load_text_file(dp).strip())

    # notes: prefer spec files if present, fallback to project.json fields
    notes_global = ""
    notes_blocks: Dict[str, str] = {}

    # Spec filenames (if user created manually)
    ng = pdir / "notes_global.txt"
    nb = pdir / "notes_blocks.json"
    if ng.exists():
        notes_global = _load_text_file(ng)
    if nb.exists():
        try:
            obj = json.loads(_load_text_file(nb))
            if isinstance(obj, dict):
                notes_blocks = {str(k): str(v) for k, v in obj.items() if str(v).strip()}
        except Exception:
            pass

    # Project meta notes
    meta = read_project_meta(project_id)
    for key in ("notes", "notes_transcription", "notes_diarization"):
        v = meta.get(key)
        if isinstance(v, dict):
            g = v.get("global")
            b = v.get("blocks")
            if isinstance(g, str) and g.strip():
                notes_global += ("\n\n" if notes_global.strip() else "") + f"[{key}]\n" + g.strip()
            if isinstance(b, dict):
                for bk, bv in b.items():
                    if isinstance(bv, str) and bv.strip():
                        notes_blocks.setdefault(str(bk), bv.strip())

    if include_sources.get("notes_global", False) and notes_global.strip():
        parts.append("## Notatki globalne\n" + notes_global.strip())
    if include_sources.get("notes_blocks", False) and notes_blocks:
        # Keep it readable
        blocks_md = "\n".join([f"- Blok {k}: {v}" for k, v in sorted(notes_blocks.items(), key=lambda x: x[0])])
        parts.append("## Notatki do bloków\n" + blocks_md)

    # Documents — auto-detect bank statements and enrich with structured data
    docs = include_sources.get("documents") or []
    if isinstance(docs, list) and docs:
        ddir = _documents_dir(project_id)
        for name in docs:
            fname = safe_filename(str(name))
            fp = (ddir / fname).resolve()
            if ddir.resolve() not in fp.parents or not fp.exists():
                continue
            cache_txt, _ = _doc_cache_paths(fp)
            if not cache_txt.exists() or cache_txt.stat().st_mtime < fp.stat().st_mtime:
                # (re)extract
                try:
                    await _extract_and_cache_document(fp)
                except Exception:
                    pass
            if cache_txt.exists() and cache_txt.stat().st_size > 0:
                doc_text = _load_text_file(cache_txt).strip()
                # Auto-detect bank statements and provide structured extraction
                if fp.suffix.lower() == ".pdf":
                    try:
                        from backend.finance.detector import is_bank_statement as _is_stmt
                        _detected, _score, _ = _is_stmt(doc_text[:10000])
                        if _detected:
                            from backend.finance.quick_extract import extract_bank_statement_quick
                            quick_data = extract_bank_statement_quick(doc_text)
                            structured = "\n".join([
                                f"## Dokument: {fname} [WYCIĄG BANKOWY — dane wyodrębnione automatycznie]\n",
                                f"- **Bank**: {quick_data.get('bank') or 'nierozpoznany'}",
                                f"- **Właściciel**: {quick_data.get('wlasciciel_rachunku') or '—'}",
                                f"- **IBAN**: {quick_data.get('nr_rachunku_iban') or '—'}",
                                f"- **Okres**: {quick_data.get('okres') or '—'}",
                                f"- **Waluta**: {quick_data.get('waluta') or 'PLN'}",
                                f"- **Saldo początkowe**: {quick_data.get('saldo_poczatkowe') or '—'}",
                                f"- **Saldo końcowe**: {quick_data.get('saldo_koncowe') or '—'}",
                                f"- **Saldo dostępne**: {quick_data.get('saldo_dostepne') or '—'}",
                                f"- **Suma uznań**: {quick_data.get('suma_uznan') or '—'}",
                                f"- **Suma obciążeń**: {quick_data.get('suma_obciazen') or '—'}",
                                f"- **Liczba transakcji**: {quick_data.get('liczba_transakcji') or '—'}",
                                "",
                                "**UWAGA**: Powyższe dane wyodrębnione regexem z nagłówka PDF. "
                                "Dla pełnej analizy finansowej użyj trybu 'Analiza wyciągu bankowego'.",
                                "",
                                "### Surowy tekst dokumentu\n",
                                doc_text[:50000],
                            ])
                            parts.append(structured)
                            continue
                    except Exception:
                        pass
                parts.append(f"## Dokument: {fname}\n" + doc_text)

    # Cap size (avoid accidental huge prompts)
    combined = "\n\n---\n\n".join([p for p in parts if p.strip()]).strip()
    max_chars = int(include_sources.get("max_chars") or 200_000)
    if len(combined) > max_chars:
        combined = combined[:max_chars] + "\n\n[... ucięto materiał: przekroczono limit znaków ...]"
    return combined


async def _run_quick_analysis(project_id: str) -> Dict[str, Any]:
    pdir = project_path(project_id)
    adir = _analysis_dir(project_id)
    qs_path = adir / "quick_summary.json"

    # Prefer transcript; fallback to diarized
    t = _load_text_file(pdir / "transcript.txt")
    d = _load_text_file(pdir / "diarized.txt")
    text = (t.strip() + "\n\n" + d.strip()).strip()

    # Fallback: load text from uploaded documents
    _source_type = "transcript"  # track source for prompt selection
    if not text:
        ddir = _documents_dir(project_id)
        doc_parts: list[str] = []
        if ddir.exists():
            for fp in sorted(ddir.iterdir()):
                if fp.is_dir() or fp.name.startswith("."):
                    continue
                cache_txt, _ = _doc_cache_paths(fp)
                if not cache_txt.exists() or cache_txt.stat().st_mtime < fp.stat().st_mtime:
                    try:
                        await _extract_and_cache_document(fp)
                    except Exception:
                        continue
                if cache_txt.exists() and cache_txt.stat().st_size > 0:
                    doc_parts.append(_load_text_file(cache_txt).strip())
        if doc_parts:
            text = "\n\n---\n\n".join(doc_parts)
            # Detect if this is a bank statement
            from backend.finance.detector import is_bank_statement as _detect_bank_stmt
            _is_bank, _, _ = _detect_bank_stmt(text[:10000])
            _source_type = "bank_statement" if _is_bank else "document"

    if not text:
        raise HTTPException(status_code=400, detail="Brak źródeł do analizy (transkrypcja, diaryzacja lub dokumenty).")

    t0 = time.time()

    if _source_type == "bank_statement":
        # Bank statement: use regex extractor (fast, accurate, no LLM needed)
        from backend.finance.quick_extract import extract_bank_statement_quick
        result = extract_bank_statement_quick(text)
        model_sel = "regex"
    else:
        try:
            model_sel = _get_model_settings().get("quick") or DEFAULT_MODELS["quick"]
            result = await quick_analyze(OLLAMA, text, model=model_sel, source_type=_source_type)
        except OllamaError as e:
            raise HTTPException(status_code=503, detail=str(e))
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Quick analysis failed: {e}")

    dt = time.time() - t0

    qs_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    # Save metadata sidecar (model + timestamps) for UI display
    try:
        meta = {
            "model": model_sel,
            "generated_at": now_iso(),
            "generation_time": round(dt, 2),
        }
        (adir / "quick_summary.meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass
    return {"status": "success", "result": result, "generation_time": round(dt, 2), "meta": meta if "meta" in locals() else None}


def _quick_analysis_task_runner(project_id: str, log_cb=None, progress_cb=None) -> dict:
    """Run quick analysis via TaskManager thread (GPU RM managed)."""
    try:
        if progress_cb:
            progress_cb(5)
        if log_cb:
            log_cb("Quick analysis (GPU RM): starting")
        res = asyncio.run(_run_quick_analysis(project_id))
        if progress_cb:
            progress_cb(100)
        if log_cb:
            log_cb("Quick analysis (GPU RM): done")
        # _run_quick_analysis returns a response dict; we store it as result too
        return res
    except Exception as e:
        if log_cb:
            log_cb(f"Quick analysis (GPU RM) failed: {e}")
        raise


@app.get("/api/analysis/quick/{project_id}")
async def api_analysis_get_quick(project_id: str) -> Any:
    """Return cached quick summary if present."""
    qs = _analysis_dir(project_id) / "quick_summary.json"
    if not qs.exists():
        raise HTTPException(status_code=404, detail="Brak szybkiej analizy")
    try:
        data = json.loads(qs.read_text(encoding="utf-8"))
        meta_path = qs.with_name("quick_summary.meta.json")
        meta = None
        if meta_path.exists():
            try:
                meta = json.loads(meta_path.read_text(encoding="utf-8"))
            except Exception:
                meta = None
        return {"status": "success", "result": data, "meta": meta}
    except Exception:
        raise HTTPException(status_code=500, detail="Nie można odczytać quick_summary.json")


@app.post("/api/analysis/quick")
async def api_analysis_quick(payload: Dict[str, Any] = Body(...)) -> Any:
    """Run quick analysis (auto after transcription) and store quick_summary.json."""
    project_id = str(payload.get("project_id") or "").strip()
    if not project_id:
        raise HTTPException(status_code=400, detail="project_id required")
    app_log(f"Quick analysis requested: project_id={project_id}")

    # If GPU RM is enabled, queue the LLM call so it does not overlap with transcription/diarization.
    if GPU_RM.enabled:
        t = GPU_RM.enqueue_python_fn("analysis_quick", project_id, _quick_analysis_task_runner, project_id)
        return {"status": "queued", "task_id": t.task_id}

    return await _run_quick_analysis(project_id)


@app.post("/api/analysis/deep")
async def api_analysis_deep(payload: Dict[str, Any] = Body(...)) -> Any:
    """Deep analysis (non-stream) - returns generated text.

    Report saving/formatting is implemented in step 4.
    """
    project_id = str(payload.get("project_id") or "").strip()
    if not project_id:
        raise HTTPException(status_code=400, detail="project_id required")

    template_ids = payload.get("template_ids") or []
    if isinstance(template_ids, str):
        template_ids = [t.strip() for t in template_ids.split(",") if t.strip()]
    if not isinstance(template_ids, list):
        template_ids = []

    custom_prompt = payload.get("custom_prompt")
    use_templates = payload.get("custom_prompt_use_templates")
    if isinstance(use_templates, str):
        use_templates = use_templates.strip().lower() not in ("0", "false", "no", "off")
    elif isinstance(use_templates, bool):
        use_templates = bool(use_templates)
    else:
        use_templates = True
    if custom_prompt and not use_templates:
        template_ids = []
    model = str(payload.get("model") or _get_model_settings().get("deep") or DEFAULT_MODELS["deep"])
    output_format = str(payload.get("output_format") or payload.get("format") or "").lower().strip()
    title = str(payload.get("title") or payload.get("report_title") or "").strip() or None
    include_sources = payload.get("include_sources") or {}
    if not isinstance(include_sources, dict):
        include_sources = {}
    include_sources = {
        "transcript": bool(include_sources.get("transcript", False)),
        "diarization": bool(include_sources.get("diarization", False)),
        "notes_global": bool(include_sources.get("notes_global", False)),
        "notes_blocks": bool(include_sources.get("notes_blocks", False)),
    }

    # Result mode for deep analysis output
    result_mode = str(payload.get("result_mode") or "replace").strip().lower()
    if result_mode not in ("replace", "append"):
        result_mode = "replace"
    append_header = str(payload.get("append_header") or "")
    if len(append_header) > 5000:
        append_header = append_header[:5000]

    if result_mode == "append":
        base_content = ""
        try:
            dl = _analysis_deep_latest_path(project_id)
            if dl.exists():
                j = _read_json_file(dl, None)
                if isinstance(j, dict):
                    base_content = str(j.get("content") or "")
        except Exception:
            base_content = ""
        if not base_content.strip():
            result_mode = "replace"
            append_header = ""
        elif not append_header:
            stamp = now_iso()[:16]
            append_header = f"\n\n## Dodatkowa analiza ({stamp}, model={model})\n\n"

    include_sources = {
        "transcript": bool(include_sources.get("transcript", False)),
        "diarization": bool(include_sources.get("diarization", False)),
        "notes_global": bool(include_sources.get("notes_global", False)),
        "notes_blocks": bool(include_sources.get("notes_blocks", False)),
    }

    # Result mode (append vs replace) + optional append header
    result_mode = str(payload.get("result_mode") or "replace").strip().lower()
    if result_mode not in ("replace", "append"):
        result_mode = "replace"
    append_header = str(payload.get("append_header") or "")
    if len(append_header) > 5000:
        append_header = append_header[:5000]

    if result_mode == "append":
        base_content = ""
        dl = _analysis_deep_latest_path(project_id)
        if dl.exists():
            j = _read_json_file(dl, None)
            if isinstance(j, dict) and isinstance(j.get("content"), str):
                base_content = j.get("content") or ""
        if not base_content.strip():
            # No base analysis -> fall back to replace
            result_mode = "replace"
            append_header = ""
        elif not append_header.strip():
            stamp = now_iso()[:16]
            append_header = f"\n\n## Dodatkowa analiza ({stamp}, model={model})\n\n"
    include_sources = {
        "transcript": bool(include_sources.get("transcript", False)),
        "diarization": bool(include_sources.get("diarization", False)),
        "notes_global": bool(include_sources.get("notes_global", False)),
        "notes_blocks": bool(include_sources.get("notes_blocks", False)),
    }

    # Result mode (replace vs append). Append works only if there is an existing deep analysis.
    result_mode = str(payload.get("result_mode") or "replace").strip().lower()
    if result_mode not in ("replace", "append"):
        result_mode = "replace"
    append_header = str(payload.get("append_header") or "")
    if len(append_header) > 5000:
        append_header = append_header[:5000]
    if result_mode == "append":
        latest_p = _analysis_deep_latest_path(project_id)
        latest = _read_json_file(latest_p, None) if latest_p.exists() else None
        base_content = str(latest.get("content") or "") if isinstance(latest, dict) else ""
        if not base_content.strip():
            result_mode = "replace"
            append_header = ""
        elif not append_header.strip():
            stamp = now_iso()[:16]
            append_header = f"\n\n## Dodatkowa analiza ({stamp}, model={model})\n\n"

    # Combine prompt templates
    try:
        instruction = PROMPTS.build_combined_prompt([str(x) for x in template_ids], str(custom_prompt) if custom_prompt else None)
        for pid in template_ids:
            try:
                PROMPTS.bump_usage(str(pid))
            except Exception:
                pass
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid prompt selection: {e}")

    sources_text = await _gather_analysis_sources(project_id, include_sources)
    if not sources_text.strip():
        raise HTTPException(status_code=400, detail="Brak źródeł do analizy")

    final_prompt = (instruction.strip() + "\n\n---\n\n" if instruction.strip() else "") + "# Materiał źródłowy\n\n" + sources_text

    # Keep a copy of the prompt used (for reproducibility)
    adir = _analysis_dir(project_id)
    ts = time.strftime("%Y%m%d_%H%M%S")
    (adir / "prompts" / f"deep_{ts}.md").write_text(final_prompt, encoding="utf-8")

    t0 = time.time()
    try:
        text = await deep_analyze(
            OLLAMA,
            final_prompt,
            model=model,
            system="Jesteś ekspertem w analizie dokumentów i rozmów. Twórz raporty po polsku, jasno i strukturalnie.",
            options={"temperature": 0.7, "num_ctx": 32768},
        )
    except OllamaError as e:
        raise HTTPException(status_code=503, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Deep analysis failed: {e}")
    dt = time.time() - t0

    # store last output
    (adir / "deep_last.md").write_text(text, encoding="utf-8")

    report_info: Optional[Dict[str, Any]] = None
    if output_format:
        try:
            res = save_report(
                reports_dir=_analysis_reports_dir(project_id),
                content=text,
                output_format=output_format,
                title=title or ("_".join(template_ids) if template_ids else "Analiza"),
                template_ids=[str(x) for x in template_ids],
                project_id=project_id,
                model=model,
            )
            download_url = f"/api/analysis/reports/{project_id}/download/{res.filename}"
            report_info = {
                "filename": res.filename,
                "format": res.format,
                "size_bytes": res.size_bytes,
                "download_url": download_url,
            }
        except ReportSaveError as e:
            raise HTTPException(status_code=400, detail=str(e))
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Report save failed: {e}")

    # Persist last deep analysis result inside the project so it is visible after reload/export/import.
    try:
        adir = _analysis_dir(project_id)
        deep_payload = {
            "content": text,
            "meta": {
                "model": model,
                "generated_at": now_iso(),
                "generation_time": round(dt, 2),
                "template_ids": [str(x) for x in template_ids],
                "custom_prompt": str(custom_prompt).strip() if custom_prompt else None,
                "custom_prompt_use_templates": bool(custom_prompt_use_templates) if custom_prompt else None,
                "include_sources": include_sources,
                "stream": False,
            },
            "report": report_info,
        }
        (adir / "deep_latest.json").write_text(json.dumps(deep_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass

    return {
        "status": "success",
        "preview": text[:4000],
        "content": text,
        "generation_time": round(dt, 2),
        "tokens_used": None,
        "report": report_info,
    }



@app.get("/api/analysis/deep/{project_id}")
async def api_analysis_get_deep(project_id: str) -> Any:
    """Return last saved deep analysis if present (content + meta)."""
    path = _analysis_dir(project_id) / "deep_latest.json"
    if not path.exists():
        raise HTTPException(status_code=404, detail="Brak głębokiej analizy")
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        raise HTTPException(status_code=500, detail="Nie można odczytać deep_latest.json")
    return {"status": "success", "result": obj}


@app.get("/api/analysis/ui_state/{project_id}")
async def api_analysis_get_ui_state(project_id: str) -> Any:
    """Per-project Analysis tab UI state (custom prompt + sources).

    Stored inside the project so it survives reload/export/import.
    """
    # Defaults: user explicitly chooses sources; transcript/diarization default OFF
    defaults = {
        "custom_prompt": "",
        "custom_prompt_use_templates": True,
        "selected_templates": [],
        "selected_docs": [],
        "result_mode": "replace",
        "sources": {
            "transcript": False,
            "diarization": False,
            "notes_global": False,
            "notes_blocks": False,
        },
    }
    path = _analysis_ui_state_path(project_id)
    data = _read_json_file(path, defaults)
    # Ensure keys exist
    if not isinstance(data, dict):
        data = defaults
    data.setdefault("custom_prompt", "")
    data.setdefault("custom_prompt_use_templates", True)
    # New fields (backward compatible)
    templates = data.get("selected_templates")
    if not isinstance(templates, list):
        templates = []
    data["selected_templates"] = [str(x) for x in templates if x is not None]

    docs = data.get("selected_docs")
    if not isinstance(docs, list):
        docs = []
    data["selected_docs"] = [str(x) for x in docs if x is not None]
    # result_mode is now fixed to append (no UI toggle)
    data["result_mode"] = "append"
    src = data.get("sources")
    if not isinstance(src, dict):
        src = {}
    merged = dict(defaults["sources"])
    merged.update({k: bool(src.get(k)) for k in merged.keys()})
    data["sources"] = merged
    return {"status": "success", "state": data}


@app.post("/api/analysis/ui_state/{project_id}")
async def api_analysis_save_ui_state(project_id: str, payload: Dict[str, Any] = Body(...)) -> Any:
    """Persist Analysis tab UI state for a project."""
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="Invalid payload")

    state = {
        "custom_prompt": str(payload.get("custom_prompt") or ""),
        "custom_prompt_use_templates": bool(payload.get("custom_prompt_use_templates", True)),
        "selected_templates": payload.get("selected_templates") if isinstance(payload.get("selected_templates"), list) else [],
        "selected_docs": payload.get("selected_docs") if isinstance(payload.get("selected_docs"), list) else [],
        "result_mode": "append",
        "sources": payload.get("sources") if isinstance(payload.get("sources"), dict) else {},
    }

    # Normalize sources
    norm_sources = {}
    for k in ("transcript", "diarization", "notes_global", "notes_blocks"):
        norm_sources[k] = bool((state.get("sources") or {}).get(k))
    state["sources"] = norm_sources

    # Normalize templates / docs
    state["selected_templates"] = [str(x) for x in (state.get("selected_templates") or []) if str(x).strip()]
    state["selected_docs"] = [str(x) for x in (state.get("selected_docs") or []) if str(x).strip()]
    if state.get("result_mode") not in ("replace", "append"):
        state["result_mode"] = "replace"

    try:
        _write_json_file(_analysis_ui_state_path(project_id), state)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Cannot save ui_state: {e}")

    return {"status": "saved", "state": state}


def _persist_deep_task_state(project_id: str, state: Dict[str, Any]) -> None:
    """Write deep task state for a project (best-effort)."""
    try:
        _write_json_file(_analysis_deep_task_state_path(project_id), state)
    except Exception:
        pass


async def _run_deep_analysis_task(
    *,
    task_id: str,
    project_id: str,
    model: str,
    template_ids: List[str],
    custom_prompt: Optional[str],
    custom_prompt_use_templates: bool,
    include_sources: Dict[str, Any],
    result_mode: str = "replace",
    append_header: str = "",
    log_cb=None,
    progress_cb=None,
) -> None:
    """Background deep analysis task.

    Streams output to project/analysis/deep_running.md so UI can reconnect/poll.
    Persists final result to deep_latest.json.
    """
    adir = _analysis_dir(project_id)
    out_path = _analysis_deep_task_output_path(project_id)

    def _log(msg: str) -> None:
        if log_cb:
            try:
                log_cb(msg)
            except Exception:
                pass

    def _prog(p: float) -> None:
        if progress_cb:
            try:
                progress_cb(p)
            except Exception:
                pass

    # Combine prompt templates
    try:
        instruction = PROMPTS.build_combined_prompt([str(x) for x in template_ids], str(custom_prompt) if custom_prompt else None)
        for pid in template_ids:
            try:
                PROMPTS.bump_usage(str(pid))
            except Exception:
                pass
    except Exception as e:
        raise RuntimeError(f"Invalid prompt selection: {e}")

    # --- Finance pipeline: intercept when wyciag_bankowy template is selected ---
    _finance_mode = "wyciag_bankowy" in [str(x) for x in template_ids]
    _finance_prompt = None

    if _finance_mode:
        _log("Tryb analizy finansowej — uruchamiam pipeline...")
        _prog(3)
        docs = include_sources.get("documents") or []
        ddir = _documents_dir(project_id)
        finance_dir = project_path(project_id) / "finance"

        # Collect all PDF paths with cached text
        pdf_paths = []
        for doc_name in docs:
            fname = safe_filename(str(doc_name))
            fp = (ddir / fname).resolve()
            if not fp.exists() or fp.suffix.lower() != ".pdf":
                continue

            cache_txt_path, _ = _doc_cache_paths(fp)
            cached_text = None
            if cache_txt_path.exists():
                try:
                    cached_text = cache_txt_path.read_text(encoding="utf-8", errors="replace")
                except Exception:
                    pass
            pdf_paths.append((fp, cached_text))

        if pdf_paths:
            _log(f"Znaleziono {len(pdf_paths)} PDF-ów do analizy finansowej")

            if len(pdf_paths) == 1:
                # Single document — use simple pipeline
                result = await run_in_threadpool(
                    run_finance_pipeline,
                    pdf_path=pdf_paths[0][0],
                    cached_text=pdf_paths[0][1],
                    save_dir=finance_dir,
                    global_dir=PROJECTS_DIR,
                    log_cb=_log,
                )
            else:
                # Multiple documents — use multi-statement pipeline with behavioral analysis
                result = await run_in_threadpool(
                    run_multi_statement_pipeline,
                    pdf_paths=pdf_paths,
                    save_dir=finance_dir,
                    global_dir=PROJECTS_DIR,
                    log_cb=_log,
                )

            if result:
                # LLM fallback for unclassified transactions (on primary result)
                try:
                    from backend.finance.llm_classifier import classify_with_llm
                    llm_model = model
                    llm_updated = await classify_with_llm(
                        result["classified"],
                        OLLAMA,
                        model=llm_model,
                        log_cb=_log,
                    )
                    if llm_updated:
                        from backend.finance.scorer import compute_score as recompute_score
                        result["score"] = recompute_score(result["classified"])
                        _log(f"Score po LLM fallback: {result['score'].total_score}/100")
                except Exception as e:
                    _log(f"LLM fallback pominięty: {e}")

                _finance_prompt = build_finance_enriched_prompt(
                    result,
                    original_instruction=instruction,
                )
                n_txns = len(result["classified"])
                score_val = result["score"].total_score
                behavioral = result.get("behavioral")
                if behavioral:
                    _log(f"Finance pipeline: {n_txns} transakcji, score={score_val}/100, "
                         f"behavioral={behavioral.total_months} mies., trajektoria={behavioral.debt_trajectory}")
                else:
                    _log(f"Finance pipeline: {n_txns} transakcji, score={score_val}/100")

        if not _finance_prompt:
            _log("Nie wykryto wyciągu bankowego w dokumentach — przechodzę do standardowej analizy.")
            _finance_mode = False

    if _finance_mode and _finance_prompt:
        final_prompt = _finance_prompt
    else:
        sources_text = await _gather_analysis_sources(project_id, include_sources)
        if not sources_text.strip():
            raise RuntimeError("Brak źródeł do analizy")
        final_prompt = (instruction.strip() + "\n\n---\n\n" if instruction.strip() else "") + "# Materiał źródłowy\n\n" + sources_text

    # keep reproducibility copy
    ts = time.strftime("%Y%m%d_%H%M%S")
    (adir / "prompts" / f"deep_{ts}.md").write_text(final_prompt, encoding="utf-8")

    # Persist task state
    state = {
        "task_id": task_id,
        "project_id": project_id,
        "model": model,
        "template_ids": [str(x) for x in template_ids],
        "custom_prompt": str(custom_prompt).strip() if custom_prompt else None,
        "custom_prompt_use_templates": bool(custom_prompt_use_templates) if custom_prompt else None,
        "include_sources": include_sources,
        "result_mode": str(result_mode or "replace"),
        "append_header": str(append_header or ""),
        "status": "running",
        "stage": "start",
        "progress": 0,
        "started_at": now_iso(),
        "updated_at": now_iso(),
    }
    _persist_deep_task_state(project_id, state)

    # Ensure model
    _log(f"Deep analysis start (task_id={task_id})")
    _log(f"Model: {model}")
    state["stage"] = "ensure_model"
    state["updated_at"] = now_iso()
    _persist_deep_task_state(project_id, state)
    _prog(2)
    await OLLAMA.ensure_model(model)

    # Stream output to file
    t0 = time.time()
    written = 0
    # Optional prefix when appending additional analysis
    prefix = ""
    eff_mode = str(result_mode or "replace").strip().lower()
    if eff_mode == "append":
        base = ""
        try:
            dl = _analysis_deep_latest_path(project_id)
            if dl.exists():
                prev = _read_json_file(dl, None)
                if isinstance(prev, dict) and prev.get("content"):
                    base = str(prev.get("content") or "")
        except Exception:
            base = ""
        if base.strip():
            ah = str(append_header or "")
            if not ah:
                # Server-generated header fallback
                stamp = now_iso()[:16]
                ah = f"\n\n## Dodatkowa analiza ({stamp}, model={model})\n\n"
            prefix = base + ah
            state["result_mode"] = "append"
            state["append_header"] = ah
            _persist_deep_task_state(project_id, state)
        else:
            # Nothing to append to -> behave like replace
            state["result_mode"] = "replace"
            state["append_header"] = ""
            _persist_deep_task_state(project_id, state)

    out_path.write_text(prefix, encoding="utf-8")
    written = len(prefix.encode("utf-8"))
    state["stage"] = "generating"
    state["updated_at"] = now_iso()
    _persist_deep_task_state(project_id, state)

    # Heuristic progress: based on output size.
    # We keep progress between 5..95 while streaming, then 100 at the end.
    def _estimate_progress(chars: int) -> int:
        # 8k chars -> ~25%, 20k -> ~50%, 40k -> ~75%, 80k -> ~92%
        if chars <= 0:
            return 5
        if chars < 8000:
            return 5 + int(chars / 8000 * 20)
        if chars < 20000:
            return 25 + int((chars - 8000) / 12000 * 25)
        if chars < 40000:
            return 50 + int((chars - 20000) / 20000 * 25)
        if chars < 80000:
            return 75 + int((chars - 40000) / 40000 * 17)
        return 92

    if _finance_mode:
        system_msg = (
            "Jesteś doświadczonym analitykiem finansowym specjalizującym się w analizie wyciągów bankowych, "
            "ocenie zdolności kredytowej i wykrywaniu anomalii transakcyjnych. "
            "Tworzysz profesjonalne raporty po polsku. Dla każdego wniosku podajesz poziom pewności. "
            "Zachowujesz ostrożność interpretacyjną — nie nadinterpretujesz danych."
        )
    else:
        system_msg = "Jesteś ekspertem w analizie dokumentów i rozmów. Twórz raporty po polsku, jasno i strukturalnie."
    msgs = [
        {"role": "system", "content": system_msg},
        {"role": "user", "content": final_prompt},
    ]

    # Use large context by default for deep.
    options = {"temperature": 0.7, "num_ctx": 32768}
    if _finance_mode:
        options["temperature"] = 0.4  # lower temp for factual financial analysis

    # Try streaming with fallback to smaller context on failure
    _ctx_sizes = [32768, 16384, 8192]
    _stream_ok = False
    for _ctx_attempt in _ctx_sizes:
        options["num_ctx"] = _ctx_attempt
        try:
            with out_path.open("a", encoding="utf-8") as f:
                async for chunk in OLLAMA.stream_chat(model=model, messages=msgs, options=options):
                    if not chunk:
                        continue
                    f.write(chunk)
                    f.flush()
                    written += len(chunk)
                    if written - state.get("_last_chars", 0) >= 300:
                        state["_last_chars"] = written
                        p = _estimate_progress(written)
                        state["progress"] = p
                        state["updated_at"] = now_iso()
                        _persist_deep_task_state(project_id, {k: v for k, v in state.items() if not str(k).startswith("_")})
                        _prog(p)
            _stream_ok = True
            break
        except Exception as _stream_err:
            if _ctx_attempt == _ctx_sizes[-1]:
                raise  # last attempt, propagate error
            _log(f"Ollama error z num_ctx={_ctx_attempt}: {_stream_err} — ponawiam z mniejszym kontekstem...")
            # Reset output file for retry
            out_path.write_text(prefix, encoding="utf-8")
            written = len(prefix.encode("utf-8"))

    dt = time.time() - t0
    state["stage"] = "finalize"
    state["progress"] = 98
    state["updated_at"] = now_iso()
    _persist_deep_task_state(project_id, {k: v for k, v in state.items() if not str(k).startswith("_")})
    _prog(98)

    # store last output
    try:
        (adir / "deep_last.md").write_text(_load_text_file(out_path), encoding="utf-8")
    except Exception:
        pass

    # Persist deep_latest.json
    try:
        content = _load_text_file(out_path)
        deep_payload = {
            "content": content,
            "meta": {
                "model": model,
                "generated_at": now_iso(),
                "generation_time": round(dt, 2),
                "template_ids": [str(x) for x in template_ids],
                "custom_prompt": str(custom_prompt).strip() if custom_prompt else None,
                "custom_prompt_use_templates": bool(custom_prompt_use_templates) if custom_prompt else None,
                "include_sources": include_sources,
                "stream": True,
            },
        }
        (adir / "deep_latest.json").write_text(json.dumps(deep_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass

    # finalize state
    state = {
        "task_id": task_id,
        "project_id": project_id,
        "model": model,
        "template_ids": [str(x) for x in template_ids],
        "status": "done",
        "stage": "done",
        "progress": 100,
        "started_at": state.get("started_at"),
        "finished_at": now_iso(),
        "generation_time": round(dt, 2),
    }
    _persist_deep_task_state(project_id, state)
    _prog(100)


def _deep_task_runner(
    *,
    task_id: str,
    proj_id: str,
    model: str,
    template_ids: List[str],
    custom_prompt: Optional[str],
    custom_prompt_use_templates: bool,
    include_sources: Dict[str, Any],
    result_mode: str = "replace",
    append_header: str = "",
    log_cb=None,
    progress_cb=None,
) -> None:
    """Sync wrapper for TASKS.start_python_fn."""
    asyncio.run(
        _run_deep_analysis_task(
            task_id=task_id,
            project_id=proj_id,
            model=model,
            template_ids=template_ids,
            custom_prompt=custom_prompt,
            custom_prompt_use_templates=custom_prompt_use_templates,
            include_sources=include_sources,
            result_mode=result_mode,
            append_header=append_header,
            log_cb=log_cb,
            progress_cb=progress_cb,
        )
    )


@app.get("/api/analysis/task_state/{project_id}")
async def api_analysis_task_state(project_id: str) -> Any:
    """Return deep-analysis task state for a project (if running)."""
    state = _read_json_file(_analysis_deep_task_state_path(project_id), None)
    if not isinstance(state, dict) or not state.get("task_id"):
        return {"status": "idle"}
    task_id = str(state.get("task_id"))
    task_obj = None
    try:
        task_obj = asdict(TASKS.get(task_id))
    except KeyError:
        task_obj = None
    except Exception:
        # best-effort
        task_obj = None
    return {"status": "ok", "state": state, "task": task_obj}


@app.get("/api/analysis/task_output/{project_id}")
async def api_analysis_task_output(project_id: str, request: Request) -> Any:
    """Return incremental deep-analysis output for a project.

    Query params:
      from: byte offset (default 0)
      max: max bytes to read (default 65536)
    """
    qp = request.query_params
    try:
        from_off = int(qp.get("from") or 0)
    except Exception:
        from_off = 0
    try:
        max_bytes = int(qp.get("max") or 65536)
    except Exception:
        max_bytes = 65536
    max_bytes = max(1024, min(max_bytes, 1024 * 1024))  # clamp 1KB..1MB

    path = _analysis_deep_task_output_path(project_id)
    if not path.exists():
        return {"chunk": "", "next": from_off, "eof": True}

    try:
        size = path.stat().st_size
        if from_off < 0:
            from_off = 0
        if from_off > size:
            from_off = size
        with path.open("rb") as f:
            f.seek(from_off)
            raw = f.read(max_bytes)
        nxt = from_off + len(raw)
        # decode best-effort; offset is in bytes so we may split a UTF-8 char
        chunk = raw.decode("utf-8", errors="ignore")
        eof = nxt >= size
        return {"chunk": chunk, "next": nxt, "eof": eof, "size": size}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Cannot read output: {e}")


@app.post("/api/analysis/start")
async def api_analysis_start(payload: Dict[str, Any] = Body(...)) -> Any:
    """Start deep analysis as a background task.

    This is more resilient than SSE streaming: switching tabs won't cancel the task.
    """
    project_id = str(payload.get("project_id") or "").strip()
    if not project_id:
        raise HTTPException(status_code=400, detail="project_id required")

    # If already running for this project, return existing task
    existing = _read_json_file(_analysis_deep_task_state_path(project_id), None)
    if isinstance(existing, dict) and existing.get("task_id"):
        ex_task_id = str(existing.get("task_id"))
        try:
            ex_task = TASKS.get(ex_task_id)
            if ex_task and ex_task.status in ("running", "queued"):
                return {"status": "already_running", "task_id": ex_task_id, "task": asdict(ex_task)}
        except KeyError:
            pass

    template_ids = payload.get("template_ids") or []
    if isinstance(template_ids, str):
        template_ids = [t.strip() for t in template_ids.split(",") if t.strip()]
    if not isinstance(template_ids, list):
        template_ids = []

    custom_prompt = payload.get("custom_prompt")
    use_templates = payload.get("custom_prompt_use_templates")
    if isinstance(use_templates, str):
        use_templates = use_templates.strip().lower() not in ("0", "false", "no", "off")
    elif isinstance(use_templates, bool):
        use_templates = bool(use_templates)
    else:
        use_templates = True
    if custom_prompt and not use_templates:
        template_ids = []

    model = str(payload.get("model") or _get_model_settings().get("deep") or DEFAULT_MODELS["deep"]).strip()
    include_sources = payload.get("include_sources") or {}
    if not isinstance(include_sources, dict):
        include_sources = {}

    # normalize sources flags (also keep selected documents list)
    docs = include_sources.get("documents", [])
    if isinstance(docs, str):
        docs = [x.strip() for x in docs.split(",") if x.strip()]
    if not isinstance(docs, list):
        docs = []
    docs = [str(x) for x in docs if x is not None and str(x).strip()]

    include_sources = {
        "transcript": bool(include_sources.get("transcript", False)),
        "diarization": bool(include_sources.get("diarization", False)),
        "notes_global": bool(include_sources.get("notes_global", False)),
        "notes_blocks": bool(include_sources.get("notes_blocks", False)),
        "documents": docs,
    }
    # Result mode: UI no longer exposes replace/append. Always append additional analysis
    # when a base deep result exists; otherwise it automatically falls back to replace.
    result_mode = "append"
    append_header = str(payload.get("append_header") or "")
    if len(append_header) > 5000:
        append_header = append_header[:5000]

    # Only allow append if there is an existing deep analysis
    if result_mode == "append":
        base_content = ""
        dl = _analysis_deep_latest_path(project_id)
        if dl.exists():
            j = _read_json_file(dl, None)
            if isinstance(j, dict):
                base_content = str(j.get("content") or "")
        if not base_content.strip():
            result_mode = "replace"
            append_header = ""
        elif not append_header:
            stamp = now_iso()[:16]
            append_header = f"\n\n## Dodatkowa analiza ({stamp}, model={model})\n\n"

    # Create a stable task id (also persisted to disk)
    task_id = uuid.uuid4().hex

    # Persist initial state so UI can immediately see it
    _persist_deep_task_state(project_id, {
        "task_id": task_id,
        "project_id": project_id,
        "model": model,
        "template_ids": [str(x) for x in template_ids],
        "custom_prompt": custom_prompt,
        "custom_prompt_use_templates": bool(use_templates),
        "include_sources": include_sources,
        "result_mode": result_mode,
        "append_header": append_header,
        "status": "queued",
        "stage": "queued",
        "progress": 0,
        "started_at": now_iso(),
        "updated_at": now_iso(),
    })

    # Start background worker (GPU RM managed if enabled)
    if GPU_RM.enabled:
        GPU_RM.enqueue_python_fn(
            "analysis",
            project_id,
            _deep_task_runner,
            task_id_override=task_id,
            task_id=task_id,
            proj_id=project_id,
            model=model,
            template_ids=[str(x) for x in template_ids],
            custom_prompt=str(custom_prompt).strip() if custom_prompt else None,
            custom_prompt_use_templates=bool(use_templates),
            include_sources=include_sources,
            result_mode=result_mode,
            append_header=append_header,
        )
    else:
        TASKS.start_python_fn(
            "analysis",
            project_id,
            _deep_task_runner,
            task_id_override=task_id,
            task_id=task_id,
            proj_id=project_id,
            model=model,
            template_ids=[str(x) for x in template_ids],
            custom_prompt=str(custom_prompt).strip() if custom_prompt else None,
            custom_prompt_use_templates=bool(use_templates),
            include_sources=include_sources,
            result_mode=result_mode,
            append_header=append_header,
        )

    return {"status": "started", "task_id": task_id}


@app.get("/api/analysis/stream")
async def api_analysis_stream(request: Request) -> Any:
    """Streaming deep analysis (SSE). Designed for EventSource usage."""
    qp = request.query_params
    project_id = str(qp.get("project_id") or "").strip()
    if not project_id:
        raise HTTPException(status_code=400, detail="project_id required")

    model = str(qp.get("model") or "").strip() or (_get_model_settings().get("deep") or DEFAULT_MODELS["deep"])
    tids = str(qp.get("template_ids") or "").strip()
    template_ids = [t.strip() for t in tids.split(",") if t.strip()]

    # Optional: custom prompt for deep analysis (can be combined with selected templates)
    custom_prompt = str(qp.get("custom_prompt") or "").strip()
    use_templates = True
    use_raw = str(qp.get("custom_prompt_use_templates") or "").strip().lower()
    if use_raw in ("0", "false", "no", "off"):  # explicit disable
        use_templates = False
    elif use_raw in ("1", "true", "yes", "on"):  # explicit enable
        use_templates = True
    if custom_prompt and not use_templates:
        # Ignore selected prompt templates when user wants custom-only instruction
        template_ids = []

    # Optional include_sources as JSON in query (best-effort)
    include_sources: Dict[str, Any] = {"transcript": True, "diarization": True, "notes_global": True, "notes_blocks": True, "documents": []}
    inc_raw = qp.get("include_sources")
    if inc_raw:
        try:
            obj = json.loads(str(inc_raw))
            if isinstance(obj, dict):
                include_sources.update(obj)
        except Exception:
            pass

    # Create a visible task (so the Logs tab can show progress for analysis).
    task = TASKS.create_task(kind="analysis", project_id=project_id)

    async def generate() -> Any:
        """SSE generator: emits chunks + progress info."""

        def _emit(*, chunk: str = "", done: bool = False, stage: str = "", progress: int | None = None) -> str:
            payload: Dict[str, Any] = {"chunk": chunk, "done": done, "task_id": task.task_id}
            if stage:
                payload["stage"] = stage
            if progress is not None:
                payload["progress"] = int(progress)
            return "data: " + json.dumps(payload, ensure_ascii=False) + "\n\n"

        # Helper: update task state
        def _set(stage: str, pct: int) -> None:
            task.progress = int(pct)
            task.status = "running" if pct < 100 else "done"
            if stage:
                task.add_log(f"{pct}% | {stage}")

        try:
            _set("Preparing prompts and sources", 2)
            yield _emit(stage="Przygotowanie…", progress=2)

            # 1) Build instruction from templates
            try:
                instruction = PROMPTS.build_combined_prompt(template_ids, custom_prompt if custom_prompt else None)
                for pid in template_ids:
                    try:
                        PROMPTS.bump_usage(str(pid))
                    except Exception:
                        pass
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Invalid prompt selection: {e}")

            _set("Gathering sources", 8)
            yield _emit(stage="Zbieranie źródeł…", progress=8)

            # 2) Gather sources text (transcript/diarization/notes/docs)
            sources_text = await _gather_analysis_sources(project_id, include_sources)
            if not sources_text.strip():
                raise HTTPException(status_code=400, detail="Brak źródeł do analizy")

            final_prompt = (instruction.strip() + "\n\n---\n\n" if instruction.strip() else "") + "# Materiał źródłowy\n\n" + sources_text

            # Keep a copy of the prompt used (for reproducibility)
            adir = _analysis_dir(project_id)
            ts = time.strftime("%Y%m%d_%H%M%S")
            (adir / "prompts" / f"stream_{ts}.md").write_text(final_prompt, encoding="utf-8")

            task.add_log(f"Model: {model}")
            task.add_log(f"Templates: {','.join(template_ids) if template_ids else '-'}")
            if custom_prompt:
                task.add_log(f"Custom prompt: {len(custom_prompt)} chars (use_templates={use_templates})")
            try:
                task.add_log(f"Sources: {json.dumps(include_sources, ensure_ascii=False)}")
            except Exception:
                pass

            # 3) Ensure model exists (auto-pull if missing)
            _set("Ensuring Ollama model (auto-pull if missing)", 15)
            yield _emit(stage="Sprawdzanie modelu Ollama…", progress=15)
            ensure_res = await OLLAMA.ensure_model(model)
            task.add_log(f"Ollama ensure_model: {ensure_res}")

            # 4) Stream generation
            _set("Generating", 20)
            yield _emit(stage="Generowanie…", progress=20)

            system_msg = "Jesteś ekspertem w analizie dokumentów i rozmów. Twórz raporty po polsku, jasno i strukturalnie."
            options = {"temperature": 0.7, "num_ctx": 32768}

            msgs = [{"role": "system", "content": system_msg}, {"role": "user", "content": final_prompt}]

            chunks = []

            chars = 0
            last_pct = 20
            last_log_bucket = 20
            async for chunk in OLLAMA.stream_chat(model=model, messages=msgs, options=options):
                if not isinstance(chunk, str) or not chunk:
                    continue
                chars += len(chunk)
                chunks.append(chunk)
                # Best-effort progress: grow with emitted characters, cap at 95%.
                pct = min(95, 20 + int(chars / 400))
                if pct != last_pct:
                    task.progress = int(pct)
                    last_pct = int(pct)
                # log every ~5%
                if last_pct - last_log_bucket >= 5:
                    last_log_bucket = last_pct
                    task.add_log(f"Progress: {last_pct}%")
                yield _emit(chunk=chunk, done=False, stage="Generowanie…", progress=last_pct)


            # Persist last deep analysis (stream) into project for reload/export/import.
            try:
                adir = _analysis_dir(project_id)
                deep_payload = {
                    "content": "".join(chunks),
                    "meta": {
                        "model": model,
                        "generated_at": now_iso(),
                        "template_ids": [str(x) for x in template_ids],
                        "custom_prompt": custom_prompt if custom_prompt else None,
                        "custom_prompt_use_templates": bool(custom_prompt_use_templates) if custom_prompt else None,
                        "include_sources": include_sources,
                        "stream": True,
                    },
                    "report": None,
                }
                (adir / "deep_latest.json").write_text(json.dumps(deep_payload, ensure_ascii=False, indent=2), encoding="utf-8")
                task.add_log(f"Saved deep_latest.json (model={model})")
            except Exception as e:
                task.add_log(f"WARNING: Failed to save deep_latest.json: {e}")

            task.progress = 100
            task.status = "done"
            task.add_log("Done")
            yield _emit(chunk="", done=True, stage="Zakończono.", progress=100)

        except HTTPException as e:
            task.status = "error"
            task.progress = max(task.progress, 100)
            task.error = str(e.detail)
            task.add_log(f"ERROR: {e.detail}")
            yield _emit(chunk=f"\n\n[ERROR] {e.detail}", done=True, stage="Błąd.", progress=task.progress)
        except Exception as e:
            task.status = "error"
            task.progress = max(task.progress, 100)
            task.error = str(e)
            task.add_log(f"ERROR: {e}")
            yield _emit(chunk=f"\n\n[ERROR] {e}", done=True, stage="Błąd.", progress=task.progress)

    return StreamingResponse(generate(), media_type="text/event-stream")


@app.get("/api/analysis/reports/{project_id}/list")
async def api_analysis_reports_list(project_id: str) -> Any:
    """List saved analysis reports (best-effort)."""
    rdir = _analysis_reports_dir(project_id)
    items: List[Dict[str, Any]] = []
    if not rdir.exists():
        return {"reports": []}

    for fp in sorted(rdir.glob("*"), key=lambda p: p.stat().st_mtime, reverse=True):
        if not fp.is_file():
            continue
        # hide metadata sidecars
        if fp.name.endswith(".meta.json"):
            continue
        if fp.suffix.lower() not in (".md", ".html", ".docx"):
            continue
        try:
            st = fp.stat()
            items.append({
                "filename": fp.name,
                "format": fp.suffix.lstrip(".").lower(),
                "size_bytes": st.st_size,
                "modified_at": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(st.st_mtime)),
                "download_url": f"/api/analysis/reports/{project_id}/download/{fp.name}",
            })
        except Exception:
            continue
    return {"reports": items}


@app.get("/api/analysis/reports/{project_id}/download/{filename}")
async def api_analysis_reports_download(project_id: str, filename: str) -> Any:
    rdir = _analysis_reports_dir(project_id)
    path = _safe_child_path(rdir, filename)
    require_existing_file(path, "Plik raportu nie istnieje.")
    return FileResponse(str(path), filename=path.name)


# --- Finance entity memory API ---

def _finance_memory(project_id: str):
    """Get EntityMemory instance for a project."""
    from backend.finance.entity_memory import EntityMemory
    finance_dir = project_path(project_id) / "finance"
    global_dir = PROJECTS_DIR / "_global"
    return EntityMemory(finance_dir, global_dir)


@app.get("/api/finance/entities/{project_id}")
def api_finance_entities_list(project_id: str, flagged: int = 0, entity_type: str = "") -> Any:
    """List known entities for a project."""
    mem = _finance_memory(project_id)
    return mem.list_entities(
        flagged_only=bool(flagged),
        entity_type=entity_type or None,
    )


@app.post("/api/finance/entities/{project_id}/flag")
def api_finance_entity_flag(project_id: str, payload: Dict[str, Any] = Body(...)) -> Any:
    """Flag/unflag an entity in the intelligence memory."""
    name = str(payload.get("name") or "").strip()
    if not name:
        return {"error": "name is required"}
    entity_type = str(payload.get("entity_type") or "").strip()
    notes = str(payload.get("notes") or "").strip()
    flagged = bool(payload.get("flagged", True))
    propagate = bool(payload.get("propagate_global", False))

    mem = _finance_memory(project_id)
    ent = mem.flag_entity(
        name=name,
        entity_type=entity_type,
        notes=notes,
        flagged=flagged,
        propagate_global=propagate,
    )
    return ent.to_dict()


@app.post("/api/finance/entities/{project_id}/unflag")
def api_finance_entity_unflag(project_id: str, payload: Dict[str, Any] = Body(...)) -> Any:
    """Remove flag from an entity."""
    name = str(payload.get("name") or "").strip()
    if not name:
        return {"error": "name is required"}
    mem = _finance_memory(project_id)
    ent = mem.unflag_entity(name)
    if ent:
        return ent.to_dict()
    return {"error": "entity not found"}


@app.delete("/api/finance/entities/{project_id}/{entity_name}")
def api_finance_entity_delete(project_id: str, entity_name: str) -> Any:
    """Delete an entity from project memory."""
    mem = _finance_memory(project_id)
    if mem.delete_entity(entity_name):
        return {"ok": True}
    return {"error": "entity not found"}


@app.get("/api/finance/parsed/{project_id}")
def api_finance_parsed(project_id: str) -> Any:
    """Get latest parsed finance data for a project."""
    finance_dir = project_path(project_id) / "finance" / "parsed"
    if not finance_dir.exists():
        return {"files": []}
    files = []
    for f in sorted(finance_dir.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
            files.append({
                "filename": f.name,
                "bank": data.get("bank", ""),
                "transactions": len(data.get("transactions", [])),
                "score": data.get("score", {}).get("total_score"),
            })
        except Exception:
            pass
    return {"files": files}


@app.post("/api/analysis/save")
async def api_analysis_save(payload: Dict[str, Any] = Body(...)) -> Any:
    """Save analysis content as TXT/HTML/DOC/DOCX/Markdown in projects/{id}/analysis/reports/."""
    project_id = str(payload.get("project_id") or "").strip()
    if not project_id:
        raise HTTPException(status_code=400, detail="project_id required")

    content = str(payload.get("content") or "")
    if not content.strip():
        raise HTTPException(status_code=400, detail="content required")

    output_format = str(payload.get("format") or payload.get("output_format") or "md").lower().strip()
    title = str(payload.get("title") or payload.get("report_title") or "").strip() or None
    model = str(payload.get("model") or "").strip() or (_get_model_settings().get("deep") or DEFAULT_MODELS["deep"])

    template_ids = payload.get("template_ids") or []
    if isinstance(template_ids, str):
        template_ids = [t.strip() for t in template_ids.split(",") if t.strip()]
    if not isinstance(template_ids, list):
        template_ids = []

    try:
        res = save_report(
            reports_dir=_analysis_reports_dir(project_id),
            content=content,
            output_format=output_format,
            title=title or ("_".join([str(x) for x in template_ids]) if template_ids else "Analiza"),
            template_ids=[str(x) for x in template_ids],
            project_id=project_id,
            model=model,
        )
    except ReportSaveError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Report save failed: {e}")

    download_url = f"/api/analysis/reports/{project_id}/download/{res.filename}"
    return {
        "status": "success",
        "report": {
            "filename": res.filename,
            "format": res.format,
            "size_bytes": res.size_bytes,
            "download_url": download_url,
        },
        "preview": content[:4000],
    }



# ---------- API: transcribe / diarize ----------

def _start_sound_detection(project_id: str, audio_path: str, model_id: str) -> Optional[str]:
    """Start sound detection as a CPU task, queued through GPU_RM when available.

    Returns task_id or None if model not installed.
    """
    # Check if model is installed
    reg = _read_sound_detection_registry()
    model_state = reg.get(model_id, {})
    if not model_state.get("downloaded"):
        app_log(f"Sound detection skipped: model '{model_id}' not installed")
        return None

    worker = ROOT / "backend" / "sound_detection_worker.py"
    if not worker.exists():
        app_log(f"Sound detection skipped: worker not found at {worker}")
        return None

    # Output file in project folder
    proj_dir = project_path(project_id)
    output_file = proj_dir / "sound_events.json"

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--action", "detect",
        "--model", model_id,
        "--audio", audio_path,
        "--output", str(output_file),
        "--threshold", "0.3",
    ]

    app_log(f"Sound detection requested: project_id={project_id}, model={model_id}")

    # Route through GPU_RM queue so CPU-only machines don't get overloaded
    # by concurrent detection + transcription/diarization processes.
    if GPU_RM.enabled:
        t = GPU_RM.enqueue_subprocess(kind="sound_detection", project_id=project_id, cmd=cmd, cwd=ROOT)
    else:
        t = TASKS.start_subprocess(kind="sound_detection", project_id=project_id, cmd=cmd, cwd=ROOT)
    return t.task_id


@app.post("/api/transcribe")
async def api_transcribe(
    project_id: str = Form(""),
    lang: str = Form("auto"),
    model: str = Form("large-v3"),
    asr_engine: str = Form("whisper"),
    audio: Optional[UploadFile] = File(None),
    sound_detection_enabled: int = Form(0),
    sound_detection_model: str = Form(""),
) -> Any:
    project_id = ensure_project(project_id or None)
    if audio is not None:
        audio_path = save_upload(project_id, audio)
    else:
        meta = read_project_meta(project_id)
        fname = str(meta.get("audio_file") or "")
        if not fname:
            raise HTTPException(status_code=400, detail="Brak pliku audio w projekcie. Wgraj plik.")
        audio_path = project_path(project_id) / fname
        require_existing_file(audio_path, "Plik audio z projektu nie istnieje.")

    # Use existing worker script (subprocess) to keep server responsive
    worker = ROOT / "backend" / "transcribe_worker.py"
    require_existing_file(worker, "Brak transcribe_worker.py")

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--audio", str(audio_path),
        "--engine", str(asr_engine or "whisper"),
        "--model", model,
        "--lang", lang,
    ]
    app_log(f"Transcription requested: project_id={project_id}, audio='{audio_path.name}', engine={asr_engine}, model={model}, lang={lang}")
    if GPU_RM.enabled:
        t = GPU_RM.enqueue_subprocess(kind="transcribe", project_id=project_id, cmd=cmd, cwd=ROOT)
    else:
        t = TASKS.start_subprocess(kind="transcribe", project_id=project_id, cmd=cmd, cwd=ROOT)

    # Run sound detection in parallel (on CPU) if enabled
    sound_task_id = None
    if sound_detection_enabled and sound_detection_model:
        sound_task_id = _start_sound_detection(project_id, str(audio_path), sound_detection_model)

    return {"task_id": t.task_id, "project_id": project_id, "sound_detection_task_id": sound_task_id}


@app.post("/api/diarize_voice")
async def api_diarize_voice(
    project_id: str = Form(""),
    diar_engine: str = Form("pyannote"),
    diar_model: str = Form(""),
    lang: str = Form("auto"),
    model: str = Form("large-v3"),
    asr_engine: str = Form("whisper"),
    audio: Optional[UploadFile] = File(None),
    sound_detection_enabled: int = Form(0),
    sound_detection_model: str = Form(""),
) -> Any:
    """Diarize voice (audio) using the selected diarization engine.

    NOTE: Engines/models must be installed in ASR Settings.
    """
    project_id = ensure_project(project_id or None)
    if audio is not None:
        audio_path = save_upload(project_id, audio)
    else:
        meta = read_project_meta(project_id)
        fname = str(meta.get("audio_file") or "")
        if not fname:
            raise HTTPException(status_code=400, detail="Brak pliku audio w projekcie. Wgraj plik.")
        audio_path = project_path(project_id) / fname
        require_existing_file(audio_path, "Plik audio z projektu nie istnieje.")

    diar_engine = str(diar_engine or "pyannote").strip().lower()
    asr_engine = str(asr_engine or "whisper").strip().lower()

    # Validate installed models to prevent auto-download outside ASR Settings
    reg = _read_asr_registry()
    if not reg:
        reg = scan_and_persist_asr_model_cache()

    def _reg_dict(key: str) -> dict:
        d = reg.get(key)
        return d if isinstance(d, dict) else {}

    reg_whisper = _reg_dict('whisper')
    reg_nemo = _reg_dict('nemo')
    reg_nemo_diar = _reg_dict('nemo_diar')
    reg_pyannote = _reg_dict('pyannote')

    # ASR model must be installed
    if asr_engine == 'whisper':
        if not reg_whisper.get(model):
            raise HTTPException(status_code=400, detail="Model Whisper nie jest zainstalowany. Zainstaluj go w Ustawieniach ASR.")
    elif asr_engine == 'nemo':
        if not reg_nemo.get(model):
            raise HTTPException(status_code=400, detail="Model NeMo ASR nie jest zainstalowany. Zainstaluj go w Ustawieniach ASR.")
    else:
        raise HTTPException(status_code=400, detail="Nieznany silnik ASR")

    # Diarization model must be installed (where applicable)
    if diar_engine == 'pyannote':
        if diar_model and not reg_pyannote.get(diar_model):
            raise HTTPException(status_code=400, detail="Model pyannote nie jest zainstalowany. Zainstaluj go w Ustawieniach ASR.")
    elif diar_engine == 'nemo_diar':
        if diar_model and not reg_nemo_diar.get(diar_model):
            raise HTTPException(status_code=400, detail="Model NeMo diarization nie jest zainstalowany. Zainstaluj go w Ustawieniach ASR.")
    else:
        raise HTTPException(status_code=400, detail="Nieznany silnik diaryzacji")

    # HF token is required only for pyannote
    s = load_settings()
    hf_token = getattr(s, "hf_token", "") or ""
    if diar_engine == 'pyannote' and not hf_token:
        raise HTTPException(status_code=400, detail="Brak tokena HF. Ustaw go w Ustawieniach.")

    worker = ROOT / "backend" / "voice_worker.py"
    require_existing_file(worker, "Brak voice_worker.py")

    cmd = [
        os.environ.get("PYTHON", __import__("sys").executable),
        str(worker),
        "--audio", str(audio_path),
        "--model", model,
        "--lang", lang,
        "--asr_engine", asr_engine,
        "--diar_engine", diar_engine,
        "--diar_model", (diar_model or ""),
        "--hf_token", hf_token,
    ]

    app_log(
        f"Voice diarization requested: project_id={project_id}, audio='{audio_path.name}', diar_engine={diar_engine}, diar_model={diar_model or '-'}, asr_engine={asr_engine}, model={model}, lang={lang}"
    )

    if GPU_RM.enabled:
        t = GPU_RM.enqueue_subprocess(kind="diarize_voice", project_id=project_id, cmd=cmd, cwd=ROOT)
    else:
        t = TASKS.start_subprocess(kind="diarize_voice", project_id=project_id, cmd=cmd, cwd=ROOT)

    # Run sound detection in parallel (on CPU) if enabled
    sound_task_id = None
    if sound_detection_enabled and sound_detection_model:
        sound_task_id = _start_sound_detection(project_id, str(audio_path), sound_detection_model)

    return {"task_id": t.task_id, "project_id": project_id, "sound_detection_task_id": sound_task_id}


@app.post("/api/diarize_text")
async def api_diarize_text(
    project_id: str = Form(""),
    text: str = Form(...),
    speakers: int = Form(2),
    method: str = Form("alternate"),
    mapping_json: str = Form(""),
) -> Any:
    project_id = ensure_project(project_id or None)
    mapping: Dict[str, str] = {}
    if mapping_json.strip():
        try:
            mapping = json.loads(mapping_json)
            if not isinstance(mapping, dict):
                mapping = {}
        except Exception:
            raise HTTPException(status_code=400, detail="Niepoprawny JSON mapowania mówców.")

    speakers = max(1, min(50, int(speakers or 2)))
    method = str(method or "alternate")

    app_log(f"Text diarization requested: project_id={project_id}, speakers={speakers}, method={method}, input_chars={len(text or '')}")

    def fn(txt: str, log_cb=None, progress_cb=None) -> Dict[str, Any]:
        res = diarize_text_simple(txt, speakers=speakers, method=method, log_cb=log_cb, progress_cb=progress_cb)
        out_text = str(res.get("text") or "")
        # Optional mapping: replace "SPK1:" -> "Jan:" etc
        if mapping:
            if log_cb:
                log_cb(f"Speaker mapping: applying {len(mapping)} replacements")
            for k, v in mapping.items():
                if not isinstance(k, str) or not isinstance(v, str):
                    continue
                # replace both "SPK1:" and "SPK1 " variants
                out_text = out_text.replace(f"{k}:", f"{v}:")
                out_text = out_text.replace(f"{k} ", f"{v} ")
        return {"kind": "diarized_text", "text": out_text}

    t = TASKS.start_python_fn("diarize_text", project_id, fn, text)
    return {"task_id": t.task_id, "project_id": project_id}



# ---------- post-processing: save default files when task done ----------
# lightweight background watcher: for finished tasks, persist outputs into project folder
# (keeps UI simple: download links point to stable filenames)
_persist_lock = threading.Lock()
_persisted: set[str] = set()


def _persist_task_outputs(t: "TaskState") -> None:
    """Persist finished task outputs into stable project files.
    We also update project.json flags so UI can show what exists.
    """
    pdir = project_path(t.project_id)
    meta = read_project_meta(t.project_id)
    wrote_any = False
    # Transcription
    if t.kind == "transcribe" and isinstance(t.result, dict):
        if ("text_ts" in t.result) or ("text" in t.result):
            out_txt = str(t.result.get("text_ts") or t.result.get("text") or "")
            (pdir / "transcript.txt").write_text(out_txt, encoding="utf-8")
            meta["has_transcript"] = True
            wrote_any = True
            # Save segments with confidence as JSON for persistence
            segments = t.result.get("segments")
            if segments and isinstance(segments, list):
                (pdir / "transcript_segments.json").write_text(
                    json.dumps(segments, indent=2, ensure_ascii=False), encoding="utf-8"
                )
            # Generate waveform peaks (best-effort, runs once per project)
            try:
                peaks_path = pdir / "peaks.json"
                if not peaks_path.exists():
                    _generate_waveform_peaks(t.project_id)
            except Exception:
                pass
            # Auto quick analysis (best-effort) similar to Whisper model auto-download.
            _schedule_quick_analysis_background(t.project_id)
    # Diarization
    if t.kind in ("diarize_voice", "diarize_text") and isinstance(t.result, dict):
        if "text" in t.result:
            text_out = str(t.result.get("text") or "")
            # Apply speaker_map if present (best-effort)
            mapping = meta.get("speaker_map") or {}
            if isinstance(mapping, dict) and mapping:
                for k, v in mapping.items():
                    if isinstance(k, str) and isinstance(v, str):
                        text_out = text_out.replace(k, v)
            (pdir / "diarized.txt").write_text(text_out, encoding="utf-8")
            meta["has_diarized"] = True
            wrote_any = True
            # Save segments with confidence as JSON for persistence
            segments = t.result.get("segments")
            if segments and isinstance(segments, list):
                (pdir / "diarized_segments.json").write_text(
                    json.dumps(segments, indent=2, ensure_ascii=False), encoding="utf-8"
                )
    if wrote_any:
        meta["updated_at"] = now_iso()
        write_project_meta(t.project_id, meta)
    else:
        # still bump timestamp for completed tasks
        meta["updated_at"] = now_iso()
        write_project_meta(t.project_id, meta)


def _schedule_quick_analysis_background(project_id: str) -> None:
    """Run quick analysis in a daemon thread if quick_summary is missing/outdated.

    If GPU RM is enabled, the job is queued as a python task so it won't overlap with transcription/diarization.
    """
    try:
        if not _get_analysis_settings().get("quick_enabled", True):
            return
    except Exception:
        pass

    try:
        pdir = project_path(project_id)
        transcript = pdir / "transcript.txt"
        if not transcript.exists():
            return
        qs = _analysis_dir(project_id) / "quick_summary.json"
        if qs.exists() and qs.stat().st_mtime >= transcript.stat().st_mtime:
            return
    except Exception:
        return

    # Avoid duplicate queued/running quick jobs for the same project
    try:
        for t in TASKS.list_tasks():
            if t.kind == "analysis_quick" and t.project_id == project_id and t.status in ("queued", "running"):
                return
    except Exception:
        pass

    if GPU_RM.enabled:
        try:
            GPU_RM.enqueue_python_fn("analysis_quick", project_id, _quick_analysis_task_runner, project_id)
            return
        except Exception as e:
            try:
                app_log(f"Auto quick analysis enqueue failed: project_id={project_id}, error={e}")
            except Exception:
                pass

    def _runner() -> None:
        try:
            asyncio.run(_run_quick_analysis(project_id))
            app_log(f"Auto quick analysis completed: project_id={project_id}")
        except Exception as e:
            # keep it silent-ish: quick analysis is optional
            try:
                app_log(f"Auto quick analysis skipped/failed: project_id={project_id}, error={e}")
            except Exception:
                pass

    threading.Thread(target=_runner, daemon=True).start()


# Chat, Admin and Tasks endpoints are in webapp/routers/ (chat.py, admin.py, tasks.py)


def _persist_loop() -> None:

    while True:
        time.sleep(1.0)
        tasks = TASKS.list_tasks()
        for t in tasks:
            if t.status != "done":
                continue
            key = t.task_id
            with _persist_lock:
                if key in _persisted:
                    continue
                _persisted.add(key)
            try:
                _persist_task_outputs(t)
            except Exception as e:
                # keep server running but leave a breadcrumb in task logs
                try:
                    t.add_log(f"PERSIST ERROR: {e}")
                except Exception:
                    pass

threading.Thread(target=_persist_loop, daemon=True).start()